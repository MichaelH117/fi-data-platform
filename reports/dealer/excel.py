"""Small openpyxl writer shared by the dealer and claims report packages.

Every workbook gets a frozen, bold header row, an auto-filter and sensible column widths. Pipeline
metadata columns (leading underscore) are never exported, and timezone-aware timestamps are made
naive because Excel cannot store them.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from pipeline import workpool as jobs


def clean(df: pd.DataFrame) -> pd.DataFrame:
    out = df[[c for c in df.columns if not str(c).startswith("_")]].copy()
    for c in out.columns:
        if isinstance(out[c].dtype, pd.DatetimeTZDtype):
            out[c] = out[c].dt.tz_localize(None)
    return out


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame], notes: list[str] | None = None) -> Path:
    """Write ``sheets`` ({sheet name: frame}) to ``path``. ``notes`` become a leading About sheet.

    Written in the background when ``python -m reports`` has a workbook pool open (pipeline/workpool.py)."""
    jobs.run(_write_workbook, path, sheets, notes)
    return path


def _write_workbook(path: Path, sheets: dict[str, pd.DataFrame], notes: list[str] | None = None) -> Path:
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        if notes:
            pd.DataFrame({"About this workbook": notes}).to_excel(xw, sheet_name="About", index=False)
        for name, df in sheets.items():
            clean(df).to_excel(xw, sheet_name=name[:31], index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            if ws.max_row > 1 and ws.title != "About":
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 60)
    return path
