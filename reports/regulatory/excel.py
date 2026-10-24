"""Excel workbooks (openpyxl, write-only so big detail tabs stay fast).

Funded-reserve filing, one workbook per filing state:
    Summary (Section 2 business activity, Section 3 computation, notes), ITD Contracts,
    Period Contracts, ITD Claims, Period Claims.
Market survey, one workbook per survey state:
    one tab per provider (numbered items x calendar year + Total), Contracts, Cancels, Claims.
"""

from __future__ import annotations

import datetime as dt
import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from pipeline import workpool as jobs
from reports.regulatory.data import COUNT_LINES, SUMMARY_COLUMN, filing_lines, line_notes, split_notes

HEAD_FILL = PatternFill("solid", start_color="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
MONEY = '#,##0.00;[Red]-#,##0.00'
COUNT = '#,##0'
DATE = "yyyy-mm-dd"
WRAP = Alignment(wrap_text=True, vertical="top")
SHEET_NAME_MAX = 31


def _py(value):
    """numpy / pandas scalars -> plain Python values openpyxl can write (NaN and NaT -> empty)."""
    if value is None or value is pd.NaT or value is pd.NA:
        return None
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    return value


def _cell(ws, value, font=None, fill=None, fmt=None, align=None) -> WriteOnlyCell:
    c = WriteOnlyCell(ws, value=_py(value))
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    return c


def _header(ws, labels: list[str]) -> list:
    return [_cell(ws, x, HEAD_FONT, HEAD_FILL) for x in labels]


def _fmt(col: str, value) -> str | None:
    if isinstance(value, dt.date):
        return DATE
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return COUNT if col.lower().endswith("count") else MONEY
    return None


def detail_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    """One flat detail tab with a styled header, frozen header row and a totals row."""
    ws = wb.create_sheet(title[:SHEET_NAME_MAX])
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns):
        ws.column_dimensions[_col_letter(i)].width = max(10, min(32, len(col) + 4))
    ws.append(_header(ws, list(df.columns)))
    for row in df.itertuples(index=False):
        ws.append([_cell(ws, v, fmt=_fmt(c, _py(v))) for c, v in zip(df.columns, row, strict=True)])
    totals = []
    for i, col in enumerate(df.columns):
        if i == 0:
            totals.append(_cell(ws, f"Total ({len(df):,} rows)", BOLD))
        elif pd.api.types.is_numeric_dtype(df[col]) and not pd.api.types.is_bool_dtype(df[col]):
            totals.append(_cell(ws, float(df[col].sum()), BOLD, fmt=COUNT if col.lower().endswith("count") else MONEY))
        else:
            totals.append(_cell(ws, None))
    ws.append(totals)


def _col_letter(i: int) -> str:
    s, i = "", i + 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def write_filing(path: Path, state: str, state_name: str, summary: pd.DataFrame, definitions: pd.DataFrame,
                 tabs: dict[str, pd.DataFrame], as_of: pd.Timestamp) -> Path:
    """Funded-reserve filing workbook for one state. ``summary`` = that state's summary rows.

    Written in the background when ``python -m reports`` has a workbook pool open (pipeline/workpool.py)."""
    jobs.run(_write_filing, path, state, state_name, summary, definitions, tabs, as_of)
    return path


def _write_filing(path: Path, state: str, state_name: str, summary: pd.DataFrame, definitions: pd.DataFrame,
                  tabs: dict[str, pd.DataFrame], as_of: pd.Timestamp) -> Path:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Summary")
    for col, width in zip("ABCDEFG", (8, 44, 20, 20, 20, 20, 80), strict=True):
        ws.column_dimensions[col].width = width
    first = summary.iloc[0]
    status = "FINAL" if bool(first.is_final) else "PRELIMINARY"
    period = summary[summary.scope == "PERIOD"].iloc[0]
    ws.append([_cell(ws, f"Funded-reserve filing: {state_name} ({state})", TITLE_FONT)])
    ws.append([_cell(ws, f"Regulatory year {period.period_label}; ITD = inception to {period.period_end:%b %d, %Y}. "
                         f"Status {status}. Method {first.method_code} at {first.reserve_rate:.0%}; "
                         f"security deposit {first.security_deposit_rate:.0%}. Data as of {as_of:%b %d, %Y}. "
                         "Synthetic data.")])
    s2, s3 = filing_lines(definitions, 2), filing_lines(definitions, 3)
    notes = line_notes(definitions)
    for provider, rows in summary.groupby("provider_name", sort=True):
        def pick(scope: str, basis: str, rows=rows) -> pd.Series:
            return rows[(rows.scope == scope) & (rows.basis == basis)].iloc[0]

        default = str(rows.loc[rows.is_default_basis, "basis"].iloc[0])
        other = "RETAIL" if default == "COVERAGE_COST" else "COVERAGE_COST"
        itd, per = pick("ITD", default), pick("PERIOD", default)
        itd_o, per_o = pick("ITD", other), pick("PERIOD", other)
        nice = {"COVERAGE_COST": "coverage cost", "RETAIL": "retail"}
        ws.append([])
        ws.append([_cell(ws, f"Provider: {provider}", BOLD)])
        ws.append([_cell(ws, "Section 2: Business activity", BOLD)])
        # the per-line note sits in column G in both sections, so notes line up down the sheet
        ws.append(_header(ws, ["Line", "Item", "ITD", "Period", "", "", "Note"]))
        for no, item, label in s2:
            col = SUMMARY_COLUMN[item]
            fmt = COUNT if item in COUNT_LINES else MONEY
            ws.append([_cell(ws, no), _cell(ws, label), _cell(ws, float(itd[col]), fmt=fmt),
                       _cell(ws, float(per[col]), fmt=fmt), _cell(ws, None), _cell(ws, None),
                       _cell(ws, notes.get(item, ""), align=WRAP)])
        ws.append([_cell(ws, "Section 3: Computation", BOLD)])
        ws.append(_header(ws, ["Line", "Item", f"ITD ({nice[default]})", f"Period ({nice[default]})",
                               f"ITD ({nice[other]})", f"Period ({nice[other]})", "Note"]))
        for no, item, label in s3:
            col = SUMMARY_COLUMN[item]
            fmt = COUNT if item in COUNT_LINES else MONEY
            ws.append([_cell(ws, no), _cell(ws, label)]
                      + [_cell(ws, float(r[col]), fmt=fmt) for r in (itd, per, itd_o, per_o)]
                      + [_cell(ws, notes.get(item, ""), align=WRAP)])
        ws.append([_cell(ws, "Notes", BOLD)])
        common, diffs = split_notes([itd, per, itd_o, per_o])
        ws.append([_cell(ws, "All"), _cell(ws, "; ".join(common), align=WRAP)])
        for scope, basis, extra in diffs:
            ws.append([_cell(ws, f"{scope} {nice[basis]}"), _cell(ws, "; ".join(extra), align=WRAP)])
    for title, df in tabs.items():
        detail_sheet(wb, title, df)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def survey_matrix(rows: pd.DataFrame) -> pd.DataFrame:
    """One provider's survey: item rows x year columns + Total (SUM, or the latest year for LAST items)."""
    years = rows.drop_duplicates("calendar_year").sort_values("calendar_year")
    m = rows.pivot_table(index=["item_no", "item_label", "aggregation"], columns="calendar_year", values="value",
                         aggfunc="sum").reset_index()
    ycols = list(years.calendar_year)
    m["Total"] = [r[ycols].sum() if r["aggregation"] == "SUM" else r[ycols[-1]] for _, r in m.iterrows()]
    m = m.rename(columns=dict(zip(ycols, years.year_label, strict=True)))
    return m.drop(columns="aggregation").rename(columns={"item_no": "Item", "item_label": "Question"})


def write_survey(path: Path, state: str, state_name: str, survey: pd.DataFrame, details: dict[str, pd.DataFrame],
                 as_of: pd.Timestamp, notes: dict[str, str] | None = None) -> Path:
    """Market-survey workbook for one state (in the background when a workbook pool is open)."""
    jobs.run(_write_survey, path, state, state_name, survey, details, as_of, notes)
    return path


def _write_survey(path: Path, state: str, state_name: str, survey: pd.DataFrame, details: dict[str, pd.DataFrame],
                  as_of: pd.Timestamp, notes: dict[str, str] | None = None) -> Path:
    wb = Workbook(write_only=True)
    for provider, rows in survey.groupby("provider_name", sort=True):
        m = survey_matrix(rows)
        ws = wb.create_sheet(str(provider)[:SHEET_NAME_MAX])
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 40
        for i in range(2, len(m.columns)):
            ws.column_dimensions[_col_letter(i)].width = 16
        ws.append([_cell(ws, f"Market survey: {state_name} ({state}), {provider}", TITLE_FONT)])
        ws.append([_cell(ws, f"Calendar years; the current year is year-to-date through the last closed posting "
                             f"period. "
                             f"Data as of {as_of:%b %d, %Y}. Synthetic data.")])
        ws.append(_header(ws, [str(c) for c in m.columns]))
        money_items = {"REVENUE_COVERAGE_COST", "REVENUE_PURCHASE_PRICE", "CLAIMS_PAID_AMOUNT"}
        items = rows.drop_duplicates("item_no").set_index("item_no")["item"]
        for r in m.itertuples(index=False):
            fmt = MONEY if items[r[0]] in money_items else COUNT
            ws.append([_cell(ws, r[0]), _cell(ws, r[1])] + [_cell(ws, float(v), fmt=fmt) for v in r[2:]])
        noted = [(no, notes[it]) for no, it in items.items() if notes and notes.get(it)]
        if noted:
            ws.append([])
            ws.append([_cell(ws, "Notes", TITLE_FONT)])
            for no, text in noted:
                ws.append([_cell(ws, int(no)), _cell(ws, text, align=WRAP)])
    for title, df in details.items():
        detail_sheet(wb, title, df)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
