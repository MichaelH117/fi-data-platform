"""Finance marts, seeds and report pages (Financial Analysts II and I).

Three layers of checks:
1. Seed and SQL rules that need nothing built (fast).
2. A hand-built mini ``core`` warehouse in ``tmp_path``: dbt builds only the finance seeds, marts and
   tests on it, and every statement value is compared with a number worked out by hand below.
   An injected cost-split break must make the close control fail.
3. On the real warehouse (skipped when it is absent): the finance marts tie to the core facts.
"""

from __future__ import annotations

import csv
import json
import os
import re
import subprocess
import sys
from html.parser import HTMLParser
from pathlib import Path

import duckdb
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
DBT_DIR = ROOT / "dbt"
SEEDS = DBT_DIR / "seeds" / "finance"
MARTS = DBT_DIR / "models" / "marts" / "finance"
DBT_TESTS = DBT_DIR / "tests" / "finance"
WAREHOUSE = Path(os.environ.get("WAREHOUSE_PATH", ROOT / "data" / "warehouse" / "warehouse.duckdb"))
DISB_TYPES = {"ADMIN_FEE", "INSURER_FEE", "PREMIUM_TAX", "COMMISSION", "CLAIMS_RESERVE", "PROFIT_SHARE_RESERVE"}
STRUCTURES = {"DEALER_OWNED_OBLIGOR", "CFC", "NCFC", "PORC", "RETRO", "PROFIT_SHARE"}
REQUIRED_MARTS = {"mart_pnl_by_product", "mart_profit_share_statement", "mart_commission_statement",
                  "mart_cession_experience", "mart_reserve_rollforward", "mart_remittance_late_fees",
                  "mart_disbursement_control_totals", "mart_funded_reserve_reconciliation",
                  # V6 (docs/V3_DESIGN.md section 9)
                  "mart_participation_contract", "mart_experience", "mart_new_reserves_by_term",
                  "mart_reserve_bridge", "mart_commission_hierarchy", "mart_rate_card", "mart_rate_card_exceptions",
                  "mart_partner_settlement", "mart_remit_register",
                  # V7 (docs/V3_DESIGN.md 10.4)
                  "mart_portfolio_facts", "mart_dealer_count_reconciliation", "mart_cancel_mileage"}
# V6 interim staging reads: core has no rate dimension and no contract rate id yet (DE Manager). Remove an
# entry when core carries it; nothing else in finance may read staging.
INTERIM_STAGING = {"mart_rate_card": {"stg_admin_core__rate", "stg_admin_core__rate_component"},
                   "mart_rate_card_exceptions": {"stg_admin_core__contract"}}


def _seed(name: str) -> list[dict]:
    with open(SEEDS / f"{name}.csv", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# --------------------------------------------------------------------------- 1. rules (no build)
def test_every_required_mart_exists_and_is_documented():
    sql = {p.stem for p in MARTS.glob("*.sql")}
    assert REQUIRED_MARTS <= sql
    yml = (MARTS / "_finance.yml").read_text(encoding="utf-8")
    for name in REQUIRED_MARTS:
        assert f"name: {name}" in yml, f"{name} missing from _finance.yml"


def test_structure_terms_cover_every_participation_structure():
    rows = _seed("participation_structure_terms")
    assert {r["structure_type"] for r in rows} == STRUCTURES
    for r in rows:
        assert 0 <= float(r["default_dealer_share_pct"]) <= 1
        assert 0 <= float(r["ceding_fee_pct"]) < 0.10
        assert int(r["min_seasoning_months"]) >= 0
        assert r["owns_claims_reserve"] in ("true", "false")
    # Only the plain profit-share structure leaves the reserve with the obligor.
    assert {r["structure_type"] for r in rows if r["owns_claims_reserve"] == "false"} == {"PROFIT_SHARE"}


def test_pnl_map_covers_every_cost_split_bucket_once():
    rows = _seed("pnl_line_map")
    assert sorted(r["disbursement_type"] for r in rows) == sorted(DISB_TYPES)
    assert [r["disbursement_type"] for r in rows if r["pnl_group"] == "RESERVE"] == ["CLAIMS_RESERVE"]


def test_aging_buckets_are_contiguous_and_cover_every_day():
    rows = sorted(_seed("remit_aging_buckets"), key=lambda r: int(r["sort_order"]))
    assert int(rows[0]["min_days"]) == 0
    for prev, nxt in zip(rows, rows[1:], strict=False):
        assert int(nxt["min_days"]) == int(prev["max_days"]) + 1, (prev, nxt)
    assert int(rows[-1]["max_days"]) >= 10_000


def test_finance_sql_reads_only_core_and_finance_seeds():
    """Marts build on core (where demo/test rows are already excluded), never on raw/staging/intermediate."""
    allowed = {p.stem for p in SEEDS.glob("*.csv")} | {p.stem for p in MARTS.glob("*.sql")}
    core = {p.stem for p in (DBT_DIR / "models" / "marts" / "core").glob("*.sql")}
    for path in [*MARTS.glob("*.sql"), *DBT_TESTS.glob("*.sql")]:
        text = path.read_text(encoding="utf-8")
        assert "source(" not in text, path.name
        for ref in re.findall(r"ref\('([^']+)'\)", text):
            if ref in INTERIM_STAGING.get(path.stem, set()):
                continue
            assert ref in allowed or ref in core, f"{path.name} refs {ref}"
            assert not ref.startswith(("stg_", "int_")), f"{path.name} refs {ref}"


def test_reports_read_only_finance_marts_and_seeds():
    for path in (ROOT / "reports" / "finance").glob("*.py"):
        text = path.read_text(encoding="utf-8")
        schemas = set(re.findall(r"\b(?:from|join)\s+([a-z_]+)\.[a-z_]+\b(?!\s+import)", text))
        assert schemas <= {"finance", "seeds"}, (path.name, schemas)


def test_term_year_buckets_are_contiguous_and_cover_every_term():
    rows = sorted(_seed("term_year_buckets"), key=lambda r: int(r["sort_order"]))
    assert int(rows[0]["min_months"]) == 0 and int(rows[-1]["max_months"]) >= 999
    for prev, nxt in zip(rows, rows[1:], strict=False):
        assert int(nxt["min_months"]) == int(prev["max_months"]) + 1, (prev, nxt)
    assert [r["bucket_label"] for r in rows[:2]] == ["0-1", "1-2"]


def test_commission_targets_cover_every_mapped_district_once():
    rows = _seed("commission_volume_targets")
    keys = [(r["district"], r["product_type"]) for r in rows]
    assert len(keys) == len(set(keys))
    assert all(int(r["monthly_target_contracts"]) >= 0 for r in rows)
    hierarchy = DBT_DIR / "seeds" / "core" / "agent_territory_hierarchy.csv"
    if hierarchy.exists():  # district -> region lives in core (DE Manager); every district has a target
        with open(hierarchy, newline="", encoding="utf-8") as f:
            districts = {r["district"] for r in csv.DictReader(f)}
        assert districts <= {r["district"] for r in rows}


def test_money_cents_puts_the_sign_before_the_symbol():
    from reports.build import money
    from reports.finance import money_cents

    assert money_cents(-1082.154) == "-$1,082.15"
    assert money_cents(99050.15) == "$99,050.15"
    assert money_cents(-0.001) == "$0.00"
    assert money_cents(None) == "n/a"
    assert money(-5689) == "-$5,689" and money(-98_000) == "-$98K"


def test_excel_frame_rounds_money_and_drops_hash_keys():
    from reports.finance import excel_format, excel_frame

    df = pd.DataFrame({"participation_entity_key": ["a" * 32], "posting_period_key": [202609],
                       "entity_name": ["E"], "entity_id": [7], "itd_case_reserve": [-4.547e-13],
                       "itd_earned_reserve": [193853.3523527172], "itd_loss_ratio": [0.512345678],
                       "contracts_written": [12.0]})
    out = excel_frame(df)
    assert list(out.columns[:3]) == ["entity_id", "entity_name", "posting_period_key"]
    assert "participation_entity_key" not in out.columns
    assert out["itd_earned_reserve"].iloc[0] == 193853.35
    assert out["itd_case_reserve"].iloc[0] == 0 and str(out["itd_case_reserve"].iloc[0]) == "0.0"
    assert excel_format("itd_earned_reserve", out["itd_earned_reserve"]) == "#,##0.00;-#,##0.00"
    assert excel_format("itd_loss_ratio", out["itd_loss_ratio"]) == "0.0%"
    assert excel_format("contracts_written", out["contracts_written"]) == "#,##0"
    assert excel_format("entity_id", out["entity_id"]) is None


def test_period_helpers():
    from reports.finance import period_label, shift_period

    assert shift_period(202401, -11) == 202302
    assert shift_period(202312, 1) == 202401
    assert period_label(202610) == "Oct 2026"
    assert period_label(None) == "n/a"


# --------------------------------------------------------------------------- 2. mini warehouse
# A tiny core star schema, hand-made so every expected number can be worked out on paper.
# Entity E1 (CFC, dealer D1, share from seed = 100%, ceding fee 2%, seasoning 12 months), contract C1:
#   reserve written 1000 (202301), earned 600 (202306) + 200 (202401), claim paid 300 + case 100 (202307).
#   ITD at 202402: earned 800, incurred 400, loss ratio 0.5, reserve balance 1000 - 300 = 700,
#   unearned 200, ceding fee 16, UW result 384, seasoned from 202401 (12 months) -> due 384.
# Entity E2 (PROFIT_SHARE, dealer D2, share 50%, seasoning 24 months), contract C2:
#   reserve 2000 (202301) - 500 cancel (202310), earned 1000, paid 900 -> UW 100, never seasoned -> due 0.
# Contract C3 (dealer D3, no participation, GAP) only shows up in the P&L / roll-forward / cash.
# Claim CL4 has no contract (contract_key / product_key '-1'): a reconciling P&L line, not a product.
# Agent A1: standard commission 150 on C1; an advance of 80 on C2 that earns down 25% (20) in 202306,
#   then C2 cancels in 202310 (chargebacks -30 and -20); the rest of the advance (60) is released then.
#   The cost split carries the same commission (150 and 80 at sale, -50 on the cancel).
# C3 is a legacy contract: 10 of cost-split commission and no agent-ledger line (the reconciling row),
#   earning 2.5 on the commission curve in 202312.
# A DENIED reinstatement must not add back any refund.
CORE_DDL = """
create schema core;
create table core.dim_posting_period as
    select cast(strftime(d, '%Y%m') as integer) as posting_period_key, cast(d as date) as period_start,
        cast(last_day(d) as date) as period_end, d < date '2024-02-01' as is_closed
    from range(date '2023-01-01', date '2024-03-01', interval 1 month) as t(d);
create table core.dim_dealer (dealer_key varchar, dealer_number varchar, dealer_name varchar,
    participation_entity_key varchar);
insert into core.dim_dealer values ('D1', 'DLR001', 'Summit Motors', 'E1'),
    ('D2', 'DLR002', 'Lakeside Auto', 'E2'), ('D3', 'DLR003', 'Harbor Cars', null);
create table core.dim_contract (contract_key varchar, dealer_key varchar);
insert into core.dim_contract values ('C1', 'D1'), ('C2', 'D2'), ('C3', 'D3');
create table core.dim_product (product_key varchar, category_code varchar, category varchar,
    subcategory_code varchar, subcategory varchar, product_code varchar);
insert into core.dim_product values ('P1', 'VSC', 'Vsc', 'VSC-PT', 'Powertrain', 'VSC-PT'),
    ('P2', 'GAP', 'Gap', 'GAP', 'GAP', 'GAP'), ('-1', 'UNKNOWN', 'Unknown', 'UNKNOWN', 'Unknown', 'UNKNOWN');
create table core.dim_participation_entity (participation_entity_key varchar, entity_id bigint,
    entity_name varchar, entity_type varchar, dealer_key varchar, tenant_key bigint, reinsurer_name varchar,
    effective_date date, profit_share_pct double);
insert into core.dim_participation_entity values
    ('E1', 1, 'Summit Reinsurance Co', 'CFC', 'D1', 1, 'Reinsurer One', date '2022-01-01', null),
    ('E2', 2, 'Lakeside Profit Share Co', 'PROFIT_SHARE', 'D2', 1, null, date '2022-06-01', 0.5);
create table core.dim_account_holder (account_holder_key varchar, holder_type varchar, ref_id bigint);
insert into core.dim_account_holder values ('AH_E1', 'DEALER_RESERVE', 1), ('AH_OB', 'OBLIGOR', 9),
    ('AH_AD', 'ADMIN', 1);
create table core.fct_transaction (txn_key varchar, source_system varchar, contract_key varchar,
    txn_type varchar, posting_period_key integer, amount double);
insert into core.fct_transaction values
    ('T1', 'admin_core', 'C1', 'SALE', 202301, 1300), ('T2', 'admin_core', 'C2', 'SALE', 202301, 2300),
    ('T3', 'admin_core', 'C2', 'CANCEL', 202310, -550), ('T4', 'legacy_admin', 'C3', 'SALE', 202302, 440),
    ('T5', 'admin_core', 'C1', 'CASH', 202402, 1300);
create table core.fct_disbursement (txn_key varchar, contract_key varchar, dealer_key varchar,
    product_key varchar, tenant_key bigint, account_holder_key varchar, source_system varchar,
    txn_type varchar, disbursement_type varchar, posting_period_key integer, amount double);
insert into core.fct_disbursement values
    ('T1', 'C1', 'D1', 'P1', 1, 'AH_E1', 'admin_core', 'SALE', 'CLAIMS_RESERVE', 202301, 1000),
    ('T1', 'C1', 'D1', 'P1', 1, 'AH_AD', 'admin_core', 'SALE', 'ADMIN_FEE', 202301, 100),
    ('T1', 'C1', 'D1', 'P1', 1, '-1', 'admin_core', 'SALE', 'COMMISSION', 202301, 150),
    ('T1', 'C1', 'D1', 'P1', 1, 'AH_E1', 'admin_core', 'SALE', 'PROFIT_SHARE_RESERVE', 202301, 50),
    ('T2', 'C2', 'D2', 'P1', 1, 'AH_OB', 'admin_core', 'SALE', 'CLAIMS_RESERVE', 202301, 2000),
    ('T2', 'C2', 'D2', 'P1', 1, 'AH_AD', 'admin_core', 'SALE', 'ADMIN_FEE', 202301, 120),
    ('T2', 'C2', 'D2', 'P1', 1, '-1', 'admin_core', 'SALE', 'COMMISSION', 202301, 80),
    ('T2', 'C2', 'D2', 'P1', 1, 'AH_E1', 'admin_core', 'SALE', 'PROFIT_SHARE_RESERVE', 202301, 100),
    ('T3', 'C2', 'D2', 'P1', 1, 'AH_OB', 'admin_core', 'CANCEL', 'CLAIMS_RESERVE', 202310, -500),
    ('T3', 'C2', 'D2', 'P1', 1, '-1', 'admin_core', 'CANCEL', 'COMMISSION', 202310, -50),
    ('T4', 'C3', 'D3', 'P2', 1, '-1', 'legacy_admin', 'SALE', 'CLAIMS_RESERVE', 202302, 400),
    ('T4', 'C3', 'D3', 'P2', 1, '-1', 'legacy_admin', 'SALE', 'ADMIN_FEE', 202302, 30),
    ('T4', 'C3', 'D3', 'P2', 1, '-1', 'legacy_admin', 'SALE', 'COMMISSION', 202302, 10);
create table core.fct_earned_disbursement (contract_key varchar, dealer_key varchar, product_key varchar,
    tenant_key bigint, disbursement_type varchar, posting_period_key integer, earned_amount double,
    earned_reserve double, earned_pct double, earned_pct_cumulative double);
insert into core.fct_earned_disbursement values
    ('C1', 'D1', 'P1', 1, 'CLAIMS_RESERVE', 202306, 600, 600, 0.6, 0.6),
    ('C1', 'D1', 'P1', 1, 'CLAIMS_RESERVE', 202401, 200, 200, 0.2, 0.8),
    ('C1', 'D1', 'P1', 1, 'ADMIN_FEE', 202306, 50, 0, 0.5, 0.5),
    ('C2', 'D2', 'P1', 1, 'CLAIMS_RESERVE', 202306, 1000, 1000, 0.5, 0.5),
    ('C2', 'D2', 'P1', 1, 'COMMISSION', 202306, 20, 0, 0.25, 0.25),
    ('C3', 'D3', 'P2', 1, 'CLAIMS_RESERVE', 202312, 100, 100, 0.25, 0.25),
    ('C3', 'D3', 'P2', 1, 'COMMISSION', 202312, 2.5, 0, 0.25, 0.25);
create table core.fct_claim_payment_line (claim_key varchar, contract_key varchar, product_key varchar,
    tenant_key bigint, posting_period_key integer, paid_amount double, pending_amount double,
    incurred_change double);
insert into core.fct_claim_payment_line values
    ('CL1', 'C1', 'P1', 1, 202307, 300, 100, 400), ('CL2', 'C2', 'P1', 1, 202308, 900, 0, 900),
    ('CL3', 'C3', 'P2', 1, 202312, 50, 0, 50), ('CL4', '-1', '-1', 1, 202312, 25, 0, 25);
create table core.fct_gross_premium (contract_key varchar, product_key varchar, tenant_key bigint,
    posting_period_key integer, contracts_written integer, coverage_cost double, retail_price double);
insert into core.fct_gross_premium values ('C1', 'P1', 1, 202301, 1, 1300, 2000),
    ('C2', 'P1', 1, 202301, 1, 2300, 3000), ('C3', 'P2', 1, 202302, 1, 440, 600);
create table core.fct_cancellation (contract_key varchar, product_key varchar, tenant_key bigint,
    posting_period_key integer, is_effective_cancel boolean, coverage_cost_refund double,
    customer_refund double);
insert into core.fct_cancellation values ('C2', 'P1', 1, 202310, true, 550, 600);
create table core.fct_reinstatement (contract_key varchar, product_key varchar, tenant_key bigint,
    posting_period_key integer, status varchar, reversal_amount double);
insert into core.fct_reinstatement values ('C2', 'P1', 1, 202312, 'DENIED', 500);
create table core.dim_agent (agent_key varchar, agent_code varchar, agent_name varchar, territory varchar);
insert into core.dim_agent values ('A1', 'AG001', 'Agent One', 'North');
create table core.fct_commission (agent_key varchar, tenant_key bigint, contract_key varchar,
    posting_period_key integer, amount double, is_advance boolean);
insert into core.fct_commission values ('A1', 1, 'C1', 202301, 150, false), ('A1', 1, 'C2', 202301, 80, true),
    ('A1', 1, 'C2', 202310, -30, false), ('A1', 1, 'C2', 202310, -20, true);
create table core.fct_cash_transaction (dealer_key varchar, tenant_key bigint, posting_period_key integer,
    amount double, remit_age_days integer, is_late boolean, late_fee double);
insert into core.fct_cash_transaction values ('D1', 1, 202302, 1300, 10, false, 0),
    ('D2', 1, 202303, 2300, 45, true, 34.5), ('D3', 1, 202303, 440, 95, true, 6.6);
"""

# V6 (docs/V3_DESIGN.md section 9) columns and rows, added on top of the V5 mini so every V5 number above
# still holds. C2 is effective in Feb 2023 although entered in Jan (the bridge's timing step); C1's cash
# arrives in 202308, so its 202307 claim is a transient claim on the partner settlement; C2's reserve
# sits with the obligor (AH_OB), so its 202308 claim is an obligor claim deducted.
# Rate card: R1 and R2 are the cards C1 and C2 were booked on; C2 is booked 50 above its card (an
# exception); R3 is on the card but unused. Agent A1 sits in district West.
V6_DDL = """
alter table core.dim_contract add column contract_number varchar;
alter table core.dim_contract add column source_system varchar;
alter table core.dim_contract add column tenant_key bigint;
alter table core.dim_contract add column product_key varchar;
alter table core.dim_contract add column agent_key varchar;
alter table core.dim_contract add column entered_date date;
alter table core.dim_contract add column effective_date date;
alter table core.dim_contract add column reportable_state varchar;
update core.dim_contract set contract_number = 'VSC0001', source_system = 'admin_core', tenant_key = 1,
    product_key = 'P1', agent_key = 'A1', entered_date = date '2023-01-10', effective_date = date '2023-01-10',
    reportable_state = 'TX' where contract_key = 'C1';
update core.dim_contract set contract_number = 'VSC0002', source_system = 'admin_core', tenant_key = 1,
    product_key = 'P1', agent_key = 'A1', entered_date = date '2023-01-20', effective_date = date '2023-02-01',
    reportable_state = 'OK' where contract_key = 'C2';
update core.dim_contract set contract_number = 'GAP0003', source_system = 'legacy_admin', tenant_key = 1,
    product_key = 'P2', agent_key = '-1', entered_date = date '2023-02-05', effective_date = date '2023-02-05',
    reportable_state = 'TX' where contract_key = 'C3';
create table core.dim_contract_vehicle (contract_key varchar, vehicle_make varchar, odometer_band varchar);
insert into core.dim_contract_vehicle values ('C1', 'Ford', '0-30K'), ('C2', 'Toyota', '30-60K'),
    ('C3', 'Ford', null);
create table core.dim_contract_terms (contract_key varchar, term_months integer, surcharge_amount double,
    options_amount double);
insert into core.dim_contract_terms values ('C1', 36, 0, 0), ('C2', 60, 0, 0), ('C3', 72, 0, 0);
create table core.dim_contract_finance (contract_key varchar, coverage_cost double);
insert into core.dim_contract_finance values ('C1', 1300), ('C2', 2300), ('C3', 440);
alter table core.dim_product add column product varchar;
alter table core.dim_product add column plan varchar;
alter table core.dim_product add column plan_id bigint;
alter table core.dim_product add column plan_code varchar;
alter table core.dim_product add column plan_year integer;
alter table core.dim_product add column program_type varchar;
update core.dim_product set product = 'Powertrain', plan = 'Powertrain 2023 100 ded', plan_id = 11,
    plan_code = 'VSC-PT-2023-D100', plan_year = 2023, program_type = 'AGENCY' where product_key = 'P1';
update core.dim_product set product = 'GAP', plan = 'GAP 2023 0 ded', plan_id = 12, plan_code = 'GAP-2023-D0',
    plan_year = 2023, program_type = 'FULL_INSURED' where product_key = 'P2';
alter table core.dim_participation_entity add column status varchar;
alter table core.dim_participation_entity add column runoff_date date;
alter table core.dim_participation_entity add column cession_type varchar;
alter table core.dim_participation_entity add column gl_class_code varchar;
update core.dim_participation_entity set status = 'ACTIVE', cession_type = 'C', gl_class_code = 'GL-1'
    where participation_entity_key = 'E1';
update core.dim_participation_entity set status = 'ACTIVE', cession_type = 'S', gl_class_code = 'GL-2'
    where participation_entity_key = 'E2';
alter table core.dim_agent add column district varchar;
alter table core.dim_agent add column region varchar;
update core.dim_agent set territory = 'West', district = 'West', region = 'West';
alter table core.fct_cancellation add column cancel_fee double;
update core.fct_cancellation set cancel_fee = 50;
alter table core.fct_cash_transaction add column cash_txn_key varchar;
alter table core.fct_cash_transaction add column remit_batch_id varchar;
alter table core.fct_cash_transaction add column contract_key varchar;
alter table core.fct_cash_transaction add column contract_entered_date date;
alter table core.fct_cash_transaction add column entered_at timestamp;
alter table core.fct_cash_transaction add column sent_at timestamp;
alter table core.fct_cash_transaction add column received_at timestamp;
alter table core.fct_cash_transaction add column posted_at timestamp;
update core.fct_cash_transaction set cash_txn_key = 'K1', remit_batch_id = 'RB1', contract_key = 'C1',
    posting_period_key = 202308 where dealer_key = 'D1';
update core.fct_cash_transaction set cash_txn_key = 'K2', remit_batch_id = 'RB2', contract_key = 'C2'
    where dealer_key = 'D2';
update core.fct_cash_transaction set cash_txn_key = 'K3', remit_batch_id = 'RB3', contract_key = 'C3'
    where dealer_key = 'D3';
create schema staging;
create table staging.stg_admin_core__rate (rate_id bigint, tenant_id bigint, plan_id bigint, term_months integer,
    term_miles integer, odometer_band varchar, class_code varchar, deductible_amount double,
    coverage_cost double, suggested_retail double, coverage_level varchar, is_deleted boolean);
insert into staging.stg_admin_core__rate values
    (1, 1, 11, 36, 36000, '0-30K', 'A', 100, 1300, 2499, 'POWERTRAIN', false),
    (2, 1, 11, 60, 60000, '30-60K', 'B', 100, 2250, 4299, 'POWERTRAIN', false),
    (3, 1, 11, 24, 24000, '60K+', 'C', 100, 900, 1699, 'POWERTRAIN', false);
create table staging.stg_admin_core__rate_component (rate_id bigint, disbursement_type varchar, amount double,
    is_deleted boolean);
insert into staging.stg_admin_core__rate_component values
    (1, 'ADMIN_FEE', 100, false), (1, 'INSURER_FEE', 30, false), (1, 'PREMIUM_TAX', 20, false),
    (1, 'COMMISSION', 150, false), (1, 'CLAIMS_RESERVE', 950, false), (1, 'PROFIT_SHARE_RESERVE', 50, false),
    (2, 'ADMIN_FEE', 120, false), (2, 'COMMISSION', 80, false), (2, 'CLAIMS_RESERVE', 1950, false),
    (2, 'PROFIT_SHARE_RESERVE', 100, false),
    (3, 'ADMIN_FEE', 100, false), (3, 'COMMISSION', 50, false), (3, 'CLAIMS_RESERVE', 750, false);
create table staging.stg_admin_core__contract (contract_number varchar, rate_id bigint, is_deleted boolean);
insert into staging.stg_admin_core__contract values ('VSC0001', 1, false), ('VSC0002', 2, false);
"""
CORE_DDL += V6_DDL

# V7 (docs/V3_DESIGN.md 10.4) columns, again on top so every V5/V6 number above still holds. Each cost-split
# line carries the group the program setup put it in (edit 21); C2's cancel was never reversed (edit 23);
# commission carries its posting time (edit 28). Portfolio facts (34 + 42) and the dealer-count
# reconciliation (43) read status, product line, tenant and dealer group; C2 cancels 85,000 miles in on a
# 60,000-mile term that started at 30,000, so it has 5,000 miles left (edit 36, band "3000 to 5999").
V7_DDL = """
alter table core.fct_disbursement add column disbursement_group varchar;
update core.fct_disbursement set disbursement_group = case disbursement_type
    when 'ADMIN_FEE' then 'ADMIN' when 'COMMISSION' then 'PRODUCER' when 'CLAIMS_RESERVE' then 'PREMIUM'
    else 'DEALER_PARTICIPATION' end;
alter table core.fct_cancellation add column cancel_key varchar;
alter table core.fct_cancellation add column dealer_key varchar;
alter table core.fct_cancellation add column cancel_date date;
alter table core.fct_cancellation add column reason_group varchar;
alter table core.fct_cancellation add column is_admin_correction boolean;
alter table core.fct_cancellation add column odometer bigint;
alter table core.fct_cancellation add column is_reversed boolean;
alter table core.fct_cancellation add column is_reversal_row boolean;
update core.fct_cancellation set cancel_key = 'X2', dealer_key = 'D2', cancel_date = date '2023-10-15',
    reason_group = 'CUSTOMER', is_admin_correction = false, odometer = 85000, is_reversed = false,
    is_reversal_row = false;
alter table core.fct_commission add column posted_at timestamp;
update core.fct_commission set posted_at = make_timestamp(posting_period_key // 100, posting_period_key % 100,
    10, 9, 0, 0);
alter table core.dim_contract add column status varchar;
alter table core.dim_contract add column category_code varchar;
alter table core.dim_contract add column subcategory_code varchar;
update core.dim_contract set status = 'INFORCE', category_code = 'VSC', subcategory_code = 'VSC-PT'
    where contract_key = 'C1';
update core.dim_contract set status = 'CANCELLED', category_code = 'VSC', subcategory_code = 'VSC-PT'
    where contract_key = 'C2';
update core.dim_contract set status = 'INFORCE', category_code = 'GAP', subcategory_code = 'GAP'
    where contract_key = 'C3';
alter table core.dim_product add column is_bundle boolean;
alter table core.dim_product add column external_product_line varchar;
update core.dim_product set is_bundle = false, external_product_line = category_code;
alter table core.dim_dealer add column dealer_group_key varchar;
alter table core.dim_dealer add column tenant_key bigint;
alter table core.dim_dealer add column city varchar;
alter table core.dim_dealer add column state varchar;
alter table core.dim_dealer add column is_current boolean;
update core.dim_dealer set dealer_group_key = case when dealer_key = 'D3' then 'G2' else 'G1' end,
    tenant_key = 1, city = 'Austin', state = 'TX', is_current = true;
create table core.dim_tenant (tenant_key bigint, reportable_name varchar);
insert into core.dim_tenant values (1, 'Tenant One');
alter table core.dim_contract_terms add column term_miles integer;
update core.dim_contract_terms set term_miles = case contract_key when 'C2' then 60000 when 'C1' then 36000 end;
alter table core.dim_contract_vehicle add column odometer bigint;
update core.dim_contract_vehicle set odometer = 30000 where contract_key = 'C2';
"""
CORE_DDL += V7_DDL


def _dbt_build(db: Path, tmp: Path) -> subprocess.CompletedProcess:
    env = {**os.environ, "WAREHOUSE_PATH": str(db), "DBT_TARGET_PATH": str(tmp / "target"),
           "DBT_LOG_PATH": str(tmp / "logs"), "DO_NOT_TRACK": "1", "DBT_SEND_ANONYMOUS_USAGE_STATS": "false"}
    cmd = [sys.executable, "-m", "dbt.cli.main", "build", "--project-dir", str(DBT_DIR), "--profiles-dir",
           str(DBT_DIR), "--select", "path:seeds/finance", "path:models/marts/finance", "path:tests/finance"]
    # Other areas' singular tests that read a finance mart also need their own marts, which the mini
    # warehouse does not have: leave them to their owners' suites.
    others = sorted(f"path:tests/{q.relative_to(DBT_DIR / 'tests').as_posix()}"
                    for q in (DBT_DIR / "tests").iterdir() if q.name != "finance")
    if others:
        cmd += ["--exclude", *others]
    return subprocess.run(cmd, cwd=ROOT, env=env, capture_output=True, text=True, timeout=600)


def _failed(tmp: Path) -> set[str]:
    results = json.loads((tmp / "target" / "run_results.json").read_text(encoding="utf-8"))["results"]
    return {r["unique_id"].split(".")[-1] for r in results if r["status"] in ("fail", "error")}


@pytest.fixture(scope="module")
def mini(tmp_path_factory) -> Path:
    tmp = tmp_path_factory.mktemp("finance")
    db = tmp / "mini.duckdb"
    with duckdb.connect(str(db)) as con:
        con.execute(CORE_DDL)
    proc = _dbt_build(db, tmp)
    assert proc.returncode == 0, proc.stdout[-4000:]
    return db


def _one(db: Path, sql: str) -> dict:
    with duckdb.connect(str(db), read_only=True) as con:
        df = con.execute(sql).df()
    assert len(df) == 1, df
    return df.iloc[0].to_dict()


def test_cession_experience_matches_hand_calculation(mini):
    r = _one(mini, "select * from finance.mart_cession_experience "
                   "where participation_entity_key = 'E1' and posting_period_key = 202402")
    assert r["itd_written_reserve"] == pytest.approx(1000)
    assert r["itd_earned_reserve"] == pytest.approx(800)
    assert r["itd_paid_claims"] == pytest.approx(300)
    assert r["itd_incurred_claims"] == pytest.approx(400)
    assert r["itd_loss_ratio"] == pytest.approx(0.5)
    assert r["reserve_balance"] == pytest.approx(700)
    assert r["unearned_reserve"] == pytest.approx(200)
    assert r["itd_ceding_fee"] == pytest.approx(16)
    assert r["itd_underwriting_result"] == pytest.approx(384)
    assert r["itd_profit_share_due"] == pytest.approx(384)
    assert r["itd_profit_share_reserve_funded"] == pytest.approx(50)
    # 14 months 202301..202402 with no gaps, though E1 has activity in only 4 of them.
    with duckdb.connect(str(mini), read_only=True) as con:
        n = con.execute("select count(*) from finance.mart_cession_experience "
                        "where participation_entity_key = 'E1'").fetchone()[0]
        dealers = {row[0] for row in con.execute(
            "select distinct dealer_key from finance.mart_cession_experience").fetchall()}
    assert n == 14
    assert dealers == {"D1", "D2"}, "a dealer without a participation entity must not appear"


def test_profit_share_waits_for_seasoning_then_posts_the_movement(mini):
    with duckdb.connect(str(mini), read_only=True) as con:
        df = con.execute("select posting_period_key, statement_status, opening_profit_share_due, "
                         "period_profit_share, closing_profit_share_due from finance.mart_profit_share_statement "
                         "where participation_entity_key = 'E1' and posting_period_key >= 202312 "
                         "order by 1").df()
    assert df["closing_profit_share_due"].tolist() == pytest.approx([0, 384, 384])
    assert df["period_profit_share"].tolist() == pytest.approx([0, 384, 0])
    assert df["statement_status"].tolist() == ["FINAL", "FINAL", "PRELIMINARY"]
    e2 = _one(mini, "select * from finance.mart_profit_share_statement "
                    "where participation_entity_key = 'E2' and posting_period_key = 202402")
    assert e2["itd_underwriting_result"] == pytest.approx(100)  # 1000 earned - 900 incurred, no ceding fee
    assert not e2["is_seasoned"]  # 24-month seasoning not met
    assert e2["closing_profit_share_due"] == pytest.approx(0)
    assert e2["due_category"] == "PAYABLE"
    assert e2["ps_reserve_shortfall"] == pytest.approx(-100)  # nothing owed, 100 of profit-share reserve held
    # Funding gap also counts the earned surplus left in the claims reserve (1000 earned - 900 incurred).
    assert e2["funding_gap"] == pytest.approx(-200)
    e1 = _one(mini, "select * from finance.mart_profit_share_statement "
                    "where participation_entity_key = 'E1' and posting_period_key = 202402")
    assert e1["due_category"] == "DEALER_ENTITY_SURPLUS" and pd.isna(e1["funding_gap"])
    assert e1["distributions_paid"] == 0 and e1["itd_profit_share_earned"] == pytest.approx(384)



def test_pnl_by_product_matches_hand_calculation(mini):
    r = _one(mini, "select * from finance.mart_pnl_by_product where product_code = 'VSC-PT' "
                   "and posting_period_key = 202306")
    assert r["earned_premium"] == pytest.approx(1670)
    assert r["earned_claims_reserve"] == pytest.approx(1600)
    assert r["earned_admin_fee"] == pytest.approx(50)
    assert r["earned_acquisition_cost"] == pytest.approx(20)
    assert r["posted_loss_ratio"] == 0 and r["loss_ratio"] == 0  # no claims posted in 202306
    with duckdb.connect(str(mini), read_only=True) as con:
        tot = con.execute("""select sum(written_coverage_cost), sum(coverage_cost_refunded), sum(earned_premium),
            sum(earned_claims_reserve), sum(incurred_claims), sum(incurred_claims) / sum(earned_claims_reserve)
            from finance.mart_pnl_by_product where product_code = 'VSC-PT'""").fetchone()
    assert tot == pytest.approx((3600, 550, 1870, 1800, 1300, 1300 / 1800))
    # Net written = written - refunded + reinstated; the DENIED reinstatement adds nothing back.
    with duckdb.connect(str(mini), read_only=True) as con:
        nw = con.execute("select sum(net_written_coverage_cost), sum(coverage_cost_reinstated), "
                         "sum(contracts_cancelled) from finance.mart_pnl_by_product").fetchone()
    assert nw == pytest.approx((1300 + 2300 + 440 - 550, 0, 1))


def test_pnl_orphan_claims_are_one_reconciling_row_not_a_product(mini):
    with duckdb.connect(str(mini), read_only=True) as con:
        df = con.execute("select product_code, category, is_unmatched_claims, claim_lines, incurred_claims, "
                         "earned_premium from finance.mart_pnl_by_product where incurred_claims <> 0 "
                         "order by product_code").df()
    assert "UNKNOWN" not in set(df["product_code"])
    um = df[df["is_unmatched_claims"]]
    assert um[["product_code", "category", "claim_lines", "incurred_claims", "earned_premium"]].values.tolist() == [
        ["UNMATCHED", "Unmatched claims", 1, 25, 0]]
    assert df["incurred_claims"].sum() == pytest.approx(400 + 900 + 50 + 25)  # still ties to the claim ledger


def test_pnl_margin_after_dealer_participation(mini):
    with duckdb.connect(str(mini), read_only=True) as con:
        r = con.execute("select sum(dealer_participation), sum(program_margin), "
                        "sum(program_margin_after_participation) from finance.mart_pnl_by_product "
                        "where product_code = 'VSC-PT' and posting_period_key = 202401").fetchone()
    # E1 (CFC on C1, VSC-PT) earns its 384 profit share in 202401: 200 earned reserve - 0 claims - 0 acq.
    assert r == pytest.approx((384, 200, 200 - 384))


def test_pnl_program_margin_is_one_minus_combined_ratio(mini):
    # One treatment of the admin fee: a program cost in both program_margin and combined_ratio.
    r = _one(mini, "select * from finance.mart_pnl_by_product where product_code = 'VSC-PT' "
                   "and posting_period_key = 202306")
    assert r["program_margin"] == pytest.approx(1670 - 0 - 20 - 50)
    assert r["program_margin"] / r["earned_premium"] == pytest.approx(1 - r["combined_ratio"])
    assert r["administrator_result"] == pytest.approx(r["earned_admin_fee"] + r["program_margin_after_participation"])
    with duckdb.connect(str(mini), read_only=True) as con:
        ep, pm, cr = con.execute("select sum(earned_premium), sum(program_margin), sum(incurred_claims + "
                                 "earned_acquisition_cost + earned_admin_fee) / sum(earned_premium) "
                                 "from finance.mart_pnl_by_product").fetchone()
    assert pm / ep == pytest.approx(1 - cr)


def test_commission_contract_counts_are_distinct_over_periods(mini):
    with duckdb.connect(str(mini), read_only=True) as con:
        df = con.execute("select posting_period_key, is_no_agent_ledger, contracts_commissioned, "
                         "contracts_first_commissioned, contracts_first_commissioned_primary "
                         "from finance.mart_commission_statement order by 1, 2").fetchall()
    # C1 and C2 first commissioned in 202301 (agent A1); legacy C3 in 202302; the 202310 chargebacks and
    # 202306 / 202312 curve-only rows add no contracts, so the sum over periods is 3 distinct contracts.
    first = {(k, nl): (a, b, c) for k, nl, a, b, c in df}
    assert first[(202301, False)] == (2, 2, 2) and first[(202302, True)] == (1, 1, 1)
    assert sum(r[4] for r in df) == 3 and sum(r[3] for r in df) == 3


def test_reserve_rollforward_rolls(mini):
    r = _one(mini, "select * from finance.mart_reserve_rollforward where category_code = 'VSC' "
                   "and posting_period_key = 202402")
    assert r["closing_funded_reserve"] == pytest.approx(3000 - 500 - 1200)
    assert r["closing_unearned_reserve"] == pytest.approx(2500 - 1800)
    oct_ = _one(mini, "select * from finance.mart_reserve_rollforward where category_code = 'VSC' "
                      "and posting_period_key = 202310")
    assert oct_["opening_funded_reserve"] == pytest.approx(3000 - 1200)
    assert oct_["reserve_cancelled"] == pytest.approx(-500)


def test_commission_statement_lines(mini):
    with duckdb.connect(str(mini), read_only=True) as con:
        df = con.execute("select posting_period_key, contracts_commissioned, commission_expensed, "
                         "commission_chargebacks, advances_paid, advance_recoveries, advances_net, net_commission, "
                         "advance_balance from finance.mart_commission_statement where not is_no_agent_ledger "
                         "order by 1").df()
        rates = con.execute("select posting_period_key, commission_written, chargeback_rate, "
                            "commission_earned_on_curve, expensed_ahead_of_curve "
                            "from finance.mart_commission_statement where not is_no_agent_ledger order by 1").df()
    # period, contracts, expensed, chargebacks, advances paid, earned down, advances net, net (cash), balance
    assert df.to_numpy().ravel().tolist() == pytest.approx([
        202301, 2, 150, 0, 80, 0, 80, 230, 80,
        202306, 0, 20, 0, 0, 20, -20, 0, 60,
        202310, 0, 60, -50, 0, 60, -60, -50, 0])
    # Expensed + chargebacks + advances net = net cash on every row; chargebacks are never positive.
    assert (df["commission_expensed"] + df["commission_chargebacks"] + df["advances_net"]
            - df["net_commission"]).abs().max() < 0.01
    assert (df["commission_chargebacks"] <= 0).all()
    # Written = paid at sale + advances paid (150 + 80); the chargeback rate divides by it for both plans,
    # so a cancel-only period has no rate. The curve earns 20 in 202306: expensed 150 in 202301 is ahead.
    assert rates["commission_written"].tolist() == pytest.approx([230, 0, 0])
    assert rates["chargeback_rate"].iloc[0] == 0 and rates["chargeback_rate"].iloc[1:].isna().all()
    assert rates["commission_earned_on_curve"].tolist() == pytest.approx([0, 20, 0])
    assert rates["expensed_ahead_of_curve"].tolist() == pytest.approx([150, 0, 60])


def test_legacy_commission_is_a_reconciling_row_and_ties_to_cost_split(mini):
    with duckdb.connect(str(mini), read_only=True) as con:
        nl = con.execute("select posting_period_key, source_system, agent_code, contracts_commissioned, "
                         "commission_written, net_commission, commission_earned_on_curve "
                         "from finance.mart_commission_statement where is_no_agent_ledger order by 1").fetchall()
        split = con.execute("select sum(amount) from core.fct_disbursement "
                            "where disbursement_type = 'COMMISSION'").fetchone()[0]
        stmt = con.execute("select sum(net_commission) from finance.mart_commission_statement").fetchone()[0]
        curve = con.execute("select sum(commission_earned_on_curve) "
                            "from finance.mart_commission_statement").fetchone()[0]
        pnl = con.execute("select sum(earned_commission) from finance.mart_pnl_by_product").fetchone()[0]
    assert nl == [(202302, "legacy_admin", "NO-LEDGER", 1, 10, 10, 0),
                  (202312, "legacy_admin", "NO-LEDGER", 0, 0, 0, 2.5)]
    assert stmt == pytest.approx(split) == pytest.approx(150 + 80 - 50 + 10)
    assert curve == pytest.approx(pnl) == pytest.approx(22.5)


def test_funded_reserve_reconciliation_has_no_plug(mini):
    r = _one(mini, "select * from finance.mart_funded_reserve_reconciliation where posting_period_key = 202402")
    # Rest of book: C3 reserve 400 - its claim 50 - the claim with no contract 25.
    assert r["rest_of_book_reserve_balance"] == pytest.approx(400 - 50 - 25)
    assert r["unmatched_claims_paid"] == pytest.approx(25)
    assert r["participation_reserve_balance"] == pytest.approx(700 + 600)
    assert r["book_closing_funded_reserve"] == pytest.approx(1300 + 325)
    assert abs(r["difference"]) < 0.005 and abs(r["earned_surplus_difference"]) < 0.005


def test_reserve_reconciliation_test_catches_a_cession_balance_error(mini, tmp_path):
    import shutil
    db = tmp_path / "broken_cession.duckdb"
    shutil.copy(mini, db)
    with duckdb.connect(str(db)) as con:
        # A one-dollar error in one entity's reserve balance: the old 'rest of book = book - participation'
        # plug would absorb it; the reconciliation computes each side on its own, so it must show.
        con.execute("update finance.mart_cession_experience set reserve_balance = reserve_balance + 1 "
                    "where participation_entity_key = 'E1' and posting_period_key = 202402")
    env = {**os.environ, "WAREHOUSE_PATH": str(db), "DBT_TARGET_PATH": str(tmp_path / "target"),
           "DBT_LOG_PATH": str(tmp_path / "logs"), "DO_NOT_TRACK": "1", "DBT_SEND_ANONYMOUS_USAGE_STATS": "false"}
    proc = subprocess.run([sys.executable, "-m", "dbt.cli.main", "build", "--project-dir", str(DBT_DIR),
                           "--profiles-dir", str(DBT_DIR), "--select", "mart_funded_reserve_reconciliation",
                           "fin_reserve_rollforward_ties_to_cession"],
                          cwd=ROOT, env=env, capture_output=True, text=True, timeout=600)
    assert proc.returncode != 0
    assert _failed(tmp_path) == {"fin_reserve_rollforward_ties_to_cession"}


def test_remittance_buckets_and_late_fees(mini):
    with duckdb.connect(str(mini), read_only=True) as con:
        rows = con.execute("select dealer_number, aging_bucket, receipts, late_fee_billed "
                           "from finance.mart_remittance_late_fees order by 1").fetchall()
    assert rows == [("DLR001", "0-15", 1, 0), ("DLR002", "31-60", 1, 34.5), ("DLR003", "91+", 1, 6.6)]


def test_controls_balance_and_cash_is_excluded(mini):
    with duckdb.connect(str(mini), read_only=True) as con:
        df = con.execute("select * from finance.mart_disbursement_control_totals").df()
    assert set(df["txn_type"]) == {"SALE", "CANCEL"}
    assert df["is_balanced"].all()
    assert df["txn_amount"].sum() == pytest.approx(1300 + 2300 - 550 + 440)


# ---- V6 marts on the mini (as of 202401, the latest closed period)
def _rows(db: Path, sql: str) -> list[tuple]:
    with duckdb.connect(str(db), read_only=True) as con:
        return con.execute(sql).fetchall()


def test_participation_map_cedes_every_contract_while_active(mini):
    assert _rows(mini, "select contract_key, participation_entity_key, entity_status, is_ceded "
                       "from finance.mart_participation_contract order by 1") == [
        ("C1", "E1", "ACTIVE", True), ("C2", "E2", "ACTIVE", True)]


def test_experience_matches_hand_calculation_and_every_view_adds_up(mini):
    # C1: net 1000, earned 800, paid 300; C2: 2000 - 500 cancel, earned 1000, paid 900, cancelled;
    # C3 (legacy GAP, no entity): 400, earned 100, paid 50. CL4 has no contract and is not experience.
    tot = _rows(mini, "select sum(net_written), sum(earned), sum(unearned), sum(claims_paid), "
                      "sum(contracts_written), sum(contracts_cancelled), sum(contracts_net), sum(claims_count), "
                      "min(as_of_period_key) from finance.mart_experience")[0]
    assert tot == pytest.approx((2900, 1900, 1000, 1250, 3, 1, 2, 3, 202401))
    for col in ("entered_year", "product_code", "dealer_key", "reportable_state", "plan", "vehicle_make",
                "effective_year", "term_months"):
        assert _rows(mini, f"select sum(net_written) from (select {col}, sum(net_written) as net_written "
                           f"from finance.mart_experience group by 1)")[0][0] == pytest.approx(2900), col
    ford = _rows(mini, "select sum(net_written), sum(earned), sum(claims_paid) from finance.mart_experience "
                       "where vehicle_make = 'Ford'")[0]
    assert ford == pytest.approx((1400, 900, 350))
    c2 = _one(mini, "select * from finance.mart_experience where dealer_key = 'D2'")
    assert c2["effective_year"] == 2023 and c2["participation_entity_key"] == "E2"
    assert c2["cancel_pct"] == pytest.approx(1.0) and c2["earned_pct"] == pytest.approx(1000 / 1500)
    assert c2["loss_ratio"] == pytest.approx(0.9)
    assert _one(mini, "select * from finance.mart_experience where dealer_key = 'D3'")[
        "participation_entity_key"] == "-1"


def test_new_reserves_by_term_bucket(mini):
    rows = _rows(mini, "select entered_month, category_code, bucket_label, contracts, new_reserve, avg_reserve, "
                       "avg_term_months, row_pct from finance.mart_new_reserves_by_term order by 1, 3")
    jan, feb = pd.Timestamp("2023-01-01").date(), pd.Timestamp("2023-02-01").date()
    assert [r[:5] for r in rows] == [(jan, "VSC", "2-3", 1, 1000), (jan, "VSC", "4-5", 1, 2000),
                                     (feb, "GAP", "5-6", 1, 400)]
    assert [r[7] for r in rows] == pytest.approx([1 / 3, 2 / 3, 1.0])
    assert rows[1][5] == pytest.approx(2000) and rows[1][6] == pytest.approx(60)


def test_reserve_bridge_walks_new_reserves_to_net_written(mini):
    # Jan: gross 3000 (C1 + C2) - C2's 2000 sale reserve + its 1500 remainder = 2500 entered; C2 is
    # effective in Feb, so -1500 timing -> 1000 by effective month. Feb: C3 400 + C2's 1500 timing in.
    rows = _rows(mini, "select month_start, gross_new_reserve, cancelled_contracts_reserve, partial_refund_remainder, "
                       "other_adjustments, net_written_entered, timing_in, timing_out, net_written_effective, "
                       "residual, contracts_gross, contracts_cancelled, contracts_net "
                       "from finance.mart_reserve_bridge order by 1")
    assert [r[1:] for r in rows] == pytest.approx([(3000, -2000, 1500, 0, 2500, 0, -1500, 1000, 0, 2, 1, 1),
                                                   (400, 0, 0, 0, 400, 1500, 0, 1900, 0, 1, 0, 1)])


def test_reserve_bridge_test_catches_a_break(mini, tmp_path):
    import shutil
    db = tmp_path / "broken_bridge.duckdb"
    shutil.copy(mini, db)
    with duckdb.connect(str(db)) as con:  # the Experience page drifts one dollar from the bridge
        con.execute("update finance.mart_experience set net_written = net_written + 1 where dealer_key = 'D1'")
    env = {**os.environ, "WAREHOUSE_PATH": str(db), "DBT_TARGET_PATH": str(tmp_path / "target"),
           "DBT_LOG_PATH": str(tmp_path / "logs"), "DO_NOT_TRACK": "1", "DBT_SEND_ANONYMOUS_USAGE_STATS": "false"}
    proc = subprocess.run([sys.executable, "-m", "dbt.cli.main", "test", "--project-dir", str(DBT_DIR),
                           "--profiles-dir", str(DBT_DIR), "--select", "assert_reserve_bridge_closes"],
                          cwd=ROOT, env=env, capture_output=True, text=True, timeout=600)
    assert proc.returncode != 0 and _failed(tmp_path) == {"assert_reserve_bridge_closes"}


def test_commission_hierarchy_rolls_up_against_targets(mini):
    # A1 (district West, region West) sold C1 and C2 (VSC) in 202301 and was paid 150 + 80 then.
    d = _one(mini, "select * from finance.mart_commission_hierarchy where level = 'DISTRICT' and node = 'West' "
                   "and product_type = 'VSC' and posting_period_key = 202301")
    target = {(r["district"], r["product_type"]): int(r["monthly_target_contracts"])
              for r in _seed("commission_volume_targets")}
    assert (d["contracts"], d["commission"], d["parent_node"]) == (2, pytest.approx(230), "West")
    assert d["target_contracts"] == target[("West", "VSC")]
    assert d["over_under_contracts"] == 2 - target[("West", "VSC")]
    n = _one(mini, "select * from finance.mart_commission_hierarchy where level = 'NATIONAL' "
                   "and product_type = 'VSC' and posting_period_key = 202301")
    assert n["target_contracts"] == sum(v for (_, pt), v in target.items() if pt == "VSC")
    # Each level's commission sums to the ledger (150 + 80 - 30 - 20); C3 has no agent: UNASSIGNED.
    for level, total in _rows(mini, "select level, sum(commission) from finance.mart_commission_hierarchy "
                                    "group by 1"):
        assert total == pytest.approx(180), level
    un = _one(mini, "select * from finance.mart_commission_hierarchy where level = 'DISTRICT' "
                    "and node = 'UNASSIGNED' and posting_period_key = 202302")
    assert un["contracts"] == 1 and pd.isna(un["target_contracts"]) and pd.isna(un["over_under_contracts"])


# ---- V7 marts on the mini (docs/V3_DESIGN.md 10.4)
def test_v7_control_totals_split_by_disbursement_group(mini):
    df = _rows(mini, "select txn_type, split_amount, group_admin + group_premium + group_dealer_participation "
                     "+ group_producer + group_tax_and_fees, group_unassigned, group_producer "
                     "from finance.mart_disbursement_control_totals order by 1")
    for txn_type, split, groups, unassigned, _ in df:
        assert groups == pytest.approx(split) and unassigned == 0, txn_type
    producer = {}  # one row per source system; the producer group is the commission bucket (150 + 80 + 10)
    for r in df:
        producer[r[0]] = producer.get(r[0], 0) + r[4]
    assert producer == {"CANCEL": pytest.approx(-50), "SALE": pytest.approx(240)}


def test_v7_bridge_shows_reversed_cancels_and_commission_is_stamped(mini):
    assert _rows(mini, "select sum(contracts_cancel_reversed), sum(reversed_cancel_reserve) "
                       "from finance.mart_reserve_bridge") == [(0, 0.0)]
    r = _one(mini, "select * from finance.mart_commission_hierarchy where level = 'DISTRICT' and node = 'West' "
                   "and product_type = 'VSC' and posting_period_key = 202301")
    assert r["commission_period_label"] == "Jan 2023" and r["as_of_period_key"] == 202402
    assert pd.Timestamp(r["data_as_of_ts"]) == pd.Timestamp("2023-10-10 09:00") and not r["is_unassigned"]
    assert _rows(mini, "select distinct is_unassigned from finance.mart_commission_hierarchy "
                       "where node = 'UNASSIGNED'") == [(True,)]


def test_v7_portfolio_facts_windows_and_cancel_methods(mini):
    facts = {(w, m): (n, d, v) for w, m, n, d, v in _rows(
        mini, "select window_id, metric_id, numerator, denominator, value from finance.mart_portfolio_facts "
              "where tenant_scope = 'COMBINED' and product_line = 'ALL'")}
    assert facts[("ITD", "CONTRACTS")][0] == 3 and facts[("ITD", "ACTIVE_DEALERS")][0] == 3
    assert facts[("ITD", "DEALER_GROUPS")][0] == 2
    assert facts[("ITD", "CANCEL_RATE_COHORT")] == (1, 3, pytest.approx(1 / 3, abs=1e-6))
    # TTM ends Jan 2024 (the latest closed period): only C3 was sold in it, and C2's cancel is dated in it
    assert facts[("TTM", "CONTRACTS")][0] == 1 and facts[("TTM", "CANCEL_RATE_FLOW")][2] == pytest.approx(1.0)
    assert facts[("CY2023", "CONTRACTS")][0] == 3
    by_line = dict(_rows(mini, "select product_line, numerator from finance.mart_portfolio_facts "
                               "where window_id = 'ITD' and metric_id = 'CONTRACTS' and tenant_scope = '1'"))
    assert by_line == {"ALL": 3, "VSC": 2, "GAP": 1}


def test_v7_dealer_counts_side_by_side_and_mileage_bands(mini):
    counts = dict(_rows(mini, "select basis, dealer_count from finance.mart_dealer_count_reconciliation "
                              "where tenant_scope = 'COMBINED'"))
    assert counts == {"TTM_NONVOIDED": 1, "TTM_INCL_VOIDED": 1, "ITD_NONVOIDED": 3, "ITD_ALL": 3,
                      "DEALER_DIM_TOTAL": 3, "CROSS_TENANT_DEDUP": 3}  # three different names: nothing merges
    m = _one(mini, "select * from finance.mart_cancel_mileage")
    assert (m["miles_used"], m["miles_remaining"], m["mileage_band_id"]) == (55000, 5000, "LT6K")
    assert not m["is_near_limit"]


def test_rate_card_fee_stack_and_reviewer_sort(mini):
    import reports.finance as fin
    r1 = _one(mini, "select * from finance.mart_rate_card where rate_id = 1")
    assert (r1["rate"], r1["admin_fee"], r1["program_fee"], r1["premium_tax"], r1["dealer_pack"],
            r1["fees_total"], r1["reserve"]) == pytest.approx((1300, 100, 30, 20, 150, 300, 1000))
    assert r1["program_type"] == "AGENCY" and r1["product_code"] == "VSC-PT"
    with duckdb.connect(str(mini), read_only=True) as con:
        d = fin.load(con)
    assert d["rate_card"]["term_months"].tolist() == [24, 36, 60]  # term first, then odometer band
    ex = _rows(mini, "select contract_number, exception_type, difference from finance.mart_rate_card_exceptions")
    assert ex == [("VSC0002", "COST_DIFFERS", 50.0)]  # booked 2300 on a 2250 card, no surcharges


def test_partner_settlement_lines(mini):
    rows = _rows(mini, "select dealer_number, posting_period_key, admin_premium_due, program_fee_adjustment, "
                       "cancel_fees, transient_claims, obligor_claims_deducted, net_due, statement_status "
                       "from finance.mart_partner_settlement order by 1, 2")
    assert rows == [("DLR001", 202301, 1300, 0, 0, 0, 0, 1300, "FINAL"),
                    ("DLR001", 202307, 0, 0, 0, 300, 0, 300, "FINAL"),       # paid before C1 was remitted
                    ("DLR002", 202301, 2300, 0, 0, 0, 0, 2300, "FINAL"),
                    ("DLR002", 202308, 0, 0, 0, 0, 900, -900, "FINAL"),      # obligor-held reserve
                    ("DLR002", 202310, -550, 0, 50, 0, 0, -500, "FINAL")]
    # Legacy C3 is not remitted through the partner.
    assert not _rows(mini, "select * from finance.mart_partner_settlement where dealer_number = 'DLR003'")


def test_remit_register_ties_to_cash(mini):
    rows = _rows(mini, "select remit_batch_id, dealer_number, contract_number, posting_period_key, amount, "
                       "late_fee from finance.mart_remit_register order by 1")
    assert rows == [("RB1", "DLR001", "VSC0001", 202308, 1300, 0), ("RB2", "DLR002", "VSC0002", 202303, 2300, 34.5),
                    ("RB3", "DLR003", "GAP0003", 202303, 440, 6.6)]


# A run-off copy of the mini: E1 goes to run-off on 2023-06-30 and its dealer D1 sells C4 in Aug 2023.
# C4's reserve (500) goes to the obligor, so it must stay out of E1's experience and show as a transfer.
RUNOFF_DDL = """
update core.dim_participation_entity set status = 'RUNOFF', runoff_date = date '2023-06-30'
    where participation_entity_key = 'E1';
insert into core.dim_contract (contract_key, dealer_key, contract_number, source_system, tenant_key, product_key,
    agent_key, entered_date, effective_date, reportable_state)
    values ('C4', 'D1', 'VSC0004', 'admin_core', 1, 'P1', 'A1', date '2023-08-15', date '2023-08-15', 'TX');
insert into core.fct_transaction values ('T6', 'admin_core', 'C4', 'SALE', 202308, 500);
insert into core.fct_disbursement values
    ('T6', 'C4', 'D1', 'P1', 1, 'AH_OB', 'admin_core', 'SALE', 'CLAIMS_RESERVE', 202308, 500, 'PREMIUM');
insert into core.fct_gross_premium values ('C4', 'P1', 1, 202308, 1, 500, 800);
"""


def test_runoff_entity_keeps_its_book_and_transfers_new_business(tmp_path):
    db = tmp_path / "runoff.duckdb"
    with duckdb.connect(str(db)) as con:
        con.execute(CORE_DDL + RUNOFF_DDL)
    proc = _dbt_build(db, tmp_path)  # every finance mart and test, incl. the ledger attribution test
    assert proc.returncode == 0, proc.stdout[-4000:]
    assert _rows(db, "select contract_key, is_ceded, is_transferred from finance.mart_participation_contract "
                     "where participation_entity_key = 'E1' order by 1") == [("C1", True, False), ("C4", False, True)]
    e1 = _one(db, "select * from finance.mart_cession_experience "
                  "where participation_entity_key = 'E1' and posting_period_key = 202402")
    assert e1["entity_status"] == "RUNOFF" and e1["cession_type"] == "C"
    assert e1["itd_written_reserve"] == pytest.approx(1000) and e1["reserve_balance"] == pytest.approx(700)
    assert (e1["itd_contracts_transferred"], e1["itd_reserve_transferred"]) == (1, pytest.approx(500))
    # The obligor's 500 is rest of book, and the reconciliation still has no difference.
    r = _one(db, "select * from finance.mart_funded_reserve_reconciliation where posting_period_key = 202402")
    assert r["rest_of_book_reserve_balance"] == pytest.approx(400 - 50 - 25 + 500)
    assert abs(r["difference"]) < 0.005


def test_injected_cost_split_break_fails_the_close_control(tmp_path):
    db = tmp_path / "broken.duckdb"
    with duckdb.connect(str(db)) as con:
        con.execute(CORE_DDL)
        con.execute("update core.fct_transaction set amount = 1300.50 where txn_key = 'T1'")
    proc = _dbt_build(db, tmp_path)
    assert proc.returncode != 0
    # The ledger no longer equals its split, and the P&L's net written no longer ties to the ledger.
    assert _failed(tmp_path) == {"fin_disbursement_controls_balanced", "fin_pnl_ties_to_core"}


def test_report_pages_and_workbook(mini, tmp_path):
    from jinja2 import Environment, FileSystemLoader, select_autoescape
    from markupsafe import Markup
    from openpyxl import load_workbook

    import reports.finance as fin
    from reports.build import HERE, Page

    with duckdb.connect(str(mini), read_only=True) as con:
        pages = fin.build(con, tmp_path)
    assert pages and all(isinstance(p, Page) for p in pages)
    assert pages[0].path == "finance/index.html" and fin.NAV_LABEL == "Finance"
    paths = {p.path for p in pages}
    assert {"finance/cession.html", "finance/profit-share.html", "finance/commissions.html", "finance/pnl.html",
            "finance/remittance.html", "finance/controls.html", "finance/experience.html",
            "finance/new-reserves.html", "finance/commission-hierarchy.html"} <= paths

    wb = load_workbook(tmp_path / "downloads" / "finance" / fin.WORKBOOK, read_only=True)
    names = wb.sheetnames
    assert names[0] == "README"
    first_output = names.index("Cession_ITD")
    assert all(n.startswith("Input_") for n in names[1:first_output]), "inputs come before outputs"

    # Render with the shell's template and check that every local link resolves.
    env = Environment(loader=FileSystemLoader(HERE / "templates"), autoescape=select_autoescape(["html", "j2"]),
                      trim_blocks=True, lstrip_blocks=True)
    tpl = env.get_template("page.html.j2")
    (tmp_path / "assets").mkdir()
    for name in ("site.css", "site.js", "plotly.min.js"):
        (tmp_path / "assets" / name).write_text("", encoding="utf-8")
    from reports import infographics

    infographics.copy_to_site(tmp_path / "assets")  # render_site copies the textbook figures the same way
    (tmp_path / "index.html").write_text("", encoding="utf-8")
    for p in pages:
        out = tmp_path / p.path
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(tpl.render(page=p, root="../" * p.path.count("/"), nav=[("index.html", "Overview")],
                                  active=p.path, as_of=pd.Timestamp("2024-02-29"), last_month="Jan 2024",
                                  chart=lambda s: Markup(s)), encoding="utf-8")

    class Links(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self.refs: list[str] = []

        def handle_starttag(self, tag, attrs):
            self.refs += [v for k, v in attrs if k in ("href", "src") and v and not v.startswith(("#", "http"))]

    broken = []
    for p in pages:
        parser = Links()
        page_file = tmp_path / p.path
        parser.feed(page_file.read_text(encoding="utf-8"))
        broken += [f"{p.path} -> {r}" for r in parser.refs
                   if not (page_file.parent / r.split("#")[0]).resolve().is_file()]
    assert not broken
    cession = (tmp_path / "finance" / "cession.html").read_text(encoding="utf-8")
    assert "Summit Reinsurance Co" in cession and "Harbor Cars" not in cession
    for p in pages:
        assert "$-" not in (tmp_path / p.path).read_text(encoding="utf-8"), f"{p.path}: sign after the $"
    pnl = (tmp_path / "finance" / "pnl.html").read_text(encoding="utf-8")
    assert "Claims not matched to a contract" in pnl and "UNMATCHED" not in pnl and "Unknown" not in pnl
    # V6: View by sections, the bridge and rate-card check on controls, run-off on cession, settlement.
    exp = (tmp_path / "finance" / "experience.html").read_text(encoding="utf-8")
    assert all(f"By {v}" in exp for v in ("Date", "Product", "Dealer", "State", "Plan", "Make", "Term"))
    controls = (tmp_path / "finance" / "controls.html").read_text(encoding="utf-8")
    assert "Reserve bridge" in controls and "Residual (must be $0.00)" in controls and "VSC0002" in controls
    assert "Run-off details" in cession and "Transfer summary" in cession
    assert "Partner settlement statement" in (tmp_path / "finance" / "remittance.html").read_text(encoding="utf-8")
    index = (tmp_path / "finance" / "index.html").read_text(encoding="utf-8")
    assert all(f'{h}"' in index for h in ("finance/experience.html", "finance/new-reserves.html",
                                          "finance/commission-hierarchy.html"))
    assert {"Experience", "New_Reserves", "Reserve_Bridge", "Commission_Hierarchy", "Rate_Card",
            "Rate_Card_Exceptions", "Partner_Settlement", "Remit_Register"} <= set(names)

    ws = wb["Cession_ITD"]
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert not any(h.endswith("_key") and h not in ("posting_period_key", "tenant_key") for h in header)


# --------------------------------------------------------------------------- 3. real warehouse
def _finance_built() -> bool:
    if not WAREHOUSE.exists():
        return False
    try:
        with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
            return con.execute("select count(*) from information_schema.tables where table_schema = 'finance' "
                               "and table_name = 'mart_cession_experience'").fetchone()[0] == 1
    except duckdb.Error:
        return False


@pytest.mark.skipif(not _finance_built(), reason="finance marts not built (run python -m pipeline.run)")
def test_real_warehouse_finance_ties_to_core():
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        q = lambda s: con.execute(s).fetchone()[0]  # noqa: E731
        assert q("select sum(earned_premium) from finance.mart_pnl_by_product") == pytest.approx(
            q("select sum(earned_amount) from core.fct_earned_disbursement"), abs=0.01)
        assert q("select sum(net_commission) from finance.mart_commission_statement "
                 "where not is_no_agent_ledger") == pytest.approx(q("select sum(amount) from core.fct_commission"),
                                                                  abs=0.01)
        # Cost-split commission = agent ledger + legacy unassigned (no commission dollar without a payee).
        assert q("select sum(net_commission) from finance.mart_commission_statement") == pytest.approx(
            q("select sum(amount) from core.fct_disbursement where disbursement_type = 'COMMISSION'"), abs=1.0)
        assert q("select sum(commission_earned_on_curve) from finance.mart_commission_statement") == pytest.approx(
            q("select sum(earned_commission) from finance.mart_pnl_by_product"), abs=0.05)
        assert q("select max(abs(difference)) + max(abs(earned_surplus_difference)) "
                 "from finance.mart_funded_reserve_reconciliation") < 0.05
        assert q("select sum(net_written_coverage_cost) from finance.mart_pnl_by_product") == pytest.approx(
            q("select sum(amount) from core.fct_transaction where txn_type in "
              "('SALE', 'AMEND', 'CANCEL', 'REINSTATE')"), abs=0.01)
        assert q("select count(*) from finance.mart_pnl_by_product where product_code = 'UNKNOWN'") == 0
        # Advance-plan agents count their contracts and earn their advances down.
        assert q("select count(*) from (select agent_key from finance.mart_commission_statement group by 1 "
                 "having sum(advances_paid) > 0 and (sum(contracts_commissioned) = 0 "
                 "or sum(advance_recoveries) = 0))") == 0
        # Profit share is paid out once a year (seed payout month) and the balance is earned - paid.
        assert q("select sum(distributions_paid) from finance.mart_profit_share_statement") > 0
        assert q("select count(*) from finance.mart_profit_share_statement as s "
                 "join seeds.profit_share_payout_schedule as p on p.structure_type = s.entity_type "
                 "where s.distributions_paid > 0.005 and s.posting_period_key % 100 <> p.payout_month") == 0
        assert q("select max(abs(itd_profit_share_earned - itd_distributions_paid - closing_profit_share_due)) "
                 "from finance.mart_profit_share_statement") < 0.005
        assert q("select sum(late_fee_billed) from finance.mart_remittance_late_fees") == pytest.approx(
            q("select sum(late_fee) from core.fct_cash_transaction"), abs=0.01)
        assert q("select count(*) from finance.mart_disbursement_control_totals where not is_balanced") == 0
        assert q("select count(distinct participation_entity_key) from finance.mart_cession_experience") == q(
            "select count(*) from core.dim_participation_entity")
        # Distributions leave the reserve: the claims-reserve part of each payout (beyond the profit-share
        # reserve) is deducted from the cession reserve balance and from the book-wide roll-forward.
        assert q("select sum(claims_reserve_released) from finance.mart_cession_experience") > 0
        assert q("select max(abs(reserve_balance - (itd_written_reserve - itd_paid_claims "
                 "- itd_claims_reserve_released))) from finance.mart_cession_experience") < 0.005
        assert q("select sum(surplus_released) from finance.mart_reserve_rollforward") == pytest.approx(
            q("select sum(claims_reserve_released) from finance.mart_cession_experience"), abs=0.05)
        # Cession closing due = statement closing due, entity by entity.
        assert q("select max(abs(c.closing_profit_share_due - s.closing_profit_share_due)) "
                 "from finance.mart_cession_experience as c join finance.mart_profit_share_statement as s "
                 "using (participation_entity_key, posting_period_key)") < 0.005


# --------------------------------------------------------------------------- 4. page logic on small frames
def _stmt_row(name: str, cat: str, due: float, **kw) -> dict:
    base = dict(entity_name=name, entity_type="PROFIT_SHARE" if cat == "PAYABLE" else "CFC",
                statement_status="FINAL", due_category=cat, itd_earned_reserve=1000.0, itd_incurred_claims=500.0,
                itd_ceding_fee=0.0, itd_underwriting_result=500.0, dealer_share_pct=0.5, is_seasoned=True,
                months_seasoned=30, opening_profit_share_due=due, period_profit_share=0.0, distributions_paid=0.0,
                closing_profit_share_due=due, funding_gap=None if cat != "PAYABLE" else -10.0,
                ps_reserve_shortfall=None if cat != "PAYABLE" else -5.0)
    base.update(kw)
    return base


def test_profit_share_payable_is_not_netted_with_overpayments():
    import reports.finance as fin
    st = pd.DataFrame([_stmt_row("A", "PAYABLE", 100.0), _stmt_row("B", "PAYABLE", -30.0),
                       _stmt_row("C", "DEALER_ENTITY_SURPLUS", 50.0), _stmt_row("D", "DEALER_ENTITY_SURPLUS", -5.0)])
    ps = pd.DataFrame({"posting_period_key": [202609], "period_profit_share": [0.0], "distributions_paid": [0.0],
                       "closing_payable": [100.0], "closing_recoverable": [-30.0], "closing_surplus": [50.0],
                       "closing_over_released": [-5.0]})
    page = fin.page_profit_share({"statement": st, "ps_monthly": ps, "t12_start": 202510, "closed": 202609})
    # The monthly table's latest row shows the same figures as the KPI cards, never netted.
    assert page.sections[0].rows[0][3:] == [fin.money(100.0), fin.money(30.0), fin.money(50.0), fin.money(5.0)]
    netted = ps.assign(closing_payable=70.0, closing_surplus=45.0)
    with pytest.raises(ValueError, match="closing_payable"):
        fin.page_profit_share({"statement": st, "ps_monthly": netted, "t12_start": 202510, "closed": 202609})
    k = {x.label: x for x in page.kpis}
    assert k["Payable to dealers"].value == fin.money(100.0)
    assert k["Overpaid, recoverable"].value == fin.money(30.0) and k["Overpaid, recoverable"].status == "warning"
    assert k["Underwriting surplus in dealer-owned entities"].value == fin.money(50.0)
    flags = [r[13] for r in page.sections[1].rows]
    assert flags == ["", "Overpaid, recoverable", "", "Over-released surplus"]


def test_controls_last_12_excludes_the_open_period_and_reconciles_reserve():
    import reports.finance as fin
    keys = [202509, 202510, 202609, 202610]
    ctl = pd.DataFrame({"posting_period_key": keys, "source_system": "admin_core", "txn_type": "SALE",
                        "txn_count": 1, "txn_amount": 1.0, "split_amount": 1.0, "difference": 0.0,
                        "rounding_txns": 0, "rounding_residual": 0.0, "is_balanced": True})
    rf = pd.DataFrame({"posting_period_key": [202609], "opening": [90.0], "written": [20.0], "paid": [5.0],
                       "released": [5.0], "closing": [100.0], "earned": [3.0], "unearned": [60.0]})
    ce = pd.DataFrame({"reserve_balance": [30.0, 10.0]})
    rr = pd.DataFrame({"participation_reserve_balance": [40.0], "rest_of_book_reserve_balance": [59.0],
                       "unmatched_claims_paid": [2.0], "book_closing_funded_reserve": [100.0],
                       "book_closing_unearned_reserve": [60.0], "earned_surplus": [40.0], "difference": [1.0],
                       "earned_surplus_difference": [0.0]})
    page = fin.page_controls({"controls": ctl, "rollforward": rf, "cession": ce, "reserve_recon": rr,
                              "t12_start": 202510, "closed": 202609})
    periods = [r[0] for r in page.sections[0].rows]
    assert periods == ["Oct 2025", "Sep 2026"] and "Oct 2026" not in periods
    assert "Oct 2025 to Sep 2026" in page.sections[0].title
    assert page.kpis[0].value == "3" and "Sep 2026" in page.kpis[0].sub
    recon = {r[0]: r[1] for r in page.sections[2].rows}
    assert recon["Participation reserve balance (2 entities, Cession page)"] == fin.money_cents(40.0)
    # The rest of book is its own figure, not book - participation, so a $1 break shows as a difference.
    assert recon["Contracts with no participation entity: reserve written less claims paid"] == fin.money_cents(59.0)
    assert recon["Difference"] == {"text": fin.money_cents(1.0), "status": "critical"}
    assert recon["Difference: closing funded - unearned - earned surplus"]["status"] == "good"


@pytest.mark.skipif(not _finance_built(), reason="finance marts not built (run python -m pipeline.run)")
def test_real_warehouse_profit_share_monthly_row_equals_statement_kpis():
    import reports.finance as fin
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        try:
            d = fin.load(con)
        except duckdb.Error as exc:  # marts built before this change: rebuild the finance slice
            pytest.skip(f"finance marts out of date: {exc}")
    fin.check_ps_monthly_ties(d["ps_monthly"], d["statement"], d["closed"])  # raises on any netting
