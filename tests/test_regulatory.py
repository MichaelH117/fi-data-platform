"""Regulatory area: counting-rule seeds, the filing / survey workbooks and HTML, and the plug-in.

Unit tests run on the seed CSVs and a tiny hand-built warehouse in tmp_path. The last tests
run against WAREHOUSE_PATH when its regulatory marts exist and skip otherwise.
"""

from __future__ import annotations

import csv
import datetime as dt
import os
from pathlib import Path

import duckdb
import openpyxl
import pandas as pd
import pytest

from reports.__main__ import broken_links
from reports.build import Page, plugin_pages
from reports.regulatory import NAV_LABEL, build, data, excel
from reports.regulatory import __main__ as reg_cli

ROOT = Path(__file__).resolve().parents[1]
SEEDS = ROOT / "dbt" / "seeds" / "regulatory"
WAREHOUSE = Path(os.environ.get("WAREHOUSE_PATH", ROOT / "data" / "warehouse" / "warehouse.duckdb"))
STATES_AND_DC = {
    "AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "DC", "FL", "GA", "HI", "ID", "IL", "IN", "IA", "KS", "KY",
    "LA", "ME", "MD", "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH", "NJ", "NM", "NY", "NC", "ND", "OH",
    "OK", "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT", "VT", "VA", "WA", "WV", "WI", "WY",
}


def rows(name: str) -> list[dict[str, str]]:
    with open(SEEDS / f"{name}.csv", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# --------------------------------------------------------------------------- seeds: state method matrix
def test_state_matrix_covers_all_50_states_and_dc_once():
    states = [r["state"] for r in rows("state_reserve_method")]
    assert len(states) == 51
    assert set(states) == STATES_AND_DC


def test_state_matrix_requirement_rows_are_complete_and_others_blank():
    for r in rows("state_reserve_method"):
        assert 1 <= int(r["period_start_month"]) <= 12, r["state"]
        assert r["notes"].strip(), r["state"]
        if r["has_requirement"] == "true":
            assert r["method_code"] in {"NET_OF_CLAIMS", "GROSS_CONSIDERATION"}, r["state"]
            assert 0 < float(r["reserve_rate"]) <= 1, r["state"]
            assert 0 <= float(r["security_deposit_rate"]) <= 1, r["state"]
            assert r["default_basis"] in {"COVERAGE_COST", "RETAIL"}, r["state"]
        else:
            assert r["has_requirement"] == "false"
            assert (r["method_code"], r["reserve_rate"], r["default_basis"]) == ("NONE", "", ""), r["state"]


def test_state_matrix_has_skipped_states_both_methods_and_a_survey():
    m = rows("state_reserve_method")
    assert any(r["has_requirement"] == "false" for r in m)
    assert {r["method_code"] for r in m} == {"NET_OF_CLAIMS", "GROSS_CONSIDERATION", "NONE"}
    assert {r["default_basis"] for r in m if r["default_basis"]} == {"COVERAGE_COST", "RETAIL"}
    assert len({r["period_start_month"] for r in m if r["has_requirement"] == "true"}) > 1  # Sep-Aug and Jan-Dec
    assert any(r["market_survey"] == "true" for r in m)


# --------------------------------------------------------------------------- seeds: status sets and definitions
def status_sets() -> dict[str, set[str]]:
    out: dict[str, set[str]] = {}
    for r in rows("report_status_sets"):
        out.setdefault(r["status_set_id"], set()).add(r["status"])
    return out


def test_status_sets_encode_the_counting_rules():
    s = status_sets()
    holds = {"SUSPENDED", "REINSTATEMENT_PENDING", "REINSTATEMENT_SUSPENDED"}  # V6: written and still in force
    assert s["WRITTEN"] == {"PENDING", "INFORCE", "CANCELLED", "EXPIRED"} | holds
    assert holds <= s["IN_FORCE"]
    assert s["CLAIM_PAID"] == {"CLOSED_FULLY_PAID", "CLOSED_PARTIALLY_PAID"}
    assert s["CLAIM_DENIED"] == {"DENIED"}
    for name in ("WRITTEN", "IN_FORCE"):  # voided and denied contracts never count
        assert not s[name] & {"VOIDED", "DENIED", "REINSTATEMENT_VOIDED", "REINSTATEMENT_DENIED"}
    assert "PENDING" not in s["IN_FORCE"]
    assert s["CLAIM_PAID"] <= s["CLAIM_ANY"] and s["CLAIM_DENIED"] <= s["CLAIM_ANY"]
    assert s["WRITTEN"] | s["IN_FORCE"] < s["CONTRACT_ANY"]  # all-contracts dollars tie to the claim count


def test_report_definitions_reference_known_sets_scopes_and_bases():
    sets = status_sets()
    defs = rows("report_definitions")
    keys = [(r["report_id"], r["line_item"], r["scope"], r["basis"]) for r in defs]
    assert len(keys) == len(set(keys))
    for r in defs:
        assert r["scope"] in {"ITD", "PERIOD"}
        assert r["basis"] in {"", "COVERAGE_COST", "RETAIL"}
        assert r["aggregation"] in {"SUM", "LAST"}
        if r["measure"] in {"DERIVED", "COMPLAINT_COUNT"}:
            continue
        assert r["status_set_id"] in sets, r
        assert r["category_filter"], r  # every counted line names its product scope


def test_funded_reserve_lines_have_both_scopes_and_both_bases():
    defs = [r for r in rows("report_definitions") if r["report_id"] == "FUNDED_RESERVE"]
    items = {r["line_item"] for r in defs}
    assert items == set(data.SUMMARY_COLUMN)  # every line maps to a summary column
    for item in items:
        assert {r["scope"] for r in defs if r["line_item"] == item} == {"ITD", "PERIOD"}, item
    gross = {(r["scope"], r["basis"]) for r in defs if r["line_item"] == "GROSS_CONSIDERATION"}
    assert gross == {("ITD", "COVERAGE_COST"), ("ITD", "RETAIL"), ("PERIOD", "COVERAGE_COST"), ("PERIOD", "RETAIL")}
    assert {r["section"] for r in defs} == {"2", "3"}
    scopes = {(r["category_filter"], r["subcategory_exclude"]) for r in defs if r["measure"] != "DERIVED"}
    assert len(scopes) == 1  # Section 2 and Section 3 describe the same book


def test_claim_count_and_amount_pairs_share_a_population_or_explain_it():
    """A count and a dollar line on different populations carry a note; a like-for-like pair exists."""
    defs = [r for r in rows("report_definitions") if r["report_id"] == "FUNDED_RESERVE"]
    for scope in ("ITD", "PERIOD"):
        line = {r["line_item"]: r for r in defs if r["scope"] == scope}
        count, block, total = line["CLAIMS_PAID_COUNT"], line["CLAIMS_PAID_AMOUNT"], line["CLAIMS_PAID_AMOUNT_ALL"]
        assert (total["date_rule"], total["status_set_id"], total["section"]) == ("PAID_IN_WINDOW", "CONTRACT_ANY", "2")
        assert (block["date_rule"], block["status_set_id"]) != (total["date_rule"], total["status_set_id"])
        assert "not comparable" in block["notes"] and "line 2.6" in block["notes"]
        assert "line 2.6" in count["notes"] and "line 3.2" in count["notes"]
        assert "all contracts" in count["label"] and "all contracts" in total["label"]
        assert "in-force block" in block["label"]


def test_current_reserve_is_the_balance_left_after_claims():
    # line 3.6 must be funded reserve minus claims paid on the same block, not the amount funded at sale
    defs = [r for r in rows("report_definitions") if r["line_item"] == "CURRENT_RESERVE"]
    assert {r["scope"] for r in defs} == {"ITD", "PERIOD"}
    for r in defs:
        assert r["measure"] == "RESERVE_HELD" and r["status_set_id"] == "IN_FORCE"
        assert r["date_rule"] == "IN_FORCE_ENTERED_IN_WINDOW"  # same block as line 3.2
        assert "minus claims paid" in r["notes"] and "line 3.2" in r["notes"], r["scope"]


def test_line_notes_and_collapsed_methodology_notes():
    defs = pd.DataFrame(rows("report_definitions"))
    notes = data.line_notes(defs)
    assert set(notes) == set(data.SUMMARY_COLUMN)
    assert notes["CLAIMS_PAID_AMOUNT"].startswith("ITD: ") and "Period: " in notes["CLAIMS_PAID_AMOUNT"]
    assert "ITD:" not in notes["NET_CONSIDERATION"]  # same note in both scopes -> printed once
    r = [pd.Series({"scope": sc, "basis": b, "notes": f"Basis: {b}; {st}; Shared method note"})
         for sc, st in (("ITD", "FINAL"), ("PERIOD", "FINAL")) for b in ("COVERAGE_COST", "RETAIL")]
    common, diffs = data.split_notes(r)
    assert common == ["FINAL", "Shared method note"]
    assert [(s, b, x) for s, b, x in diffs][0] == ("ITD", "COVERAGE_COST", ["Basis: COVERAGE_COST"])
    assert len(diffs) == 4


def test_market_survey_items_are_numbered_and_use_both_revenue_bases():
    defs = [r for r in rows("report_definitions") if r["report_id"] == "MARKET_SURVEY"]
    assert sorted(int(r["line_no"]) for r in defs) == list(range(1, len(defs) + 1))
    assert {r["scope"] for r in defs} == {"PERIOD"}
    assert {r["basis"] for r in defs if r["measure"] == "CONTRACT_AMOUNT"} == {"COVERAGE_COST", "RETAIL"}
    assert next(r for r in defs if r["line_item"] == "IN_FORCE")["aggregation"] == "LAST"


def test_market_survey_written_amounts_are_not_labelled_as_collected():
    """WRITTEN includes PENDING (not remitted) and CANCELLED (refunds not netted): never call it collected."""
    defs = [r for r in rows("report_definitions") if r["report_id"] == "MARKET_SURVEY"]
    written = {s["status"] for s in rows("report_status_sets") if s["status_set_id"] == "WRITTEN"}
    assert "PENDING" in written
    for r in defs:
        if r["measure"] == "CONTRACT_AMOUNT" and r["status_set_id"] == "WRITTEN":
            assert "collected" not in r["label"].lower(), r["label"]
            assert r["label"].startswith("Gross written")
            assert "pending" in r["notes"].lower() and "cancel" in r["notes"].lower()


def test_survey_window_ends_at_last_closed_period_not_as_of():
    sql = (ROOT / "dbt/models/marts/regulatory/mart_regulatory_windows.sql").read_text(encoding="utf-8")
    assert "last_closed_end" in sql and "as_of_date) as window_end" not in sql
    assert (ROOT / "dbt/tests/regulatory/reg_windows_within_closed_periods.sql").exists()


# --------------------------------------------------------------------------- a tiny warehouse in tmp_path
AS_OF = dt.date(2026, 10, 23)


def _summary_rows() -> pd.DataFrame:
    m = {r["state"]: r for r in rows("state_reserve_method")}["TX"]
    out = []
    for scope, start in (("ITD", None), ("PERIOD", dt.date(2025, 9, 1))):
        for basis in ("COVERAGE_COST", "RETAIL"):
            gross = (1000.0 if scope == "ITD" else 400.0) * (2 if basis == "RETAIL" else 1)
            paid = 300.0 if scope == "ITD" else 0.0
            rate, dep = float(m["reserve_rate"]), float(m["security_deposit_rate"])
            reserve = round(rate * max(gross - paid, 0), 2)
            out.append({
                "state": "TX", "provider_key": "p1", "provider_name": "Obligor 1 Service Corp", "scope": scope,
                "basis": basis, "is_default_basis": basis == m["default_basis"],
                "period_label": "Sep 01, 2025 to Aug 31, 2026", "period_start": start,
                "period_end": dt.date(2026, 8, 31), "is_final": True, "method_code": m["method_code"],
                "reserve_rate": rate, "security_deposit_rate": dep, "contracts_in_force": 2.0,
                "new_contracts_written": 3.0 if scope == "ITD" else 1.0, "complaints_received": 0.0,
                "claims_paid_count": 1.0 if scope == "ITD" else 0.0, "claims_denied_count": 0.0,
                "claims_paid_amount_all": 450.0 if scope == "ITD" else 120.0,
                "gross_consideration": gross, "claims_paid_amount": paid, "net_consideration": gross - paid,
                "calculated_reserve": reserve, "security_deposit": round(dep * reserve, 2),
                "current_reserve": 250.0, "reserve_shortfall": max(reserve - 250.0, 0), "mapped_contracts": 1,
                "notes": f"Basis: {basis}; {'FINAL' if scope == 'ITD' else 'PRELIMINARY'}; "
                         "Genuine zero: no complaints received in the window; Method note",
            })
    return pd.DataFrame(out)


def _survey_rows() -> pd.DataFrame:
    defs = [r for r in rows("report_definitions") if r["report_id"] == "MARKET_SURVEY"]
    out = []
    for year in (2025, 2026):
        for r in defs:
            out.append({
                "state": "CA", "provider_key": "p3", "provider_name": "Obligor 3 Contract Services",
                "calendar_year": year, "year_label": "2026 YTD (Jan-Sep)" if year == 2026 else "2025",
                "year_start": dt.date(year, 1, 1),
                "year_end": dt.date(2026, 9, 30) if year == 2026 else dt.date(2025, 12, 31),
                "is_final": year == 2025, "item_no": int(r["line_no"]), "item": r["line_item"],
                "item_label": r["label"], "basis": r["basis"] or None, "aggregation": r["aggregation"],
                "value": float(int(r["line_no"]) * (year - 2024)),
            })
    return pd.DataFrame(out)


def _contracts() -> pd.DataFrame:
    base = {c: None for c in data.CONTRACT_COLS}
    out = []
    for i, (state, entered) in enumerate((("TX", dt.date(2024, 5, 1)), ("TX", dt.date(2026, 1, 2)),
                                          ("TX", dt.date(2025, 10, 1)), ("CA", dt.date(2025, 3, 3)))):
        r = dict(base, state=state, provider_name="Obligor", provider_basis="CONTRACT_OBLIGOR", dealer_number="D1",
                 dealer_name="Summit Motors", insurer_name="Insurer North", contract_number=f"VSC{i:05d}",
                 status="INFORCE", category="Vsc", subcategory="Powertrain", product="Powertrain 2024",
                 entered_date=entered, effective_date=entered, expiration_date=dt.date(2030, 1, 1),
                 cancel_date=None, coverage_cost=500.0, retail_price=1000.0, coverage_cost_refund=0.0,
                 reserve_amount=125.0, is_in_force_at_period_end=state == "TX", contract_count=1)
        r.update(contract_key=f"k{i}", is_filing_product_scope=True, is_written_itd=state == "TX",
                 is_written_in_period=state == "TX" and entered >= dt.date(2025, 9, 1))
        out.append(r)
    return pd.DataFrame(out)


def _claims() -> pd.DataFrame:
    r = {c: None for c in data.CLAIM_COLS if c != "paid"}
    r.update(state="TX", provider_name="Obligor", dealer_number="D1", insurer_name="Insurer North",
             contract_number="VSC00000", claim_number="1001", category="Vsc", subcategory="Powertrain",
             product="Powertrain 2024", reported_date=dt.date(2025, 2, 1), closed_date=dt.date(2025, 3, 1),
             status="CLOSED_FULLY_PAID", substatus="PAID", claim_count=1)
    r.update(paid_itd=300.0, paid_in_period=0.0, paid_to_date=300.0, is_reported_itd=True,
             is_reported_in_period=False, is_filing_product_scope=True,
             period_start=dt.date(2025, 9, 1), period_end=dt.date(2026, 8, 31))
    claims = pd.DataFrame([r, dict(r, state="CA", claim_number="1002", reported_date=dt.date(2025, 6, 1),
                                   paid_itd=None, paid_in_period=None, is_reported_itd=None,
                                   is_reported_in_period=None, period_start=None, period_end=None)])
    return claims


@pytest.fixture
def tiny_warehouse(tmp_path: Path) -> Path:
    path = tmp_path / "wh.duckdb"
    con = duckdb.connect(str(path))
    for schema in ("seeds", "core", "regulatory"):
        con.execute(f"create schema {schema}")
    for name in ("state_reserve_method", "report_status_sets", "report_definitions"):
        con.execute(f"create table seeds.{name} as select * from read_csv(?, header = true, all_varchar = false)",
                    [str(SEEDS / f"{name}.csv")])
    con.execute("create table core.dim_reporting_date as select date '2026-10-23' as as_of_date")
    frames = {"regulatory.mart_funded_reserve_summary": _summary_rows(),
              "regulatory.mart_market_survey": _survey_rows(),
              "regulatory.mart_funded_reserve_contracts": _contracts(),
              "regulatory.mart_funded_reserve_claims": _claims(),
              "core.fct_cancellation": pd.DataFrame([{"contract_key": "k3", "cancel_date": dt.date(2025, 6, 1),
                                                      "reason": "Customer Request", "method": "PRORATED",
                                                      "initiated_by": "CUSTOMER", "refund_pct": 0.5,
                                                      "customer_refund": 200.0, "coverage_cost_refund": 100.0,
                                                      "cancel_fee": 50.0}])}
    for table, df in frames.items():
        con.register("frame", df)
        con.execute(f"create table {table} as select *, current_timestamp as _built_at from frame")
        con.unregister("frame")
    con.close()
    return path


def test_build_writes_workbooks_and_pages(tiny_warehouse: Path, tmp_path: Path):
    out = tmp_path / "site"
    with duckdb.connect(str(tiny_warehouse), read_only=True) as con:
        pages = build(con, out)
    paths = [p.path for p in pages]
    assert paths[0] == "regulatory/index.html"
    assert {"regulatory/funded-reserve-tx.html", "regulatory/market-survey-ca.html",
            "regulatory/report-definitions.html"} <= set(paths)
    assert all(isinstance(p, Page) for p in pages)
    assert all(p.nav == "regulatory/index.html" for p in pages[1:])  # sub-pages keep the nav tab active

    wb = openpyxl.load_workbook(out / "downloads/regulatory/funded_reserve_tx_20260831.xlsx", read_only=True)
    assert wb.sheetnames == ["Summary", "ITD Contracts", "Period Contracts", "ITD Claims", "Period Claims"]
    summary = [r for r in wb["Summary"].iter_rows(values_only=True) if r]
    # V7 edit 52: column A is the form question the line answers (Q2.1, Q3.4, ...)
    by_label = {r[1]: r for r in summary if len(r) > 2 and isinstance(r[0], str) and r[0][:1] == "Q"}
    assert by_label["Contracts in force at period end"][0] == "Q2.1"
    assert by_label["Contracts in force at period end"][2:4] == (2, 2)
    assert by_label["New contracts written"][2:4] == (3, 1)
    # Section 3: default basis first (coverage cost for TX), retail beside it; reserve = 40% x (1000 - 300)
    assert by_label["Calculated reserve"][2] == pytest.approx(280.0)
    assert by_label["Gross consideration (in force)"][4] == pytest.approx(2000.0)
    # line 2.6 pairs the 2.4 count with all-contracts dollars; 3.2 says it is the in-force block only
    assert by_label["Claims paid, all contracts (amount)"][2:4] == (pytest.approx(450.0), pytest.approx(120.0))
    assert "not comparable" in by_label["Total claims paid, in-force block only"][6]
    assert "line 2.6" in by_label["Claims paid, all contracts (count)"][6]
    # methodology notes: the shared text once, then only what differs by scope / basis
    notes = [r for r in summary if r[0] == "All" or (isinstance(r[0], str) and r[0].startswith(("ITD ", "PERIOD ")))]
    assert sum("Genuine zero" in str(r[1]) for r in notes) == 1
    assert notes[0][0] == "All" and "Method note" in notes[0][1]
    itd = list(wb["ITD Contracts"].iter_rows(values_only=True))
    assert itd[0][0] == "State" and len(itd) == 1 + 3 + 1  # header, 3 TX contracts, totals row
    assert itd[-1][0].startswith("Total (3 rows)")
    period = list(wb["Period Contracts"].iter_rows(values_only=True))
    assert len(period) == 1 + 2 + 1

    sv = openpyxl.load_workbook(out / "downloads/regulatory/market_survey_ca_2026.xlsx", read_only=True)
    assert sv.sheetnames == ["Obligor 3 Contract Services", "Contracts", "Cancels", "Claims"]
    grid = list(sv["Obligor 3 Contract Services"].iter_rows(values_only=True))
    assert grid[2] == ("Item", "Question", "2025", "2026 YTD (Jan-Sep)", "Total")
    item1, item2 = grid[3], grid[4]
    assert item1[-1] == item1[-2]  # in force: total = latest year (LAST)
    assert item2[-1] == item2[2] + item2[3]  # flows: total = sum (SUM)
    # the seed notes travel with the workbook: items 3 and 4 say they are gross written, not cash
    start = next(i for i, r in enumerate(grid) if r and r[0] == "Notes") + 1
    noted = {r[0]: r[1] for r in grid[start:] if r and isinstance(r[0], int)}
    assert "not cash collected" in noted[3] and "not cash collected" in noted[4]
    html = next(p for p in pages if p.path == "regulatory/market-survey-ca.html")
    assert html.sections[0].cols[-1] == "Note"


def test_survey_matrix_totals_follow_the_aggregation_rule():
    m = excel.survey_matrix(_survey_rows())
    last = m[m.Item == 1].iloc[0]
    flow = m[m.Item == 2].iloc[0]
    assert last["Total"] == last["2026 YTD (Jan-Sep)"]
    assert flow["Total"] == flow["2025"] + flow["2026 YTD (Jan-Sep)"]


def test_plugin_runs_through_the_shell_hook(tiny_warehouse: Path, tmp_path: Path):
    import reports.regulatory as mod

    with duckdb.connect(str(tiny_warehouse), read_only=True) as con:
        pages, nav = plugin_pages(con, tmp_path / "site", [("regulatory", mod)])
    assert nav == [("regulatory/index.html", NAV_LABEL)]
    assert len(pages) == 4


def test_standalone_cli_renders_html_without_broken_links(tiny_warehouse: Path, tmp_path: Path):
    out = tmp_path / "site"
    written = reg_cli.run(tiny_warehouse, out)
    assert len(written) == 4
    assert broken_links(out) == []
    html = (out / "regulatory/funded-reserve-tx.html").read_text(encoding="utf-8")
    assert "Section 2" in html and "Section 3" in html and "Methodology notes" in html
    assert "funded_reserve_tx_20260831.xlsx" in html
    index = (out / "regulatory/index.html").read_text(encoding="utf-8")
    assert "States with no filing requirement" in index


def test_plugin_skips_when_marts_are_missing(tmp_path: Path):
    with duckdb.connect(str(tmp_path / "empty.duckdb")) as con:
        assert build(con, tmp_path / "site") == []


# --------------------------------------------------------------------------- against a built warehouse
def _has_marts() -> bool:
    if not WAREHOUSE.exists():
        return False
    try:
        with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
            return data.available(con)
    except duckdb.Error:  # locked by a writer, or an older file
        return False


needs_marts = pytest.mark.skipif(not _has_marts(), reason="regulatory marts not built in WAREHOUSE_PATH")


@needs_marts
def test_warehouse_filing_skips_states_without_requirement_and_obeys_the_method():
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        s = con.execute("select * from regulatory.mart_funded_reserve_summary").df()
        m = con.execute("select * from seeds.state_reserve_method").df().set_index("state")
    assert set(s.state) <= set(m.index[m.has_requirement])
    assert len(s) and set(s.scope) == {"ITD", "PERIOD"} and set(s.basis) == {"COVERAGE_COST", "RETAIL"}
    base = s.gross_consideration.where(s.method_code == "GROSS_CONSIDERATION", s.net_consideration)
    expected = (s.state.map(m.reserve_rate) * base.clip(lower=0)).round(2)
    assert (s.calculated_reserve - expected).abs().max() < 0.01 + 1e-6  # DuckDB half-up vs pandas half-even
    assert s.notes.str.contains("Basis:").all()
    # basis changes Section 3 only
    k = ["state", "provider_key", "scope"]
    sec2 = s.groupby(k)[["contracts_in_force", "new_contracts_written", "claims_paid_count"]].nunique()
    assert (sec2 == 1).all().all()


@needs_marts
def test_warehouse_build_produces_every_filing(tmp_path: Path):
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        pages = build(con, tmp_path)
        states = [r[0] for r in con.execute(
            "select distinct state from regulatory.mart_funded_reserve_summary order by 1").fetchall()]
    files = sorted(p.name for p in (tmp_path / "downloads" / "regulatory").glob("funded_reserve_*.xlsx"))
    assert [f.split("_")[2].upper() for f in files] == states
    assert {f"regulatory/funded-reserve-{s.lower()}.html" for s in states} <= {p.path for p in pages}


# --------------------------------------------------------------------------- shortfall wording (V3)
def test_shortfall_cell_says_filed_for_final_and_action_only_when_open():
    from reports.regulatory.pages import shortfall_cell
    assert shortfall_cell(0.0, True)["status"] == "good"
    filed = shortfall_cell(309_393.7, True)
    assert "Filed – shortfall" in filed["text"] and filed["status"] == "warning"
    assert "Action needed" not in filed["text"]
    open_ = shortfall_cell(10.0, False)
    assert "Action needed" in open_["text"] and open_["status"] == "critical"


def test_shortfall_cause_names_the_basis_effect_and_the_expected_action():
    from reports.regulatory.pages import shortfall_cause
    assert shortfall_cause("RETAIL", 0.0) == ""
    retail = shortfall_cause("RETAIL", 1.0)
    assert "retail" in retail and "coverage cost" in retail and "top-up deposit" in retail
    assert "top-up deposit" in shortfall_cause("COVERAGE_COST", 1.0)


def test_index_labels_totals_and_splits_shortfall_by_basis(tiny_warehouse: Path, tmp_path: Path):
    with duckdb.connect(str(tiny_warehouse), read_only=True) as con:
        pages = build(con, tmp_path / "site")
    index = pages[0]
    kpis = {k.label: k for k in index.kpis}
    for label in ("Calculated reserve", "Reserve shortfall"):
        assert "period ends vary; default basis per state" in kpis[label].sub
    # TX (coverage cost) is short 280 - 250 = 30 on a Final filing: filed, not "Action needed"
    short = kpis["Reserve shortfall"]
    assert "retail-basis states $0" in short.sub and "coverage-cost states $30" in short.sub
    assert short.status == "warning"
    table = next(s for s in index.sections if s.title == "Funded-reserve filings by state")
    row = dict(zip(table.cols, table.rows[0], strict=True))
    assert row["Basis"] == "coverage cost" and "Filed – shortfall" in row["Shortfall"]["text"]
    assert "top-up deposit" in row["Shortfall cause"]["text"]
    split = next(s for s in index.sections if s.title == "Shortfall by basis")
    assert [r[0] for r in split.rows] == ["Retail-basis states", "Coverage-cost states"]
    assert split.rows[1][1] == "TX" and split.rows[0][1] == "none"
    tx = next(p for p in pages if p.path == "regulatory/funded-reserve-tx.html")
    why = next(s for s in tx.sections if s.title == "Why this filing shows a shortfall")
    assert "top-up deposit" in why.caption and why.rows[2][1]["status"] == "warning"
    assert any(k.label == "Reserve shortfall" and "Filed – shortfall" in k.sub for k in tx.kpis)


@needs_marts
def test_warehouse_retail_basis_shortfall_is_explained_on_its_pages(tmp_path: Path):
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        pages = {p.path: p for p in build(con, tmp_path)}
        s = con.execute("select state, basis, sum(reserve_shortfall) short from regulatory.mart_funded_reserve_summary "
                        "where scope = 'ITD' and is_default_basis group by all").df()
    for r in s[(s.short > 0) & (s.basis == "RETAIL")].itertuples():
        page = pages[f"regulatory/funded-reserve-{r.state.lower()}.html"]
        why = next(x for x in page.sections if x.title == "Why this filing shows a shortfall")
        assert "Basis effect" in why.caption and "top-up deposit" in why.caption


# --------------------------------------------------------------------------- V6 edit 5: state in-force list
def _inforce_rows() -> pd.DataFrame:
    base = {c: None for c in data.INFORCE_COLS}
    out = []
    for i, (state, year, form) in enumerate((("TX", 2022, None), ("TX", 2024, "VSC-24-TX"),
                                             ("TX", 2025, "VSC-25-TX"), ("CA", 2025, "VSC-25-CA"))):
        out.append(dict(base, state=state, contract_date=dt.date(year, 3, 1), contract_number=f"VSC{i:05d}",
                        product_name="Powertrain 2024 - FI", category="Vsc", subcategory="Powertrain",
                        program_type="FULL_INSURED", dealer_number="D1", dealer_name="Summit Motors",
                        lienholder_name="Harbor Bank", current_form_code=form,
                        form_code_source="document" if form else "missing", obligor_name="Obligor",
                        obligor_kind="internal_provider", status="INFORCE", effective_date=dt.date(year, 3, 1),
                        expiration_date=dt.date(2030, 1, 1), contract_count=1, list_as_of_date=dt.date(2026, 10, 23)))
    return pd.DataFrame(out)


@pytest.fixture
def inforce_warehouse(tiny_warehouse: Path) -> Path:
    with duckdb.connect(str(tiny_warehouse)) as con:
        con.register("frame", _inforce_rows())
        con.execute("create table regulatory.mart_state_inforce_list as select *, current_timestamp as _built_at "
                    "from frame")
    return tiny_warehouse


def test_inforce_list_columns_cover_the_compliance_ask():
    # contract date, number, product, dealer, lienholder name / address / state, form code, obligor
    need = {"contract_date", "contract_number", "product_name", "dealer_name", "lienholder_name",
            "lienholder_address", "lienholder_state", "current_form_code", "obligor_name", "obligor_kind"}
    assert need <= set(data.INFORCE_COLS)
    sql = (ROOT / "dbt/models/marts/regulatory/mart_state_inforce_list.sql").read_text(encoding="utf-8")
    assert "status_set_id = 'IN_FORCE'" in sql and "is_filing_product_scope" in sql  # rules from seeds/filing
    assert "'inferred'" not in sql  # the form code is never guessed from product x state


def test_form_code_coverage_counts_missing_by_contract_year():
    cov = data.form_code_coverage(_inforce_rows())
    by_year = cov.set_index("Contract year")
    assert by_year.loc[2022, "Missing form code count"] == 1 and by_year.loc[2022, "Form code share"] == 0
    assert by_year.loc[2025, "With form code count"] == 2 and by_year.loc[2025, "Form code share"] == 1


def test_build_writes_inforce_workbook_and_page_per_state(inforce_warehouse: Path, tmp_path: Path):
    out = tmp_path / "site"
    with duckdb.connect(str(inforce_warehouse), read_only=True) as con:
        pages = {p.path: p for p in build(con, out)}
    assert {"regulatory/inforce-list-tx.html", "regulatory/inforce-list-ca.html"} <= set(pages)
    wb = openpyxl.load_workbook(out / "downloads/regulatory/inforce_service_contracts_tx_20261023.xlsx",
                                read_only=True)
    assert wb.sheetnames == ["In Force Contracts", "By Product", "Form Code Coverage"]
    rows_ = list(wb["In Force Contracts"].iter_rows(values_only=True))
    assert list(rows_[0]) == list(data.INFORCE_COLS.values())
    assert len(rows_) == 1 + 3 + 1  # header, 3 TX contracts, totals row
    cov = list(wb["Form Code Coverage"].iter_rows(values_only=True))
    assert (list(cov[-1]) + [None] * 5)[4] is None  # the share column has no (meaningless) total
    tx = pages["regulatory/inforce-list-tx.html"]
    kpis = {k.label: k for k in tx.kpis}
    assert kpis["Contracts in force"].value == "3" and kpis["Missing form code"].status == "warning"
    index = pages["regulatory/index.html"]
    section = next(s for s in index.sections if s.title == "State in-force contract lists")
    assert [r[1] for r in section.rows] == ["1", "3"]  # CA, TX


def test_inforce_pages_have_no_broken_links(inforce_warehouse: Path, tmp_path: Path):
    out = tmp_path / "site"
    written = reg_cli.run(inforce_warehouse, out)
    assert len(written) == 6
    assert broken_links(out) == []


def test_plugin_without_inforce_mart_keeps_the_filings_only(tiny_warehouse: Path, tmp_path: Path):
    with duckdb.connect(str(tiny_warehouse), read_only=True) as con:
        assert not data.inforce_available(con)
        paths = [p.path for p in build(con, tmp_path / "site")]
    assert not any("inforce-list" in p for p in paths)


def _has_inforce() -> bool:
    if not _has_marts():
        return False
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        return data.inforce_available(con)


@pytest.mark.skipif(not _has_inforce(), reason="mart_state_inforce_list not built in WAREHOUSE_PATH")
def test_warehouse_inforce_list_follows_the_filing_rules():
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        bad = con.execute("""
            select count(*) from regulatory.mart_state_inforce_list as l
            left join regulatory.mart_funded_reserve_contracts as c using (contract_key)
            where not c.is_filing_product_scope
               or l.status not in (select status from seeds.report_status_sets where status_set_id = 'IN_FORCE')
               or l.effective_date > l.list_as_of_date
               or coalesce(c.cancel_date, date '9999-12-31') <= l.list_as_of_date
               or (l.form_code_source = 'missing') <> (l.current_form_code is null)
        """).fetchone()[0]
        n = con.execute("select count(*) from regulatory.mart_state_inforce_list").fetchone()[0]
    assert n > 0 and bad == 0
