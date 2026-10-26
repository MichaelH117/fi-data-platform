"""Dealer & claims area: seeds, dbt models, report plug-ins (reports.dealer, reports.claims).

Unit tests build a tiny DuckDB with the mart columns the reports read, so they pass without a built
warehouse. Tests against the real marts skip when the dealer/claims marts have not been built.
"""

from __future__ import annotations

import csv
import os
import re
from pathlib import Path

import duckdb
import pandas as pd
import pytest
from jinja2 import Environment, FileSystemLoader, select_autoescape
from markupsafe import Markup

import reports.claims as claims_report
import reports.dealer as dealer_report
from reports.__main__ import broken_links
from reports.build import HERE as REPORTS_DIR
from reports.build import Page, discover_plugins
from reports.dealer import watch as dealer_watch
from reports.dealer.excel import clean, write_workbook

ROOT = Path(__file__).resolve().parents[1]
SEEDS = ROOT / "dbt" / "seeds" / "dealer"
DEALER_MODELS = ROOT / "dbt" / "models" / "marts" / "dealer"
CLAIMS_MODELS = ROOT / "dbt" / "models" / "marts" / "claims"
WAREHOUSE = Path(os.environ.get("WAREHOUSE_PATH", ROOT / "data" / "warehouse" / "warehouse.duckdb"))
LAST_MONTH = pd.Timestamp("2026-09-01")


def seed(name: str) -> list[dict]:
    with (SEEDS / f"{name}.csv").open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# --------------------------------------------------------------------------- seeds (counting rules)
def test_status_sets_never_count_voided_or_denied():
    rows = seed("dealer_report_status_sets")
    reports = {r["report_id"] for r in rows}
    assert reports >= {"DEALER_PRODUCTION", "DEALER_PENDING", "SPECIALTY_LINE", "SURCHARGE_AUDIT"}
    assert not {r["status"] for r in rows} & {"VOIDED", "DENIED", "REINSTATEMENT_VOIDED", "REINSTATEMENT_DENIED"}
    assert {r["status"] for r in rows if r["report_id"] == "DEALER_PENDING"} == {"PENDING"}
    assert len({(r["report_id"], r["status"]) for r in rows}) == len(rows)


def test_surcharge_rules_one_amount_per_code():
    rows = seed("dealer_surcharge_rules")
    amounts: dict[str, set[str]] = {}
    for r in rows:
        amounts.setdefault(r["surcharge_code"], set()).add(r["expected_amount"])
        assert r["vehicle_attribute"] in {"fuel_type", "drive_type", "is_commercial", "has_title_brand",
                                          "class_cylinders"}, r
    assert all(len(v) == 1 for v in amounts.values()), amounts
    assert {"DIESEL", "4X4", "HYBRID_EV", "TURBO", "COMMERCIAL", "BRANDED_TITLE"} == set(amounts)


def test_age_bands_are_contiguous():
    rows = seed("dealer_age_bands")
    for band_set in {r["band_set"] for r in rows}:
        bands = sorted((int(r["band_order"]), int(r["min_days"]), int(r["max_days"]))
                       for r in rows if r["band_set"] == band_set)
        assert bands[0][1] == 0, band_set
        for (_, _, hi), (_, lo, _) in zip(bands, bands[1:], strict=False):
            assert lo == hi + 1, f"{band_set}: gap or overlap at {hi}->{lo}"


def test_remittance_aging_uses_the_one_shared_band_seed():
    """One remittance aging-band definition (finance seed remit_aging_buckets) for both remittance marts."""
    assert "REMITTANCE" not in {r["band_set"] for r in seed("dealer_age_bands")}
    sql = (DEALER_MODELS / "mart_remittance_aging.sql").read_text(encoding="utf-8")
    assert "ref('remit_aging_buckets')" in sql and "dealer_age_bands" not in sql
    assert (ROOT / "dbt" / "tests" / "dealer" / "assert_remittance_late_fees_agree_with_finance.sql").is_file()


def test_scorecard_watch_rules_live_in_the_seed():
    rules = {r["rule_id"]: r["rule_value"] for r in seed("dealer_claim_kpi_rules")}
    sql = (DEALER_MODELS / "mart_dealer_scorecard.sql").read_text(encoding="utf-8")
    ae_sql = (DEALER_MODELS / "mart_dealer_early_claim_ae.sql").read_text(encoding="utf-8")
    for rule in ("scorecard_watch_loss_ratio", "scorecard_min_t12_earned", "scorecard_watch_late_rate",
                 "early_claim_headline_days"):
        assert rule in rules and float(rules[rule]) > 0, rule
        assert f"'{rule}'" in sql, rule
    for rule in AE_RULES:
        assert rule in rules and float(rules[rule]) > 0, rule
        assert f"'{rule}'" in ae_sql, rule
    # the old share-of-claims watch is gone: one early-claim watch rule of record
    assert "scorecard_watch_early_claim_rate" not in rules and "scorecard_watch_early_claim_rate" not in sql
    assert 0 < float(rules["scorecard_watch_loss_ratio"]) < 2


AE_RULES = ("scorecard_early_claim_ae_window_days", "scorecard_early_claim_min_exposed",
            "scorecard_early_claim_min_ae", "scorecard_early_claim_min_z")


def test_early_claim_watch_seed_matches_actuarial_rule():
    """The scorecard's early-claim watch is the Actuarial page's A/E test: same window and thresholds."""
    import inspect

    from actuarial.emergence import early_claim_by_dealer

    rules = {r["rule_id"]: float(r["rule_value"]) for r in seed("dealer_claim_kpi_rules")
             if r["rule_id"].startswith("scorecard_")}
    sig = inspect.signature(early_claim_by_dealer).parameters
    assert rules["scorecard_early_claim_min_exposed"] == sig["min_exposed"].default
    assert rules["scorecard_early_claim_ae_window_days"] in sig["windows"].default
    src = inspect.getsource(early_claim_by_dealer)
    assert f'["actual_to_expected"] >= {rules["scorecard_early_claim_min_ae"]:g}' in src
    assert f'["z_score"] >= {rules["scorecard_early_claim_min_z"]:g}' in src
    sql = (DEALER_MODELS / "mart_dealer_scorecard.sql").read_text(encoding="utf-8")
    assert "ref('mart_dealer_early_claim_ae')" in sql
    assert "coalesce(ae.is_flagged, false) as is_watch_early_claims" in sql
    assert (ROOT / "dbt" / "tests" / "dealer" / "assert_scorecard_early_claim_watch_matches_ae.sql").is_file()


def test_early_claim_labels_name_the_basis(mini_wh, tmp_path):
    """Share of claims (claims page, scorecard column) and A/E per contract (watch) never share a label."""
    card, _ = dealer_report.page_scorecard(mini_wh, tmp_path, LAST_MONTH)
    cols = card.sections[0].cols
    assert dealer_report.EARLY_SHARE_LABEL in cols and any(c.startswith(dealer_report.EARLY_AE_LABEL) for c in cols)
    assert "Early claims" not in cols
    assert "same dealer list as the Actuarial page" in card.sections[0].caption
    pages = claims_report.build(mini_wh, tmp_path)
    labels = [k.label for p in pages for k in p.kpis]
    assert "Early-claim share of claims (90 days), T12" in labels
    assert not any(lbl.startswith("Early claims (") for lbl in labels)


def test_early_claim_windows_seed_matches_report_and_model():
    rules = {r["rule_id"]: r["rule_value"] for r in seed("dealer_claim_kpi_rules")}
    windows = [int(x) for x in rules["early_claim_windows"].split("|")]
    assert windows == claims_report.WINDOWS
    sql = (CLAIMS_MODELS / "mart_claims_kpi.sql").read_text(encoding="utf-8")
    assert f"{{% set windows = {windows} %}}" in sql


def test_specialty_rv_requires_chassis_vin():
    rows = {r["subcategory_code"]: r for r in seed("dealer_specialty_lines")}
    assert rows["RV"]["requires_chassis_vin"] == "true"


# --------------------------------------------------------------------------- dbt model hygiene
def test_marts_read_core_only_except_documented_pii():
    """Marts build on core (and seeds); only documented exceptions reach into staging/intermediate."""
    allowed = {
        "mart_dealer_pending_contracts.sql": {"int_contract_key_bridge", "stg_admin_core__contract",
                                              "stg_legacy_admin__contracts", "stg_network_feed__deals"},
        "mart_surcharge_audit.sql": {"stg_admin_core__contract_surcharge"},
        "mart_claim_summary.sql": {"stg_admin_core__claim_component", "int_claim_number_bridge"},
        "mart_claim_component_lines.sql": {"stg_admin_core__claim_component", "int_claim_number_bridge"},
    }
    for path in [*DEALER_MODELS.glob("*.sql"), *CLAIMS_MODELS.glob("*.sql")]:
        refs = set(re.findall(r"ref\('([^']+)'\)", path.read_text(encoding="utf-8")))
        upstream = {r for r in refs if r.startswith(("stg_", "int_"))}
        assert upstream <= allowed.get(path.name, set()), f"{path.name} reads {upstream}"


def test_only_pending_mart_flagged_pii():
    for yml in (DEALER_MODELS / "_dealer.yml", CLAIMS_MODELS / "_claims.yml"):
        blocks = yml.read_text(encoding="utf-8").split("\n  - name: ")[1:]
        flagged = [b.split("\n")[0].strip() for b in blocks if "contains_pii: true" in b]
        assert flagged == (["mart_dealer_pending_contracts"] if yml.parent == DEALER_MODELS else []), flagged


# --------------------------------------------------------------------------- tiny warehouse fixture
def _months(n: int = 30) -> list[pd.Timestamp]:
    return list(pd.date_range(end=LAST_MONTH + pd.offsets.MonthBegin(1), periods=n, freq="MS"))


@pytest.fixture()
def mini_wh(tmp_path: Path) -> duckdb.DuckDBPyConnection:
    con = duckdb.connect(str(tmp_path / "mini.duckdb"))
    for s in ("core", "dealer", "claims"):
        con.execute(f"create schema {s}")
    months = _months()
    dealers = [("D1", "Summit Motors"), ("D2", "Harbor Auto")]
    reporting = pd.DataFrame({"as_of_date": [pd.Timestamp("2026-10-23")], "last_complete_month": [LAST_MONTH]})
    prod = pd.DataFrame([
        dict(tenant_key=1, dealer_key=k, dealer_number=k, dealer_name=n, dealer_state="TX", category_code=cat,
             subcategory_code=sub, subcategory=sub, month_start=m, entered_month_key=int(m.strftime("%Y%m")),
             contracts_written=3 + i, contracts_pending=1, contracts_in_force=1, contracts_cancelled=1,
             coverage_cost=1000.0, retail_price=2000.0, dealer_markup=1000.0)
        for i, (k, n) in enumerate(dealers) for m in months for cat, sub in [("VSC", "VSC-PT"), ("GAP", "GAP")]
    ])
    pend = pd.DataFrame([
        dict(contract_key=f"c{i}", tenant_key=1, tenant_code="T1", administrator_name="Admin One",
             dealer_key=dealers[i % 2][0], dealer_number=dealers[i % 2][0], dealer_name=dealers[i % 2][1],
             contract_number=f"VSC-{i:06d}", customer_first_name="Pat", customer_last_name=f"Sample{i}",
             vin_token="abc123", product_code="VSC-PT", product="Powertrain", plan="P1", subcategory_code="VSC-PT",
             date_sold=pd.Timestamp("2026-09-01") - pd.Timedelta(days=d), effective_date=pd.Timestamp("2026-09-01"),
             lienholder_name=None if i % 2 else "Redwood Lending", provider_name="Obligor 1 Service Corp",
             retail_price=1500.0, coverage_cost=700.0, days_pending=d,
             pending_age_band="90+ days" if d > 90 else "0-15 days", pending_age_band_order=5 if d > 90 else 1,
             as_of_date=pd.Timestamp("2026-10-23"))
        for i, d in enumerate([5, 10, 95, 120])
    ])
    sur = pd.DataFrame([
        dict(contract_key=f"s{i}", tenant_key=1, dealer_key="D1", dealer_number="D1", dealer_name="Summit Motors",
             contract_number=f"VSC-{100 + i}", contract_status="INFORCE", subcategory_code="VSC-PT",
             contract_date=pd.Timestamp("2026-08-01"), vin_token="v", fuel_type="DIESEL", drive_type="2WD",
             is_commercial=False, title_brand=None, surcharge_code=code, is_expected=f != "NOT_EXPECTED",
             is_charged=f != "MISSING", finding=f, expected_reason="x",
             expected_amount=0.0 if f == "NOT_EXPECTED" else 150.0, charged_amount=0.0 if f == "MISSING" else 150.0,
             amount_variance=0.0)
        for i, (code, f) in enumerate([("DIESEL", "OK"), ("DIESEL", "MISSING"), ("4X4", "NOT_EXPECTED")])
    ])
    sur.loc[sur["finding"] == "MISSING", "amount_variance"] = -150.0
    sur.loc[sur["finding"] == "NOT_EXPECTED", "amount_variance"] = 150.0
    spec = pd.DataFrame([
        dict(contract_key=f"r{i}", tenant_key=1, dealer_key="D2", dealer_number="D2", dealer_name="Harbor Auto",
             contract_number=f"RV-{i}", status="INFORCE", subcategory_code=sub, line_name=line, product="RV plan",
             entered_date=m, effective_date=m, entered_month=m, vin_token="vt",
             chassis_vin_token=None if miss else "ct",
             has_chassis_vin=not miss, requires_chassis_vin=sub == "RV", is_chassis_vin_missing=miss and sub == "RV",
             vehicle_make="Coachline", vehicle_model="Voyager 30", vehicle_year=2024, new_used="NEW",
             purchase_price=90000.0, coverage_cost=3000.0, retail_price=5000.0, claim_count=1, claims_paid=500.0)
        for i, (sub, line, m, miss) in enumerate([("RV", "Recreational vehicle", LAST_MONTH, False),
                                                  ("RV", "Recreational vehicle", months[3], True),
                                                  ("PSP", "Powersports", LAST_MONTH, True)])
    ])
    remit = pd.DataFrame([
        dict(tenant_key=1, dealer_key=k, dealer_number=k, dealer_name=n, posting_period_key=int(m.strftime("%Y%m")),
             age_band_order=b, age_band=lbl, cash_transactions=10, late_transactions=2 if b > 1 else 0,
             remitted_amount=5000.0, late_fees=40.0 if b > 1 else 0.0, avg_remit_age_days=12.0, max_remit_age_days=50)
        for k, n in dealers for m in months[-13:-1] for b, lbl in [(1, "0-30 days"), (2, "31-60 days")]
    ])
    kpi_rows = []
    for m in months:
        for sub in ("VSC-PT", "GAP"):
            r = dict(tenant_key=1, subcategory_code=sub, month_start=m, exposure_years=100.0, claims_reported=12,
                     claims_denied=2, claim_count=10, paid_claim_count=8, open_claims=2, open_claims_aged_30_plus=1,
                     paid_amount_on_paid_claims=8000.0, paid_amount=8200.0, incurred_amount=9000.0,
                     turnaround_days_total=80, turnaround_claims=10)
            r.update({f"early_claims_{w}": min(10, w // 30) for w in claims_report.WINDOWS})
            kpi_rows.append(r)
    kpi = pd.DataFrame(kpi_rows)
    kpi["claim_frequency"] = kpi["claim_count"] / kpi["exposure_years"]
    comp = pd.DataFrame([dict(component_code=c, status_group=g, month_start=m, loss_month_key=int(m.strftime("%Y%m")),
                              component_sources="claim_component", claims=3, component_lines=4, parts_amount=500.0,
                              labor_amount=200.0, labor_hours=2.0, component_cost=700.0 * (j + 1))
                         for m in months[-12:] for j, c in enumerate(["ENGINE", "TRANSMISSION"])
                         for g in ("paid", "denied")])
    summ = pd.DataFrame([dict(claim_key="k1", status_group="open", open_age_band="30-59 days", incurred_amount=900.0,
                              open_age_days=40),
                         dict(claim_key="k2", status_group="paid", open_age_band=None, incurred_amount=500.0,
                              open_age_days=None)])
    score = pd.DataFrame([
        dict(tenant_key=1, dealer_key=k, dealer_number=k, dealer_name=n, dealer_state="TX", month_start=m,
             month_key=int(m.strftime("%Y%m")), t12_contracts_written=100, t12_cancel_rate=0.1,
             t12_earned_claims_reserve=50000.0, t12_incurred_claims=50000.0 * lr, t12_loss_ratio=lr,
             t12_loss_ratio_peer_pct=pp, t12_early_claim_rate=0.02, t12_late_remit_rate=0.1,
             t12_avg_days_late=None if i else 4.0, t12_late_fees=120.0,
             early_claim_actual_to_expected=1.0 + i,
             is_watch_loss_ratio=lr > 0.8, is_watch_late_remit=False, is_watch_early_claims=False,
             is_watch=lr > 0.8)
        for i, (k, n, lr, pp) in enumerate([("D1", "Summit Motors", 0.9, 1.0), ("D2", "Harbor Auto", 0.4, 0.0)])
        for m in months
    ])
    access = pd.DataFrame([dict(user_key=f"u{k}", user_login=f"{k.lower()}@example.test", role="DEALER",
                                tenant_key=1, tenant_code="T1", dealer_key=k, dealer_number=k) for k, _ in dealers])
    for name, df in [("core.dim_reporting_date", reporting), ("core.dim_user_dealer_access", access),
                     ("dealer.mart_dealer_production", prod),
                     ("dealer.mart_dealer_scorecard", score),
                     ("dealer.mart_dealer_pending_contracts", pend), ("dealer.mart_surcharge_audit", sur),
                     ("dealer.mart_specialty_line", spec), ("dealer.mart_remittance_aging", remit),
                     ("claims.mart_claims_kpi", kpi), ("claims.mart_claim_components", comp),
                     ("claims.mart_claim_summary", summ)]:
        con.register("_df", df)
        con.execute(f"create table {name} as select * from _df")
        con.unregister("_df")
    yield con
    con.close()


def render(pages: list[Page], out: Path) -> None:
    env = Environment(loader=FileSystemLoader(REPORTS_DIR / "templates"), autoescape=select_autoescape(["html", "j2"]),
                      trim_blocks=True, lstrip_blocks=True)
    tpl = env.get_template("page.html.j2")
    (out / "assets").mkdir(parents=True, exist_ok=True)
    for a in ("site.css", "site.js", "plotly.min.js"):
        (out / "assets" / a).write_text("", encoding="utf-8")
    (out / "index.html").write_text("<!doctype html>", encoding="utf-8")  # the shell's home page (brand link)
    paths = {p.path for p in pages}
    nav = [(h, lbl) for h, lbl in [(dealer_report.HUB, dealer_report.NAV_LABEL),
                                   (claims_report.PAGE, claims_report.NAV_LABEL)] if h in paths]
    for p in pages:
        (out / p.path).parent.mkdir(parents=True, exist_ok=True)
        html = tpl.render(page=p, root="../" * p.path.count("/"), nav=nav, active=p.nav or p.path,
                          as_of=pd.Timestamp("2026-10-23"), last_month="Sep 2026", chart=lambda s: Markup(s))
        (out / p.path).write_text(html, encoding="utf-8")


# --------------------------------------------------------------------------- report plug-ins
def test_plugins_are_discovered_by_the_shell():
    found = dict(discover_plugins())
    assert found.get("dealer") is dealer_report and found.get("claims") is claims_report


def test_dealer_build_pages_excel_and_links(mini_wh, tmp_path):
    out = tmp_path / "site"
    pages = dealer_report.build(mini_wh, out)
    assert all(isinstance(p, Page) for p in pages)
    assert pages[0].path == dealer_report.HUB
    assert {p.path for p in pages} == {dealer_report.HUB, *[r[0] for r in dealer_report.REPORTS]}
    assert all(p.nav == dealer_report.HUB for p in pages[1:])
    for name in ("dealer_production", "dealer_pending_contracts", "surcharge_audit", "specialty_line",
                 "remittance_aging", "dealer_scorecard"):
        assert (out / "downloads" / "dealer" / f"{name}.xlsx").is_file(), name
    render(pages, out)
    assert broken_links(out) == []


def test_production_kpi_and_pivot(mini_wh, tmp_path):
    page, xlsx = dealer_report.page_production(mini_wh, tmp_path, LAST_MONTH)
    # 2 dealers x 2 subcats x 12 months; dealer D1 writes 3/row, D2 writes 4/row.
    assert page.kpis[0].value == f"{(3 + 4) * 2 * 12:,}"
    summary = pd.read_excel(xlsx, sheet_name="Summary T12")
    assert summary["Total"].sum() == (3 + 4) * 2 * 12
    assert list(pd.ExcelFile(xlsx).sheet_names) == ["About", "Summary T12", "Pivot by month", "Production"]


def test_pending_report_counts_and_fake_names(mini_wh, tmp_path):
    page, xlsx = dealer_report.page_pending(mini_wh, tmp_path)
    assert page.kpis[0].value == "4"
    assert page.kpis[3].value == "2" and page.kpis[3].status == "warning"
    detail = pd.read_excel(xlsx, sheet_name="Pending contracts")
    assert len(detail) == 4 and detail["days_pending"].is_monotonic_decreasing
    assert "customer_last_name" in detail.columns and not any(c.startswith("_") for c in detail.columns)


def test_surcharge_findings(mini_wh, tmp_path):
    page, xlsx = dealer_report.page_surcharge(mini_wh, tmp_path)
    k = {x.label: x for x in page.kpis}
    assert k["Missing surcharges"].value == "1" and k["Missing surcharges"].status == "critical"
    assert k["Charged but not expected"].value == "1"
    exc = pd.read_excel(xlsx, sheet_name="Exceptions")
    assert set(exc["finding"]) == {"MISSING", "NOT_EXPECTED"}


def test_specialty_line_flags_missing_chassis_and_monthly_tab(mini_wh, tmp_path):
    page, xlsx = dealer_report.page_specialty(mini_wh, tmp_path, LAST_MONTH)
    k = {x.label: x for x in page.kpis}
    assert k["RV missing chassis VIN"].value == "1"  # the powersports row needs no chassis VIN
    assert k["RV missing chassis VIN"].status == "critical"
    monthly = pd.read_excel(xlsx, sheet_name="Monthly")
    assert len(monthly) == 2  # entered in the last complete month only
    assert len(pd.read_excel(xlsx, sheet_name="All time")) == 3


def test_remittance_t12(mini_wh, tmp_path):
    page, _ = dealer_report.page_remittance(mini_wh, tmp_path, LAST_MONTH)
    k = {x.label: x for x in page.kpis}
    assert k["Late fees, T12"].value == f"${2 * 12 * 40:,}"
    assert k["Late remittances, T12"].value == f"{(2 * 12 * 2) / (2 * 12 * 20):.1%}"


def test_scorecard_watch_and_portfolio_loss_ratio(mini_wh, tmp_path):
    page, xlsx = dealer_report.page_scorecard(mini_wh, tmp_path, LAST_MONTH)
    k = {x.label: x for x in page.kpis}
    assert k["Dealers on watch"].value == "1" and k["Dealers on watch"].status == "warning"
    assert k["Watch: loss ratio"].value == "1" and k["Watch: late remittance"].value == "0"
    assert k["Portfolio loss ratio on claims reserve, T12"].value == f"{(0.9 + 0.4) / 2:.1%}"  # summed, not averaged
    assert k["Watch: loss ratio"].sub == "T12, minimum $10K earned claims reserve"
    first = page.sections[0].rows[0]
    assert first[0].startswith("D1") and first[-1]["text"] == "Loss ratio"  # watch dealers sort first
    assert len(pd.read_excel(xlsx, sheet_name="Scorecard T12")) == 2


def test_claims_kpi_states_loss_month_basis_and_unmatched(mini_wh, tmp_path):
    mini_wh.execute("alter table claims.mart_claims_kpi add column claims_unmatched integer default 1")
    k = {x.label: x for x in claims_report.build(mini_wh, tmp_path / "site")[0].kpis}
    assert "by loss month" in k["Denial rate, T12"].sub
    assert "incl. 24 with no matched contract" in k["Denial rate, T12"].sub  # 2 subcats x 12 months


def test_claims_kpis_recomputed_from_sums(mini_wh, tmp_path):
    pages = claims_report.build(mini_wh, tmp_path / "site")
    assert [p.path for p in pages] == [claims_report.PAGE]
    k = {x.label: x for x in pages[0].kpis}
    assert k["Claim frequency, T12"].value == "0.100"  # 10 claims / 100 exposure years
    assert k["Claim severity, T12"].value == "$1,000"  # 8000 / 8
    assert k["Denial rate, T12"].value == f"{2 / 12:.1%}"
    assert k["Average turnaround, T12"].value == "8.0 days"
    assert (tmp_path / "site" / "downloads" / "claims" / "claims_kpi.xlsx").is_file()
    render(pages, tmp_path / "site")
    assert broken_links(tmp_path / "site") == []


def test_kpis_from_never_averages_ratios():
    df = pd.DataFrame({c: [0.0, 0.0] for c in claims_report.SUMS})
    df.loc[0, ["claim_count", "exposure_years"]] = [1, 1]  # frequency 1.0
    df.loc[1, ["claim_count", "exposure_years"]] = [1, 99]  # frequency ~0.01
    total = claims_report.kpis_from(df.sum().to_frame().T).iloc[0]
    assert total["claim_frequency"] == pytest.approx(2 / 100)
    assert pd.isna(total["claim_severity"])  # 0 paid claims -> n/a, not 0 or inf


def test_excel_writer_drops_metadata_and_tz(tmp_path):
    df = pd.DataFrame({"a": [1], "_built_at": [pd.Timestamp("2026-01-01", tz="UTC")],
                       "t": [pd.Timestamp("2026-01-01", tz="UTC")]})
    assert list(clean(df).columns) == ["a", "t"]
    path = write_workbook(tmp_path / "x" / "book.xlsx", {"Data": df}, notes=["n"])
    assert pd.ExcelFile(path).sheet_names == ["About", "Data"]


# --------------------------------------------------------------------------- against the built marts
def _marts_built() -> bool:
    if not WAREHOUSE.exists():
        return False
    try:
        with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
            n = con.execute("select count(*) from information_schema.tables where table_schema in ('dealer','claims') "
                            "and table_name in ('mart_dealer_production','mart_claims_kpi')").fetchone()[0]
        return n == 2
    except duckdb.Error:
        return False


needs_marts = pytest.mark.skipif(not _marts_built(), reason="dealer/claims marts not built (dbt build)")


@needs_marts
def test_real_marts_reconcile_to_core():
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        pend = con.execute("select count(*) from dealer.mart_dealer_pending_contracts").fetchone()[0]
        core_pend = con.execute("select count(*) from core.dim_contract where status = 'PENDING'").fetchone()[0]
        assert pend == core_pend
        claims, core_claims = con.execute(
            "select (select sum(claims_reported) from claims.mart_claims_kpi), (select count(*) from core.dim_claim)"
        ).fetchone()
        assert claims == core_claims
        paid, core_paid = con.execute(
            "select (select sum(paid_amount) from claims.mart_claim_summary), "
            "(select sum(paid_amount) from core.fct_claim_payment_line)").fetchone()
        assert paid == pytest.approx(core_paid, abs=1.0)
        demo = con.execute(
            "select count(*) from dealer.mart_dealer_production p join intermediate.int_reportable_scope s "
            "using (dealer_number) where not s.is_reportable and s.is_demo_dealer").fetchone()[0]
        assert demo == 0


@needs_marts
def test_real_plugins_build(tmp_path):
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        pages = dealer_report.build(con, tmp_path) + claims_report.build(con, tmp_path)
    render(pages, tmp_path)
    assert broken_links(tmp_path) == []
    assert len(pages) == 8


@needs_marts
def test_real_scorecard_ties_to_production_and_claims_basis():
    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        if not con.execute("select count(*) from information_schema.tables where table_schema = 'dealer' "
                           "and table_name = 'mart_dealer_scorecard'").fetchone()[0]:
            pytest.skip("mart_dealer_scorecard not built yet")
        dup = con.execute("select count(*) - count(distinct (tenant_key, dealer_key, month_key)) "
                          "from dealer.mart_dealer_scorecard").fetchone()[0]
        assert dup == 0
        card, prod = con.execute(
            "select (select sum(contracts_written) from dealer.mart_dealer_scorecard), "
            "(select sum(p.contracts_written) from dealer.mart_dealer_production p, core.dim_reporting_date r "
            " where p.dealer_key <> '-1' and p.month_start <= date_trunc('month', r.as_of_date))").fetchone()
        assert card == prod


@needs_marts
def test_real_early_claim_watch_agrees_with_actuarial():
    """Scorecard early-claim watch (as-of month) == actuarial early_claim_by_dealer flags, same window."""
    from actuarial.data import load_inputs
    from actuarial.emergence import early_claim_by_dealer

    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        if not con.execute("select count(*) from information_schema.tables where table_schema = 'dealer' "
                           "and table_name = 'mart_dealer_early_claim_ae'").fetchone()[0]:
            pytest.skip("mart_dealer_early_claim_ae not built yet")
        as_of = pd.Timestamp(con.execute("select as_of_date from core.dim_reporting_date").fetchone()[0]).date()
        ae = con.execute("select * from dealer.mart_dealer_early_claim_ae where reference_date = ?", [as_of]).df()
        card = con.execute("select dealer_number, is_watch_early_claims from dealer.mart_dealer_scorecard "
                           "where month_start = date_trunc('month', cast(? as date))", [as_of]).df()
        inputs = load_inputs(con)
    window = int(ae["window_days"].iloc[0])
    c = inputs.contracts[inputs.contracts["effective_date"] <= pd.Timestamp(as_of)].copy()
    for col in ("category_code", "subcategory_code", "new_used"):
        c[col] = c[col].fillna("UNKNOWN")
    py = early_claim_by_dealer(c, inputs.claims, as_of, windows=(window,))
    j = py.merge(ae, on="dealer_number", suffixes=("_py", "_mart"), how="outer", indicator=True)
    assert (j["_merge"] == "both").all()
    assert (j["contracts_exposed_py"] == j["contracts_exposed_mart"]).all()
    assert (j["early_claim_contracts_py"] == j["early_claim_contracts_mart"]).all()
    assert ((j["expected_early_claims_py"] - j["expected_early_claims_mart"]).abs() < 1e-6).all()
    flagged_py = set(py.loc[py["is_flagged"], "dealer_number"])
    assert flagged_py == set(ae.loc[ae["is_flagged"], "dealer_number"])
    assert flagged_py == set(card.loc[card["is_watch_early_claims"], "dealer_number"])


# --------------------------------------------------------------------------- one dealer watch on every page
def _add_tiny_dealer(con: duckdb.DuckDBPyConnection) -> None:
    """D3: worst loss ratio of all but a tiny book, below scorecard_min_t12_earned, so never on watch."""
    con.execute("""insert into dealer.mart_dealer_scorecard by name
        select 1 as tenant_key, 'D3' as dealer_key, 'D3' as dealer_number, 'Tiny Cars' as dealer_name,
            'TX' as dealer_state, month_start, month_key, 14 as t12_contracts_written, 0.0 as t12_cancel_rate,
            1933.17 as t12_earned_claims_reserve, 3313.8 as t12_incurred_claims, 1.714 as t12_loss_ratio,
            null as t12_loss_ratio_peer_pct, 0.0 as t12_early_claim_rate, 0.0 as t12_late_remit_rate,
            null as t12_avg_days_late, 0.0 as t12_late_fees, false as is_watch_loss_ratio,
            false as is_watch_late_remit, false as is_watch_early_claims, false as is_watch
        from dealer.mart_dealer_scorecard where dealer_key = 'D1'""")


def _overview_data(con: duckdb.DuckDBPyConnection):
    from types import SimpleNamespace

    from reports.data import load_dealer_watch

    dealers = pd.DataFrame({"dealer_code": ["D1", "D2", "D3"], "t12_loss_ratio": [0.9, 0.4, 1.714]})
    cur, rules, dealers = load_dealer_watch(con, LAST_MONTH.date(), dealers)
    return SimpleNamespace(dealer_watch=cur, dealer_watch_rules=rules, dealers=dealers)


def test_overview_watch_count_equals_scorecard_count(mini_wh, tmp_path):
    from reports.build import overview_dealer_watch_kpi

    _add_tiny_dealer(mini_wh)
    page, _ = dealer_report.page_scorecard(mini_wh, tmp_path, LAST_MONTH)
    card = {x.label: x for x in page.kpis}
    d = _overview_data(mini_wh)
    tile = overview_dealer_watch_kpi(d)
    assert tile.label == "Dealers above 80% LR"
    assert tile.value == card["Watch: loss ratio"].value == "1"  # D3's 171% on $1.9K is not counted
    assert tile.sub.startswith("T12, minimum $10K earned claims reserve")
    assert f"{card['Dealers on watch'].value} of 3 on any scorecard watch" in tile.sub
    # The Dealers page flag column reads the same mart flag.
    assert d.dealers.set_index("dealer_code")["is_watch_loss_ratio"].to_dict() == {"D1": True, "D2": False,
                                                                                  "D3": False}


def test_overview_watchlist_ranks_only_dealers_with_minimum_volume(mini_wh):
    from reports.build import overview_dealer_watchlist

    _add_tiny_dealer(mini_wh)
    s = overview_dealer_watchlist(_overview_data(mini_wh))
    names = [r[0]["text"] if isinstance(r[0], dict) else r[0] for r in s.rows]
    assert names == ["D1 Summit Motors", "D2 Harbor Auto"]  # D3 has the worst ratio but too little volume
    assert s.rows[0][-1]["text"] == "Above 80%" and s.rows[1][-1] == ""
    assert 'href="dealer/scorecard.html"' in str(s.note)
    assert "earned premium" not in s.caption.lower()


def test_dealer_flags_never_use_a_python_threshold():
    src = (REPORTS_DIR / "build.py").read_text(encoding="utf-8") + (REPORTS_DIR / "data.py").read_text("utf-8")
    assert "is_high_loss_ratio" not in src
    assert "mart_dealer_loss_ratio" not in src.split("def load_dealer_watch")[1].split("def load(")[0]


def test_scorecard_copy_says_claims_reserve_not_earned_premium(mini_wh, tmp_path):
    """One label, one meaning: the scorecard's denominator is the earned claims reserve."""
    sql = (DEALER_MODELS / "mart_dealer_scorecard.sql").read_text(encoding="utf-8")
    final = sql.rsplit("from flagged", 1)[0].rsplit("\nselect", 1)[1]
    assert "earned_premium" not in final and "t12_earned_claims_reserve" in final
    pages = dealer_report.build(mini_wh, tmp_path)
    render(pages, tmp_path)
    html = (tmp_path / "dealer" / "scorecard.html").read_text(encoding="utf-8").lower()
    assert "earned premium" not in html
    assert dealer_watch.LR_LABEL.lower() in html


@needs_marts
def test_real_overview_watch_equals_scorecard(tmp_path):
    from reports.build import overview_dealer_watch_kpi, overview_dealer_watchlist
    from reports.data import load_dealer_watch

    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        _, last = dealer_report.as_of(con)
        try:
            page, _ = dealer_report.page_scorecard(con, tmp_path, last)
        except duckdb.Error:
            pytest.skip("mart_dealer_scorecard not built yet")
        codes = con.execute("select dealer_code from gold.dim_dealer where is_current and dealer_key <> '-1'").df()
        cur, rules, dealers = load_dealer_watch(con, last.date(), codes)
    from types import SimpleNamespace

    d = SimpleNamespace(dealer_watch=cur, dealer_watch_rules=rules, dealers=dealers)
    card = {x.label: x.value for x in page.kpis}
    assert overview_dealer_watch_kpi(d).value == card["Watch: loss ratio"]
    assert str(int(dealers["is_watch_loss_ratio"].sum())) == card["Watch: loss ratio"]
    listed = {(r[0]["text"] if isinstance(r[0], dict) else r[0]).split()[0] for r in overview_dealer_watchlist(d).rows}
    small = set(cur.loc[cur["t12_earned_claims_reserve"] < rules["scorecard_min_t12_earned"], "dealer_number"])
    assert listed and not listed & small


# --------------------------------------------------------------------------- one dealer deliverable, one label
CANCEL_LABELS = (dealer_watch.CANCEL_RATE_ITD, dealer_watch.CANCEL_RATE_T12, dealer_watch.CANCEL_RATE_MONTH)


def test_every_cancel_rate_label_is_qualified():
    """No page says just 'Cancel rate': each label names its basis (sale cohort) and window."""
    from publish import packets

    assert len(set(CANCEL_LABELS)) == 3 and all("by sale cohort" in lbl for lbl in CANCEL_LABELS)
    labels = {c.label for cols in (packets.KPI_COLS, packets.PROGRAM_DEALER_COLS,
                                   *[spec.cols for spec in packets.TABLE_SPECS.values()])
              for c in cols if c.key == "cancel_rate"}
    assert labels == {dealer_watch.CANCEL_RATE_T12, dealer_watch.CANCEL_RATE_MONTH}
    assert {c.label for c in packets.KPI_COLS if c.key == "cancel_rate"} == {dealer_watch.CANCEL_RATE_T12}
    for src in (REPORTS_DIR / "build.py", REPORTS_DIR / "dealer" / "__init__.py", ROOT / "publish" / "packets.py"):
        text = src.read_text(encoding="utf-8")
        assert not re.search(r'"(Cancel rate|Cancellation rate)"', text), src


def test_v1_flat_dealer_pages_are_retired():
    """Each dealer has one deliverable (the tenant packet); the shell builds no site/dealers/<dealer>.html."""
    import reports.build as shell

    assert not hasattr(shell, "page_dealer")
    src = (REPORTS_DIR / "build.py").read_text(encoding="utf-8")
    assert 'f"dealers/{r.dealer_code}.html"' not in src and 'f"dealers/{r.dealer_number}.html"' not in src
    assert "mart_cancellations where dealer_code" not in (REPORTS_DIR / "data.py").read_text(encoding="utf-8")


def _dealers_page_data(con, codes=("D1", "D2")):
    from types import SimpleNamespace

    from reports.data import load_dealer_watch

    dealers = pd.DataFrame({"dealer_code": list(codes), "dealer_name": [f"Dealer {c}" for c in codes],
                            "city": "Austin", "state": "TX", "dealer_status": "ACTIVE", "t12_loss_ratio": 0.5,
                            "itd_loss_ratio": 0.5, "t12_claim_frequency": 0.1, "contracts_in_force": 10})
    cur, rules, dealers = load_dealer_watch(con, LAST_MONTH.date(), dealers)
    return SimpleNamespace(dealer_watch=cur, dealer_watch_rules=rules, dealers=dealers)


def test_overview_and_dealers_page_link_to_the_packet_and_show_scorecard_cancel_rate(mini_wh, tmp_path):
    from reports.build import overview_dealer_watchlist, page_dealers

    mini_wh.execute("update dealer.mart_dealer_scorecard set t12_cancel_rate = 0.25 where dealer_key = 'D2'")
    d = _dealers_page_data(mini_wh)
    watch = overview_dealer_watchlist(d)
    assert {r[0]["href"] for r in watch.rows} == {"dealers/T1/D1/index.html", "dealers/T1/D2/index.html"}
    page = page_dealers(d)
    table = page.sections[1]
    i = table.cols.index(dealer_watch.CANCEL_RATE_T12)
    shown = {r[0]["href"]: r[i] for r in table.rows}
    assert shown == {"dealers/T1/D1/index.html": "10.0%", "dealers/T1/D2/index.html": "25.0%"}
    card, _ = dealer_report.page_scorecard(mini_wh, tmp_path, LAST_MONTH)
    s1 = card.sections[0]
    j = s1.cols.index(dealer_watch.CANCEL_RATE_T12)
    assert {r[0].split()[0]: r[j] for r in s1.rows} == {"D1": "10.0%", "D2": "25.0%"}


def test_dealer_without_a_packet_gets_no_link(mini_wh):
    from reports.build import page_dealers

    mini_wh.execute("delete from core.dim_user_dealer_access where dealer_key = 'D2'")
    rows = page_dealers(_dealers_page_data(mini_wh)).sections[1].rows
    assert isinstance(rows[0][0], dict) and rows[1][0] == "D2 Dealer D2"


@needs_marts
def test_real_same_dealer_same_label_same_value(tmp_path):
    """For every dealer, 'Cancel rate, by sale cohort, T12' shows one number on the Dealers page, the dealer
    scorecard and the dealer's packet; and the overview's portfolio tile carries the other (ITD) label."""
    from publish import packets as pk
    from reports.build import page_dealers
    from reports.data import load_dealer_monthly, load_dealer_watch, load_dealers

    with duckdb.connect(str(WAREHOUSE), read_only=True) as con:
        try:
            _, last = dealer_report.as_of(con)
            card, _ = dealer_report.page_scorecard(con, tmp_path, last)
            dealers = load_dealers(con, load_dealer_monthly(con, last.date()),
                                   (last - pd.DateOffset(months=11)).date(), last.date())
            cur, rules, dealers = load_dealer_watch(con, last.date(), dealers)
            per = pk.periods(con)
            packet_rows = pk.packets(con)
            kpi_sql = pk.TABLE_SPECS["kpi"].sql
            packet = {}
            for p in packet_rows:
                prm = {"dealer_key": p.dealer_key, "tenant_key": p.tenant_key, **per.params()}
                df = con.execute(kpi_sql, {k: v for k, v in prm.items() if f"${k}" in kpi_sql}).df()
                if len(df):
                    packet[f"dealers/{p.tenant_code}/{p.dealer_number}/index.html"] = pk.fmt_value(
                        df["cancel_rate"].iloc[0], "pct")
        except duckdb.Error as exc:
            pytest.skip(f"dealer scorecard / packet inputs not built: {exc}")
    from types import SimpleNamespace

    assert per.t12_end == last.date()  # the packet's T12 ends where the scorecard's does
    d = SimpleNamespace(dealer_watch=cur, dealer_watch_rules=rules, dealers=dealers)
    table = page_dealers(d).sections[1]
    i = table.cols.index(dealer_watch.CANCEL_RATE_T12)
    site = {r[0]["href"]: r[i] for r in table.rows if isinstance(r[0], dict)}
    assert site and set(site) <= set(packet)
    mismatched = {k: (v, packet[k].replace("-", "n/a")) for k, v in site.items() if v != packet[k].replace("-", "n/a")}
    assert not mismatched
    s1 = card.sections[0]
    j = s1.cols.index(dealer_watch.CANCEL_RATE_T12)
    by_dealer = {k.split("/")[2]: v for k, v in site.items()}
    assert all(by_dealer[r[0].split()[0]] == r[j] for r in s1.rows)
    links = dealer_watch.packet_links(cur)
    assert links.get("D1001", "").endswith("/D1001/index.html")


# --------------------------------------------------------------------------- V6 edits 10, 18, status split
def _proper_case_sql(expr: str, exceptions: str) -> str:
    """Render the dbt macro proper_case (macros/text.sql) outside dbt."""
    text = (ROOT / "dbt" / "macros" / "text.sql").read_text(encoding="utf-8")
    return str(Environment().from_string(text).module.proper_case(expr, exceptions))


def test_proper_case_exceptions_seed_is_lower_case_and_unique():
    rows = seed("proper_case_exceptions")
    tokens = [r["token"] for r in rows]
    assert tokens == [t.lower() for t in tokens] and len(set(tokens)) == len(tokens)
    assert all(r["replacement"].lower() == r["token"] for r in rows), "a replacement only changes the case"
    assert {"llc", "ii", "iii", "jr"} <= set(tokens)


def test_proper_case_macro_names_and_addresses():
    con = duckdb.connect()
    rows = seed("proper_case_exceptions")
    con.execute("create table ex as select map(?, ?) as exceptions",
                [[r["token"] for r in rows], [r["replacement"] for r in rows]])
    sql = _proper_case_sql("t.x", "ex.exceptions")
    cases = {
        "JOHN SMITH-JONES II": "John Smith-Jones II",
        "o'neil motors llc": "O'Neil Motors LLC",
        "  ann   lee  jr ": "Ann Lee Jr",
        "bank of the north na": "Bank of The North NA",
        "of counsel lending": "Of Counsel Lending",
        "123 nw main st": "123 NW Main St",
        "": None,
        None: None,
    }
    con.execute("create table t (x varchar)")
    con.executemany("insert into t values (?)", [[k] for k in cases])
    got = dict(con.execute(f"select t.x, {sql} from t cross join ex").fetchall())
    assert got == cases


def test_pending_mart_formats_packet_columns():
    sql = (DEALER_MODELS / "mart_dealer_pending_contracts.sql").read_text(encoding="utf-8")
    for col in ("customer_first_name", "customer_last_name", "cobuyer_name", "cobuyer_address_line1",
                "cobuyer_city", "lienholder_name", "provider_name"):
        assert re.search(r"proper_case\('[a-z]+\.[a-z0-9_]+'\) }} as " + col, sql), col
    assert "cb.deal_id = p.deal_id" in sql and "administrator_name" in sql
    assert set(dealer_report.PENDING_PACKET_COLUMNS) >= {
        "cobuyer_name", "cobuyer_address_line1", "cobuyer_city", "cobuyer_state", "cobuyer_zip",
        "lienholder_name", "provider_name", "administrator_name", "customer_state"}


def test_pending_workbook_has_fixed_columns_filter_and_frozen_header(mini_wh, tmp_path):
    from openpyxl import load_workbook

    _, xlsx = dealer_report.page_pending(mini_wh, tmp_path)
    detail = pd.read_excel(xlsx, sheet_name="Pending contracts")
    assert list(detail.columns) == dealer_report.PENDING_PACKET_COLUMNS  # missing columns come out empty
    ws = load_workbook(xlsx)["Pending contracts"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == ws.dimensions


def test_dealer_packet_xlsx_filters_every_column(tmp_path):
    from datetime import date

    from openpyxl import load_workbook

    import publish.packets as pk

    spec = pk.TABLE_SPECS["pending_contracts"]
    assert {c.key for c in spec.cols} >= {"provider_name", "administrator_name"}
    per = pk.Periods(date(2026, 10, 21), date(2025, 10, 1), date(2026, 9, 1))
    df = pd.DataFrame([{c.key: None for c in spec.cols} | {"contract_number": "VSC-1", "days_pending": 3}])
    path = tmp_path / "p.xlsx"
    assert pk._xlsx(path, [("Pending", "Pending contracts", "as of", df, spec.cols)], per)
    ws = load_workbook(path)["Pending"]
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref == f"A4:{ws.cell(row=4, column=len(spec.cols)).column_letter}5"


def test_production_split_comes_from_the_status_seed():
    rows = seed("dealer_report_status_sets")
    written = {r["status"] for r in rows if r["report_id"] == "DEALER_PRODUCTION"}
    buckets = {b: {r["status"] for r in rows if r["report_id"] == b}
               for b in ("DEALER_PROD_PENDING", "DEALER_PROD_IN_FORCE", "DEALER_PROD_CANCELLED")}
    assert all(buckets.values())
    flat = [s for v in buckets.values() for s in v]
    assert len(flat) == len(set(flat)), "a status counts in at most one bucket"
    assert set(flat) <= written
    assert {"SUSPENDED"} <= buckets["DEALER_PROD_IN_FORCE"]
    assert {"REINSTATEMENT_PENDING", "REINSTATEMENT_SUSPENDED"} <= buckets["DEALER_PROD_CANCELLED"]
    sql = (DEALER_MODELS / "mart_dealer_production.sql").read_text(encoding="utf-8")
    assert not re.search(r"status = '[A-Z_]+'", sql), "no hard-coded status in the production split"


def test_claim_components_aggregate_the_row_level_export():
    sql = (CLAIMS_MODELS / "mart_claim_components.sql").read_text(encoding="utf-8")
    assert set(re.findall(r"ref\('([^']+)'\)", sql)) == {"mart_claim_component_lines"}
    assert (ROOT / "dbt" / "tests" / "claims" / "assert_claim_exports_tie_to_summary.sql").is_file()


def _add_export_marts(con: duckdb.DuckDBPyConnection) -> None:
    """Give the fixture the V6 export columns: new marts get one row, the summary gets its new columns."""
    for _, _, mart, cols, _ in claims_report.EXPORTS:
        have = {r[0] for r in con.execute(f"describe {mart}").fetchall()} if mart.endswith("summary") else None
        if have is not None:
            for c in cols:
                if c not in have:
                    con.execute(f"alter table {mart} add column {c} varchar")
            con.execute(f"update {mart} set claim_number = 'CL-' || claim_key")
        else:
            sel = ", ".join(f"'CL-1' as {c}" if c == "claim_number" else f"1 as {c}" for c in cols)
            con.execute(f"create table {mart} as select {sel}, 'x' as component_line_id, '-1' as _built_at")


def test_claims_row_level_exports_fixed_columns(mini_wh, tmp_path):
    out = tmp_path / "site"
    pages = claims_report.build(mini_wh, out)  # the fixture's pre-V6 summary has no claim_number: no exports
    assert not any(s.title == "Row-level exports" for s in pages[0].sections)
    _add_export_marts(mini_wh)
    pages = claims_report.build(mini_wh, out)
    sec = [s for s in pages[0].sections if s.title == "Row-level exports"]
    assert len(sec) == 1 and len(sec[0].rows) == len(claims_report.EXPORTS)
    for stem, sheet, _, cols, _ in claims_report.EXPORTS:
        path = out / "downloads" / "claims" / f"{stem}.xlsx"
        book = pd.ExcelFile(path)
        assert book.sheet_names == ["About", sheet]
        assert list(pd.read_excel(path, sheet_name=sheet).columns) == cols
    render(pages, out)
    assert broken_links(out) == []


def test_claims_exports_skip_quietly_without_their_marts(tmp_path):
    con = duckdb.connect()
    con.execute("create schema claims")
    assert claims_report.write_exports(con, tmp_path, "Sep 2026") == []
