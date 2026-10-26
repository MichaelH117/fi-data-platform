"""Per-dealer and per-program publish packets (docs/V3_DESIGN.md section 3, per-dealer publish rule).

Dealer packets: one per (tenant, reportable dealer) in ``core.dim_user_dealer_access`` - the same
user -> dealer, tenant map the semantic RLS table uses - written to
``<packets_dir>/<tenant_code>/<dealer_number>/``. Program packets: one per tenant (white-label partner
program), written to ``<packets_dir>/<tenant_code>/`` (``index.html`` + ``program.xlsx``), listing that
program's dealers with links to their packets.

Every table reads a certified mart, never a fact, so a packet KPI is the same number the reports show
("one certified source per KPI"):

    kpi                 headline strip: T12 written, cancel rate, loss ratio and its peer percentile,
                        in force, pending, late fees        dealer.mart_dealer_scorecard (+ production,
                                                            pending marts) + peer_benchmark_t12
    production          month by month, T12                 dealer.mart_dealer_scorecard
    pending_contracts   PENDING contracts, no customer data dealer.mart_dealer_pending_contracts
    remittance_aging    shared seed bands (remit_aging_buckets), latest closed period and T12
                                                            dealer.mart_remittance_aging
    claims_kpi          T12 claims, early claims, loss ratio dealer.mart_dealer_scorecard
    program_claims      program packet only, by product line claims.mart_claims_kpi (tenant grain;
                                                            it has no dealer grain, so dealer packets
                                                            take claims from the scorecard)

Every query is filtered by BOTH dealer_key and tenant_key (program queries by tenant_key), so a
packet can only ever hold its own dealer's rows. CSVs keep machine column names; HTML and xlsx use
readable headers, money / count / percent formats and a period label on every table. No hand exports.
"""

from __future__ import annotations

import html
import json
import logging
import math
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import duckdb
import pandas as pd

from pipeline import workpool

log = logging.getLogger("publish.packets")

_DEALER = "dealer_key = $dealer_key and tenant_key = $tenant_key"
_T12 = "month_start between $t12_start and $t12_end"


@dataclass(frozen=True)
class Col:
    key: str
    label: str
    fmt: str = "text"  # text | int | money | pct | ratio | date


@dataclass(frozen=True)
class Spec:
    title: str
    period: str  # t12 | asof | aging | program_t12
    sql: str
    cols: tuple[Col, ...]


KPI_COLS: tuple[Col, ...] = (
    Col("contracts_written", "Contracts written", "int"),
    Col("cancel_rate", "Cancel rate, by sale cohort, T12", "pct"),  # = reports.dealer.watch.CANCEL_RATE_T12
    Col("loss_ratio", "Loss ratio", "pct"),
    Col("peer_loss_ratio", "Peer loss ratio (VSC)", "pct"),
    Col("loss_ratio_peer_pct", "Loss-ratio peer percentile", "pct"),
    Col("contracts_in_force", "In force", "int"),
    Col("pending_contracts", "Pending contracts", "int"),
    Col("late_fees", "Late fees", "money"),
)

TABLE_SPECS: dict[str, Spec] = {
    "kpi": Spec(
        "Headline KPIs", "t12",
        f"""
        with s as (select * from dealer.mart_dealer_scorecard where {_DEALER} and {_T12})
        select cast(sum(contracts_written) as bigint) as contracts_written,
               cast(sum(contracts_cancelled) as bigint) as contracts_cancelled,
               sum(contracts_cancelled) / nullif(sum(contracts_written), 0) as cancel_rate,
               round(sum(earned_claims_reserve), 2) as earned_claims_reserve,
               round(sum(incurred_claims), 2) as incurred_claims,
               sum(incurred_claims) / nullif(sum(earned_claims_reserve), 0) as loss_ratio,
               (select max(t12_loss_ratio_peer_pct) from dealer.mart_dealer_scorecard
                 where {_DEALER} and month_start = $t12_end) as loss_ratio_peer_pct,
               (select cast(coalesce(sum(contracts_in_force), 0) as bigint) from dealer.mart_dealer_production
                 where {_DEALER}) as contracts_in_force,
               (select count(*) from dealer.mart_dealer_pending_contracts where {_DEALER}) as pending_contracts,
               cast(sum(remittances) as bigint) as remittances,
               round(sum(late_fees), 2) as late_fees,
               cast(sum(claims_reported) as bigint) as claims_reported
        from s
        having count(*) > 0
        """,
        KPI_COLS,
    ),
    "production": Spec(
        "Production", "t12",
        f"""
        select strftime(month_start, '%b %Y') as month,
               cast(contracts_written as bigint) as contracts_written,
               cast(contracts_cancelled as bigint) as contracts_cancelled,
               cancel_rate, round(coverage_cost, 2) as coverage_cost, round(retail_price, 2) as retail_price,
               round(earned_claims_reserve, 2) as earned_claims_reserve,
               round(incurred_claims, 2) as incurred_claims, loss_ratio
        from dealer.mart_dealer_scorecard
        where {_DEALER} and {_T12}
        order by month_start
        """,
        (Col("month", "Month"), Col("contracts_written", "Contracts written", "int"),
         Col("contracts_cancelled", "Contracts cancelled", "int"),
         Col("cancel_rate", "Cancel rate, by sale cohort, month", "pct"),
         Col("coverage_cost", "Coverage cost", "money"), Col("retail_price", "Retail price", "money"),
         Col("earned_claims_reserve", "Earned claims reserve", "money"),
         Col("incurred_claims", "Incurred claims", "money"),
         Col("loss_ratio", "Loss ratio", "pct")),
    ),
    "pending_contracts": Spec(
        "Pending contracts", "asof",
        f"""
        select contract_number, product, plan, date_sold, cast(days_pending as bigint) as days_pending,
               pending_age_band, round(coverage_cost, 2) as coverage_cost, round(retail_price, 2) as retail_price,
               provider_name, administrator_name
        from dealer.mart_dealer_pending_contracts
        where {_DEALER}
        order by days_pending desc, contract_number
        """,
        (Col("contract_number", "Contract"), Col("product", "Product"), Col("plan", "Plan"),
         Col("date_sold", "Date sold", "date"), Col("days_pending", "Days pending", "int"),
         Col("pending_age_band", "Age band"), Col("coverage_cost", "Coverage cost", "money"),
         Col("retail_price", "Retail price", "money"), Col("provider_name", "Provider"),
         Col("administrator_name", "Administrator")),
    ),
    "remittance_aging": Spec(
        "Remittance aging", "aging",
        f"""
        select age_band_order, age_band,
               cast(coalesce(sum(cash_transactions) filter (where posting_period_key = $p_end), 0) as bigint)
                   as period_remittances,
               round(coalesce(sum(remitted_amount) filter (where posting_period_key = $p_end), 0), 2)
                   as period_amount,
               round(coalesce(sum(late_fees) filter (where posting_period_key = $p_end), 0), 2) as period_late_fees,
               cast(sum(cash_transactions) as bigint) as t12_remittances,
               cast(sum(late_transactions) as bigint) as t12_late_remittances,
               round(sum(remitted_amount), 2) as t12_amount,
               round(sum(late_fees), 2) as t12_late_fees
        from dealer.mart_remittance_aging
        where {_DEALER} and posting_period_key between $p_start and $p_end
        group by all
        order by age_band_order
        """,
        (Col("age_band", "Days to remit"), Col("period_remittances", "{p} remittances", "int"),
         Col("period_amount", "{p} amount", "money"), Col("period_late_fees", "{p} late fees", "money"),
         Col("t12_remittances", "T12 remittances", "int"), Col("t12_late_remittances", "T12 late", "int"),
         Col("t12_amount", "T12 amount", "money"), Col("t12_late_fees", "T12 late fees", "money")),
    ),
    "claims_kpi": Spec(
        "Claims KPIs", "t12",
        f"""
        select cast(sum(claims_reported) as bigint) as claims_reported,
               cast(sum(early_claims) as bigint) as early_claims,
               sum(early_claims) / nullif(sum(claims_reported), 0) as early_claim_rate,
               round(sum(incurred_claims), 2) as incurred_claims,
               round(sum(earned_claims_reserve), 2) as earned_claims_reserve,
               sum(incurred_claims) / nullif(sum(earned_claims_reserve), 0) as loss_ratio
        from dealer.mart_dealer_scorecard
        where {_DEALER} and {_T12}
        having count(*) > 0
        """,
        (Col("claims_reported", "Claims reported", "int"), Col("early_claims", "Early claims", "int"),
         Col("early_claim_rate", "Early-claim rate", "pct"), Col("incurred_claims", "Incurred claims", "money"),
         Col("earned_claims_reserve", "Earned claims reserve", "money"), Col("loss_ratio", "Loss ratio", "pct")),
    ),
    # V7 edit 25: this month's pending vs active, with the three months before it.
    "activations": Spec(
        "Current month activations", "asof",
        f"""
        select strftime(month_start, '%b %Y') || case when is_current_month then ' (in progress)' else '' end
                   as month,
               product_family,
               cast(sum(contracts_pending) as bigint) as contracts_pending,
               cast(sum(contracts_active) as bigint) as contracts_active
        from dealer.mart_dealer_activations
        where {_DEALER}
        group by month_start, is_current_month, product_family
        order by month_start desc, product_family
        """,
        (Col("month", "Month"), Col("product_family", "Product family"),
         Col("contracts_pending", "Still pending", "int"), Col("contracts_active", "Active", "int")),
    ),
    # V7 edit 46: what the store is owed for each remittance, line by line (total = sum of the sections).
    "incentive_statement": Spec(
        "Incentive statement by remittance", "t12",
        f"""
        select remit_batch_id, remit_received_date, product_type,
               round(sum(commissions), 2) as commissions, round(sum(cancels), 2) as cancels,
               round(sum(reinstatements), 2) as reinstatements, round(sum(amendments), 2) as amendments,
               round(sum(other_adjustments), 2) as other_adjustments, round(sum(total_due), 2) as total_due
        from dealer.mart_dealer_incentive_statement
        where {_DEALER} and statement_status = 'REMITTED'
          and remit_received_date between $t12_start and last_day($t12_end)
        group by all
        order by remit_received_date desc, remit_batch_id, product_type
        """,
        (Col("remit_batch_id", "Remittance"), Col("remit_received_date", "Received", "date"),
         Col("product_type", "Product type"), Col("commissions", "New-sale commission", "money"),
         Col("cancels", "Cancels", "money"), Col("reinstatements", "Reinstatements", "money"),
         Col("amendments", "Amendments", "money"), Col("other_adjustments", "Other adjustments", "money"),
         Col("total_due", "Total due", "money")),
    ),
}

# Sections added after a mart existed (V7): an older warehouse without the mart gets an empty table
# instead of a failed packet run.
OPTIONAL_TABLES = {"activations", "incentive_statement"}

PROGRAM_CLAIMS = Spec(
    "Program claims KPIs by product line", "program_t12",
    f"""
    select coalesce(subcategory_code, 'All products') as product_line,
           cast(sum(claims_reported) as bigint) as claims_reported,
           cast(sum(claims_denied) as bigint) as claims_denied,
           sum(claims_denied) / nullif(sum(claims_reported), 0) as denial_rate,
           cast(sum(paid_claim_count) as bigint) as paid_claims,
           round(sum(paid_amount), 2) as paid_amount,
           sum(paid_amount_on_paid_claims) / nullif(sum(paid_claim_count), 0) as claim_severity,
           sum(claim_count) / nullif(sum(exposure_years), 0) as claim_frequency
    from claims.mart_claims_kpi
    where tenant_key = $tenant_key and {_T12}
    group by grouping sets ((subcategory_code), ())
    order by grouping(subcategory_code), 1
    """,
    (Col("product_line", "Product line"), Col("claims_reported", "Claims reported", "int"),
     Col("claims_denied", "Denied", "int"), Col("denial_rate", "Denial rate", "pct"),
     Col("paid_claims", "Paid claims", "int"), Col("paid_amount", "Paid", "money"),
     Col("claim_severity", "Severity", "money"), Col("claim_frequency", "Frequency (per exposure year)", "ratio")),
)

# Back-compat name: the table list the packet is made of.
TABLES: dict[str, str] = {name: s.sql for name, s in TABLE_SPECS.items()}

PROGRAM_KPI_COLS: tuple[Col, ...] = (
    Col("dealers", "Dealers", "int"),
    Col("contracts_written", "Contracts written", "int"),
    Col("contracts_in_force", "In force", "int"),
    Col("earned_claims_reserve", "Earned claims reserve", "money"),
    Col("incurred_claims", "Incurred claims", "money"),
    Col("loss_ratio", "Loss ratio", "pct"),
    Col("late_fees", "Late fees", "money"),
    Col("pending_contracts", "Pending contracts", "int"),
)
PROGRAM_SUM_KEYS = ("contracts_written", "contracts_cancelled", "contracts_in_force", "earned_claims_reserve",
                    "incurred_claims", "late_fees", "pending_contracts", "remittances", "claims_reported")
PROGRAM_DEALER_COLS: tuple[Col, ...] = (
    Col("dealer_number", "Dealer"), Col("dealer_name", "Name"), Col("contracts_written", "Written", "int"),
    Col("cancel_rate", "Cancel rate, by sale cohort, T12", "pct"), Col("loss_ratio", "Loss ratio", "pct"),
    Col("loss_ratio_peer_pct", "Peer percentile", "pct"), Col("contracts_in_force", "In force", "int"),
    Col("pending_contracts", "Pending", "int"), Col("late_fees", "Late fees", "money"),
)


@dataclass(frozen=True)
class Packet:
    tenant_key: int
    tenant_code: str
    reportable_name: str
    primary_color: str
    secondary_color: str
    dealer_key: str
    dealer_number: str
    dealer_name: str


@dataclass(frozen=True)
class Periods:
    as_of: date
    t12_start: date  # first month of the trailing 12 complete months
    t12_end: date  # last complete month

    @property
    def p_start(self) -> int:
        return self.t12_start.year * 100 + self.t12_start.month

    @property
    def p_end(self) -> int:
        return self.t12_end.year * 100 + self.t12_end.month

    @property
    def month_label(self) -> str:
        return f"{self.t12_end:%b %Y}"

    @property
    def t12_label(self) -> str:
        return f"T12 to {self.t12_end:%b %Y}"

    def label(self, kind: str) -> str:
        return {
            "t12": f"{self.t12_label} ({self.t12_start:%b %Y} - {self.t12_end:%b %Y})",
            "program_t12": f"{self.t12_label}, all dealers in the program",
            "asof": f"As of {self.as_of:%Y-%m-%d}",
            "aging": f"{self.month_label} posting period and {self.t12_label}; bands from the shared aging seed",
        }[kind]

    def params(self) -> dict[str, object]:
        return {"t12_start": self.t12_start, "t12_end": self.t12_end, "p_start": self.p_start, "p_end": self.p_end}


def periods(con: duckdb.DuckDBPyConnection) -> Periods:
    as_of, last = con.execute("select as_of_date, last_complete_month from core.dim_reporting_date").fetchone()
    last = pd.Timestamp(last).date()
    start = (pd.Timestamp(last) - pd.DateOffset(months=11)).date()
    return Periods(pd.Timestamp(as_of).date(), start, last)


def packets(con: duckdb.DuckDBPyConnection) -> list[Packet]:
    rows = con.execute(
        """
        select distinct a.tenant_key, t.tenant_code, t.reportable_name, coalesce(t.primary_color, '#1f3a5f'),
               coalesce(t.secondary_color, '#e8eef5'), a.dealer_key, a.dealer_number, d.dealer_name
        from core.dim_user_dealer_access a
        join core.dim_tenant t using (tenant_key)
        join core.dim_dealer d using (dealer_key)
        order by 2, 7
        """
    ).fetchall()
    return [Packet(*r) for r in rows]


def _dealer_state(con: duckdb.DuckDBPyConnection, dealer_key: str) -> str | None:
    row = con.execute("select state from core.dim_dealer where dealer_key = ?", [dealer_key]).fetchone()
    return row[0] if row else None


def peer_loss_ratio(peer: pd.DataFrame | None, state: str | None) -> tuple[float | None, str]:
    """T12 peer loss ratio from the suppressed benchmark: the dealer's state cell, else all states."""
    if peer is None or peer.empty:
        return None, ""
    cells = peer[(peer["coverage_tier"] == "ALL") & (peer["term_band"] == "ALL") & ~peer["suppressed"].astype(bool)]
    for scope, label in ((state, f"{state} peers"), ("ALL", "all-state peers")):
        hit = cells[cells["state"] == scope]
        if scope and len(hit) and pd.notna(hit["loss_ratio"].iloc[0]):
            return float(hit["loss_ratio"].iloc[0]), label
    return None, ""


def packet_tables(con: duckdb.DuckDBPyConnection, p: Packet, per: Periods | None = None,
                  peer: pd.DataFrame | None = None) -> dict[str, pd.DataFrame]:
    per = per or periods(con)
    params = {"dealer_key": p.dealer_key, "tenant_key": p.tenant_key, **per.params()}
    out: dict[str, pd.DataFrame] = {}
    for name, spec in TABLE_SPECS.items():
        sql_params = {k: v for k, v in params.items() if f"${k}" in spec.sql}
        try:
            out[name] = con.execute(spec.sql, sql_params).df()
        except duckdb.CatalogException:
            if name not in OPTIONAL_TABLES:
                raise
            out[name] = pd.DataFrame(columns=[c.key for c in spec.cols])
    kpi = out["kpi"]
    if len(kpi):
        value, scope = peer_loss_ratio(peer, _dealer_state(con, p.dealer_key))
        kpi["peer_loss_ratio"] = value
        kpi["peer_scope"] = scope
    return out


# --------------------------------------------------------------------------- formatting
def _missing(v: object) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or v is pd.NA or v is pd.NaT


def _nonzero(body: str) -> bool:
    return any(ch in "123456789" for ch in body)


def _no_negative_zero(text: str) -> str:
    """'-0.0%' -> '0.0%': a tiny negative that rounds to zero is shown as zero."""
    return text[1:] if text.startswith("-") and not _nonzero(text) else text


def fmt_value(v: object, fmt: str) -> str:
    if _missing(v):
        return "-"
    if fmt == "int":
        return f"{int(v):,}"
    if fmt == "money":  # sign before the symbol (-$5.00, never $-5.00); a value that rounds to zero has no sign
        body = f"{abs(float(v)):,.2f}"
        return f"{'-' if float(v) < 0 and _nonzero(body) else ''}${body}"
    if fmt == "pct":
        return _no_negative_zero(f"{float(v):.1%}")
    if fmt == "ratio":
        return _no_negative_zero(f"{float(v):.4f}")
    if fmt == "date":
        return f"{pd.Timestamp(v):%Y-%m-%d}"
    return str(v)


XLSX_FORMATS = {"int": "#,##0", "money": '"$"#,##0.00', "pct": "0.0%", "ratio": "0.0000", "date": "yyyy-mm-dd"}
XLSX_DIGITS = {"money": 2, "pct": 3, "ratio": 4}  # decimals each format shows (pct: 0.0% = 3 on the fraction)


def _labels(cols: tuple[Col, ...], per: Periods) -> list[str]:
    return [c.label.format(p=per.month_label) for c in cols]


def _table_html(df: pd.DataFrame, cols: tuple[Col, ...], per: Periods, links: dict[str, str] | None = None) -> str:
    if not len(df):
        return "<p class='empty'>No rows for this period.</p>"
    head = "".join(f"<th class='{'num' if c.fmt != 'text' and c.fmt != 'date' else ''}'>{html.escape(lab)}</th>"
                   for c, lab in zip(cols, _labels(cols, per), strict=True))
    body = []
    for rec in df.to_dict("records"):
        cells = []
        for c in cols:
            text = html.escape(fmt_value(rec.get(c.key), c.fmt))
            if links and c.key == "dealer_number" and rec.get(c.key) in links:
                text = f"<a href='{html.escape(links[rec[c.key]])}'>{text}</a>"
            cells.append(f"<td class='{'num' if c.fmt not in ('text', 'date') else ''}'>{text}</td>")
        body.append("<tr>" + "".join(cells) + "</tr>")
    return (f"<div class='scroll'><table class='t'><thead><tr>{head}</tr></thead>"
            f"<tbody>{''.join(body)}</tbody></table></div>")


# --------------------------------------------------------------------------- look and explanations (V7)
# Packets use the F&IP site theme: the colour tokens and the tile / status / explain rules are read from
# reports/static/site.css at build time (one source, so a token change reaches every packet), plus a few
# packet-only rules below. A program keeps only a thin accent stripe and a label edge in its own colour; text
# is never drawn on a program colour. Black text on brat green only.
STATIC = Path(__file__).resolve().parents[1] / "reports" / "static"
SITE_CSS = STATIC / "site.css"
WORDMARK = STATIC / "fip-wordmark.svg"
SITE_RULE_PREFIXES = (".kpi", ".explain", ".status")
FONTS = ("https://fonts.googleapis.com/css2?family=Barlow+Condensed:wght@600;700;800&family=Inter:wght@400;500;600"
         "&family=JetBrains+Mono:wght@400;600;700&display=swap")
PACKET_CSS = """
body { margin: 0; padding: 0 16px 40px; background: var(--bg); color: var(--text); font: 15px/1.55 var(--font); }
a { color: var(--brand-ink); text-underline-offset: 3px; }
.wrap { max-width: 1180px; margin: 0 auto; }
header.top { display: flex; flex-wrap: wrap; align-items: center; gap: 14px 20px; margin: 0 -16px 20px;
  padding: 16px; background: var(--panel); border-bottom: 1px solid var(--line); border-top: 4px solid var(--program); }
header.top svg { height: 40px; width: auto; flex: none; }
.who { display: grid; gap: 4px; min-width: 0; }
.who h1 { margin: 0; font: 700 26px/1.1 var(--font-head); letter-spacing: 0.01em; color: var(--text); }
.who p { margin: 0; color: var(--text-2); font-size: 13.5px; }
.program { justify-self: start; padding: 2px 8px; border-left: 4px solid var(--program); background: var(--panel-2);
  color: var(--text); font: 600 12px/1.5 var(--font-mono); letter-spacing: 0.04em; }
.intro { max-width: 80ch; color: var(--text-2); margin: 0 0 18px; }
h2 { margin: 30px 0 2px; font: 700 21px/1.2 var(--font-head); letter-spacing: 0.02em; color: var(--text); }
.period { margin: 0 0 10px; color: var(--muted); font: 12.5px/1.4 var(--font-mono); }
.kpis { grid-template-columns: repeat(auto-fill, minmax(260px, 1fr)); }
.kpi-compare { color: var(--text); font-size: 13px; line-height: 1.4; }
.kpi-compare b { color: var(--muted); font: 600 10.5px/1.3 var(--font-mono); letter-spacing: 0.06em;
  text-transform: uppercase; }
.note, .empty { color: var(--muted); font-size: 13px; max-width: 90ch; }
.scroll { overflow-x: auto; border: 1px solid var(--line-soft);
  background: var(--panel); }
table.t { border-collapse: collapse; width: 100%; font-size: 13.5px; }
table.t th { background: var(--panel-2); color: var(--text); text-align: left; font: 600 11.5px/1.35 var(--font-mono);
  letter-spacing: 0.04em; border-bottom: 1px solid var(--line); }
table.t th, table.t td { padding: 6px 10px; white-space: nowrap; }
table.t td { border-bottom: 1px solid var(--line-soft); color: var(--text-2); }
table.t .num { text-align: right; font-variant-numeric: tabular-nums; font-family: var(--font-mono); }
@media (max-width: 560px) { .kpis { grid-template-columns: 1fr; } header.top svg { height: 32px; } }
"""


def _top_rules(css: str) -> list[tuple[str, str]]:
    """(selector, full rule) for each top-level rule; rules inside @media and other at-rules are skipped."""
    css = re.sub(r"/\*.*?\*/", "", css, flags=re.S)
    out: list[tuple[str, str]] = []
    depth, start, sel = 0, 0, ""
    for i, ch in enumerate(css):
        if ch == "{":
            if depth == 0:
                sel = css[start:i].strip()
                rule_start = start
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0:
                if not sel.startswith("@"):
                    out.append((sel, css[rule_start:i + 1].strip()))
                start = i + 1
    return out


def theme_css(site_css: Path = SITE_CSS) -> str:
    """The site's dark colour tokens and its KPI tile, status and explain rules, then the packet rules."""
    keep = []
    if site_css.is_file():
        for sel, rule in _top_rules(site_css.read_text(encoding="utf-8")):
            parts = [s.strip() for s in sel.split(",")]
            if sel == ":root" or any(p.startswith(SITE_RULE_PREFIXES) for p in parts):
                keep.append(rule)
    return "\n".join(keep) + PACKET_CSS


def wordmark_svg() -> str:
    """The F&IP wordmark, inline (works with no site next to the packet). Black text on brat green."""
    svg = WORDMARK.read_text(encoding="utf-8") if WORDMARK.is_file() else ""
    svg = re.sub(r"<title[^>]*>.*?</title>", "", svg, flags=re.S)
    return svg.replace('aria-labelledby="fip-t"', 'aria-label="F&amp;IP"').strip() or "<b>F&amp;IP</b>"


# Packet tile -> row in dbt/seeds/core/metric_explanations.csv (the same three lines the site shows).
PACKET_METRICS = {
    "contracts_written": "scorecard_contracts_written", "cancel_rate": "scorecard_cancel_rate",
    "loss_ratio": "scorecard_loss_ratio", "peer_loss_ratio": "peer_loss_ratio",
    "loss_ratio_peer_pct": "loss_ratio_peer_percentile", "contracts_in_force": "contracts_in_force",
    "pending_contracts": "dlr_pending_contracts", "late_fees": "late_fees",
}
PROGRAM_METRICS = {"contracts_written": "contracts_written", "contracts_in_force": "contracts_in_force",
                   "loss_ratio": "posted_loss_ratio", "late_fees": "late_fees",
                   "pending_contracts": "dlr_pending_contracts"}
# Status rules for packet tiles whose seed row has no threshold. PROPOSED in the V7 fix round; the Financial
# Analysts confirm or change them (docs/V3_DESIGN.md 10.8). Loss ratio uses the seed thresholds; late payments
# use seed rule scorecard_watch_late_rate.
PACKET_RULES = {
    "cancel_watch_pts": 0.05,  # Watch up to 5 points above the program's own cancel rate; Needs action beyond
    "peer_watch_pts": 0.10,  # Watch up to 10 points above the peer loss ratio; Needs action beyond
    "rank_watch": 0.50, "rank_alert": 0.80,  # loss-ratio rank: worse than the middle dealer = Watch, top fifth = action
    "pending_watch_days": 60, "pending_alert_days": 90,  # oldest pending contract
}
STATUS_WORDS = {"good": "On track", "warning": "Watch", "critical": "Needs action"}
_ICONS = {
    "good": '<path d="M2 6.5l2.6 2.5L10 3.2" fill="none" stroke="currentColor" stroke-width="2" '
            'stroke-linecap="square"/>',
    "warning": '<path d="M2 6h8" stroke="currentColor" stroke-width="2.4"/>',
    "critical": '<path d="M6 1.2L11.2 10.6H.8z" fill="currentColor"/><path d="M6 4.6v3M6 8.6v.9" stroke="#000" '
                'stroke-width="1.3"/>',
}


@dataclass(frozen=True)
class Bench:
    """What a dealer's tiles are compared with: its program's totals and the seed late-payment rule."""
    program_cancel_rate: float | None = None
    program_written_per_dealer: float | None = None
    late_rate_rule: float = 0.50


def explanations():
    """The site's explanation catalog (reports/explain.py), or None when the report package is not there."""
    try:
        from reports import explain
    except ImportError:
        return None
    return explain.load()


def late_rate_rule(con: duckdb.DuckDBPyConnection) -> float:
    try:
        row = con.execute("select rule_value from seeds.dealer_claim_kpi_rules "
                          "where rule_id = 'scorecard_watch_late_rate'").fetchone()
        if row:
            return float(row[0])
    except duckdb.Error:
        pass
    seed = Path(__file__).resolve().parents[1] / "dbt" / "seeds" / "dealer" / "dealer_claim_kpi_rules.csv"
    if seed.is_file():
        rules = pd.read_csv(seed, dtype=str).set_index("rule_id")["rule_value"]
        if "scorecard_watch_late_rate" in rules:
            return float(rules["scorecard_watch_late_rate"])
    return 0.50


def _num(v: object) -> float | None:
    return None if _missing(v) else float(v)


def _pct(v: float | None) -> str:
    return fmt_value(v, "pct")


def _sum(df: pd.DataFrame | None, col: str) -> float:
    if df is None or not len(df) or col not in df:
        return 0.0
    return float(pd.to_numeric(df[col], errors="coerce").fillna(0).sum())


def kpi_status(key: str, row: dict, tables: dict[str, pd.DataFrame], bench: Bench, cat=None) -> tuple[str, str]:
    """(status, "compared to" sentence) for one dealer packet tile; status is good / warning / critical or ""."""
    v = _num(row.get(key))
    lr = _num(row.get("loss_ratio"))
    r = PACKET_RULES
    if key == "contracts_written":
        avg = bench.program_written_per_dealer
        return ("good" if v else "warning"), (f"Program average per dealer: {avg:,.0f}." if avg is not None else "")
    if key == "cancel_rate":
        base = bench.program_cancel_rate
        if v is None or base is None:
            return "", ""
        st = "good" if v <= base else "warning" if v <= base + r["cancel_watch_pts"] else "critical"
        return st, f"Program average: {_pct(base)}. Watch up to {r['cancel_watch_pts'] * 100:.0f} points above it."
    if key == "loss_ratio":
        mid = PACKET_METRICS[key]
        e = cat.metrics.get(mid) if cat is not None else None
        if e is None or e.watch_threshold is None or e.alert_threshold is None:
            return "", ""
        return cat.status(mid, v), (f"Watch above {_pct(e.watch_threshold)}, "
                                    f"needs action above {_pct(e.alert_threshold)}.")
    if key == "peer_loss_ratio":
        scope = row.get("peer_scope") or "peers"
        if v is None:
            return "", "Not published: too few similar dealers to share it safely."
        if lr is None:
            return "", f"Average for {scope}."
        st = "good" if lr <= v else "warning" if lr <= v + r["peer_watch_pts"] else "critical"
        return st, f"Yours: {_pct(lr)} vs {_pct(v)} for {scope}. Lower is better."
    if key == "loss_ratio_peer_pct":
        if v is None:
            return "", "Not ranked yet: not enough earned reserve in 12 months."
        st = "good" if v <= r["rank_watch"] else "warning" if v < r["rank_alert"] else "critical"
        return st, "50% is the middle dealer; 0% is the lowest loss ratio. Lower is better."
    if key == "contracts_in_force":
        return ("good" if v else "warning"), (
            f"Written in the last 12 months: {fmt_value(row.get('contracts_written'), 'int')}.")
    if key == "pending_contracts":
        pend = tables.get("pending_contracts")
        oldest = (float(pd.to_numeric(pend["days_pending"], errors="coerce").max())
                  if pend is not None and len(pend) else 0.0)
        st = ("good" if oldest <= r["pending_watch_days"] else
              "warning" if oldest <= r["pending_alert_days"] else "critical")
        return st, (f"Oldest one: {oldest:,.0f} days. Watch past {r['pending_watch_days']} days, "
                    f"needs action past {r['pending_alert_days']}." if v else "Nothing waiting.")
    if key == "late_fees":
        aging = tables.get("remittance_aging")
        paid, late = _sum(aging, "t12_remittances"), _sum(aging, "t12_late_remittances")
        share = late / paid if paid else 0.0
        st = "good" if not v else "critical" if share > bench.late_rate_rule else "warning"
        return st, (f"{late:,.0f} of {paid:,.0f} payments were late in 12 months. "
                    f"Needs action when more than {bench.late_rate_rule:.0%} are late.")
    return "", ""


def _status_html(status: str) -> str:
    if not status:
        return ""
    return (f"<p class='status status-{status}'><svg class='status-icon' viewBox='0 0 12 12' aria-hidden='true'>"
            f"{_ICONS[status]}</svg><span>{STATUS_WORDS[status]}</span></p>")


def _explain_html(cat, metric_id: str) -> str:
    e = cat.metrics.get(metric_id) if cat is not None and metric_id else None
    if e is None:
        return ""
    lines = "".join(f"<div><dt>{html.escape(label)}</dt><dd>{html.escape(text)}</dd></div>"
                    for label, text in e.lines() if text)
    return f"<dl class='explain'>{lines}</dl>"


def _kpi_strip(row: dict, cols: tuple[Col, ...], note: str = "", status=None, cat=None,
               metrics: dict[str, str] | None = None) -> str:
    """KPI tiles: label, big value, status word + shape, a "compared to" line and the three explain lines."""
    metrics = metrics or {}
    tiles = []
    for c in cols:
        value = fmt_value(row.get(c.key), c.fmt)
        st, cmp = status(c.key) if status else ("", "")
        mid = metrics.get(c.key, "")
        val = ("<p class='kpi-value kpi-empty'>Not enough data yet</p>" if value == "-"
               else f"<p class='kpi-value'>{html.escape(value)}</p>")
        tiles.append(
            (f'<div class="kpi" data-metric="{html.escape(mid)}">' if mid else "<div class='kpi'>")
            + f"<p class='kpi-label'>{html.escape(c.label)}</p>{val}"
            + (f"<div class='kpi-foot'>{_status_html(st)}</div>" if st else "")
            + (f"<p class='kpi-compare'><b>Compared to</b><br>{html.escape(cmp)}</p>" if cmp else "")
            + f"{_explain_html(cat, mid)}</div>")
    return f"<div class='kpis'>{''.join(tiles)}</div>" + (f"<p class='note'>{html.escape(note)}</p>" if note else "")


def _page(title: str, brand: str, tint: str, header: str, body: str) -> str:
    """``brand`` is the program's own colour, used only for the thin top stripe and the label edge. ``tint`` is
    kept for callers; the dark theme never draws text on a program colour."""
    del tint
    program = brand if re.fullmatch(r"#[0-9A-Fa-f]{3,8}", brand or "") else "#8ACE00"
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)}</title>
<meta name="color-scheme" content="dark">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="{FONTS}">
<style>
{theme_css()}
:root {{ --program: {program}; }}
</style></head><body><div class="wrap">
<header class="top">{wordmark_svg()}{header}</header>
{body}
</div></body></html>
"""


def _section(title: str, period: str, inner: str) -> str:
    return f"<h2>{html.escape(title)}</h2>\n<p class='period'>{html.escape(period)}</p>\n{inner}"


def _peer_note(kpi: pd.DataFrame) -> str:
    if not len(kpi):
        return ""
    scope = kpi["peer_scope"].iloc[0] if "peer_scope" in kpi else ""
    return ("How these are worked out: loss ratio = incurred claims / earned claims reserve. Peer loss ratio: VSC "
            f"contracts, {scope or 'not published (suppressed)'}, from the suppressed T12 peer benchmark. "
            "Peer percentile ranks this dealer's T12 loss ratio among all dealers with enough earned reserve "
            "(0% = lowest).")


def _header(kicker: str, title: str, detail: str) -> str:
    return (f"<div class='who'><span class='program'>{html.escape(kicker)}</span><h1>{html.escape(title)}</h1>"
            f"<p>{detail}</p></div>")


def _html(p: Packet, tables: dict[str, pd.DataFrame], per: Periods, bench: Bench | None = None, cat=None) -> str:
    bench = bench or Bench()
    kpi = tables["kpi"]
    row = kpi.iloc[0].to_dict() if len(kpi) else {}
    parts = [
        "<p class='intro'>Your store's last 12 complete months on one page. Each headline number has a status word, "
        "what it is compared with, and three short lines on what it means and why it is asked for. The tables "
        "below hold the detail; the same tables are in packet.xlsx.</p>",
        _section("Headline KPIs", per.label("t12"),
                 _kpi_strip(row, KPI_COLS, _peer_note(kpi), lambda k: kpi_status(k, row, tables, bench, cat),
                            cat, PACKET_METRICS) if len(kpi)
                 else "<p class='empty'>No activity in this period.</p>")]
    for name, spec in TABLE_SPECS.items():
        if name != "kpi":
            parts.append(_section(spec.title, per.label(spec.period), _table_html(tables[name], spec.cols, per)))
    header = _header(p.reportable_name, f"Dealer packet: {p.dealer_number} {p.dealer_name or ''}".strip(),
                     f"F&amp;I Data Platform &middot; {html.escape(per.t12_label)} &middot; data as of "
                     f"{per.as_of:%Y-%m-%d} &middot; <a href='../index.html'>program packet</a>")
    return _page(f"{p.dealer_number} packet", p.primary_color, p.secondary_color, header, "\n".join(parts))


def _xlsx(path: Path, sheets: list[tuple[str, str, str, pd.DataFrame, tuple[Col, ...]]], per: Periods) -> bool:
    """sheets: (sheet name, title, period label, frame, columns). Readable headers and number formats."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError:
        return False
    wb = Workbook()
    wb.remove(wb.active)
    for sheet, title, period, df, cols in sheets:
        ws = wb.create_sheet(sheet[:31])
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([period])
        ws.append([])
        ws.append(_labels(cols, per))
        for cell in ws[4]:
            cell.font = Font(bold=True)
        for rec in df.to_dict("records"):
            row = []
            for c in cols:
                v = rec.get(c.key)
                if _missing(v):
                    v = None
                elif c.fmt == "int":
                    v = int(v)
                elif c.fmt in ("money", "pct", "ratio"):
                    # round to what the format shows (+ 0.0 turns -0.0 into 0.0): Excel picks the negative
                    # section from the stored value, so -0.004 would otherwise display as -$0.00
                    v = round(float(v), XLSX_DIGITS[c.fmt]) + 0.0
                elif c.fmt == "date":
                    v = pd.Timestamp(v).date()
                row.append(v)
            ws.append(row)
        for j, c in enumerate(cols, start=1):
            letter = ws.cell(row=4, column=j).column_letter
            ws.column_dimensions[letter].width = max(12, len(_labels(cols, per)[j - 1]) + 2)
            if c.fmt in XLSX_FORMATS:
                for r in range(5, ws.max_row + 1):
                    ws.cell(row=r, column=j).number_format = XLSX_FORMATS[c.fmt]
        # header is row 4 (title, period, blank above it): freeze it and filter every column (V6 edit 10)
        ws.freeze_panes = "A5"
        if len(df):
            ws.auto_filter.ref = f"A4:{ws.cell(row=4, column=len(cols)).column_letter}{ws.max_row}"
        else:
            ws.append(["No rows for this period."])
    wb.save(path)
    return True


def _xlsx_async(path: Path, sheets: list[tuple[str, str, str, pd.DataFrame, tuple[Col, ...]]], per: Periods) -> bool:
    """``_xlsx`` in the background when build_packets has a pool open (pipeline/workpool.py)."""
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    workpool.run(_xlsx, path, sheets, per)
    return True


def write_packet(out_dir: Path, p: Packet, tables: dict[str, pd.DataFrame], per: Periods,
                 bench: Bench | None = None, cat=None) -> Path:
    d = out_dir / p.tenant_code / p.dealer_number
    d.mkdir(parents=True, exist_ok=True)
    for name, df in tables.items():
        df.to_csv(d / f"{name}.csv", index=False)
    (d / "index.html").write_text(_html(p, tables, per, bench, cat), encoding="utf-8")
    sheets = [(name, spec.title, per.label(spec.period), tables[name], spec.cols) for name, spec in TABLE_SPECS.items()]
    has_xlsx = _xlsx_async(d / "packet.xlsx", sheets, per)
    meta = {"tenant_code": p.tenant_code, "dealer_number": p.dealer_number, "dealer_key": p.dealer_key,
            "as_of": str(per.as_of), "period": per.t12_label, "tables": {k: len(v) for k, v in tables.items()},
            "xlsx": has_xlsx}
    (d / "packet.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    return d


# --------------------------------------------------------------------------- program packets
def program_tables(con: duckdb.DuckDBPyConnection, members: list[tuple[Packet, dict[str, pd.DataFrame]]],
                   per: Periods) -> dict[str, pd.DataFrame]:
    """Program (tenant) tables built from its dealer packets' KPI rows, so totals tie by construction,
    plus claims.mart_claims_kpi at tenant grain."""
    rows = []
    for p, t in members:
        k = t["kpi"].iloc[0].to_dict() if len(t["kpi"]) else {}
        rows.append({"dealer_number": p.dealer_number, "dealer_name": p.dealer_name, **k})
    dealers = pd.DataFrame(rows, columns=["dealer_number", "dealer_name", *PROGRAM_SUM_KEYS, "cancel_rate",
                                          "loss_ratio", "loss_ratio_peer_pct"]) if rows else pd.DataFrame()
    total: dict[str, object] = {"dealers": len(members)}
    for key in PROGRAM_SUM_KEYS:
        total[key] = float(pd.to_numeric(dealers[key], errors="coerce").fillna(0).sum()) if len(dealers) else 0.0
    ecr = total["earned_claims_reserve"]
    total["loss_ratio"] = total["incurred_claims"] / ecr if ecr else None
    tenant_key = members[0][0].tenant_key if members else -1
    claims = con.execute(PROGRAM_CLAIMS.sql, {"tenant_key": tenant_key, "t12_start": per.t12_start,
                                              "t12_end": per.t12_end}).df()
    return {"program_kpi": pd.DataFrame([total]), "dealers": dealers, "program_claims": claims}


def _program_status(tables: dict[str, pd.DataFrame], cat):
    """Program tiles get a status only where the seed has a rule (loss ratio)."""
    row = tables["program_kpi"].iloc[0].to_dict()

    def status(key: str) -> tuple[str, str]:
        if key != "loss_ratio" or cat is None:
            return "", ""
        e = cat.metrics.get(PROGRAM_METRICS[key])
        if e is None or e.watch_threshold is None or e.alert_threshold is None:
            return "", ""
        return cat.status(e.metric_id, _num(row.get(key))), (
            f"Watch above {_pct(e.watch_threshold)}, needs action above {_pct(e.alert_threshold)}.")
    return status


def bench_for(members: list[tuple[Packet, dict[str, pd.DataFrame]]], late_rule: float) -> Bench:
    """A program's own cancel rate and average contracts per dealer, from its dealer packets' KPI rows."""
    written = sum(_sum(t["kpi"], "contracts_written") for _, t in members)
    cancelled = sum(_sum(t["kpi"], "contracts_cancelled") for _, t in members)
    return Bench(cancelled / written if written else None, written / len(members) if members else None, late_rule)


def write_program(out_dir: Path, tenant: Packet, tables: dict[str, pd.DataFrame], per: Periods, cat=None) -> Path:
    d = out_dir / tenant.tenant_code
    d.mkdir(parents=True, exist_ok=True)
    for name, df in tables.items():
        df.to_csv(d / f"program_{name}.csv", index=False)
    links = {n: f"{n}/index.html" for n in tables["dealers"].get("dealer_number", [])}
    body = "\n".join([
        _section("Program totals", per.label("program_t12"),
                 _kpi_strip(tables["program_kpi"].iloc[0].to_dict(), PROGRAM_KPI_COLS,
                            "Written, earned, incurred and late fees are T12 sums of the dealer packets below; "
                            f"in force and pending are as of {per.as_of:%Y-%m-%d}. "
                            "Loss ratio = incurred / earned claims reserve.", _program_status(tables, cat),
                            cat, PROGRAM_METRICS)),
        _section("Dealers", per.label("t12"), _table_html(tables["dealers"], PROGRAM_DEALER_COLS, per, links)),
        _section(PROGRAM_CLAIMS.title, per.label("program_t12"),
                 _table_html(tables["program_claims"], PROGRAM_CLAIMS.cols, per)),
    ])
    header = _header(tenant.reportable_name, f"Partner program packet: {tenant.tenant_code}",
                     f"F&amp;I Data Platform &middot; {len(tables['dealers'])} dealers &middot; "
                     f"{html.escape(per.t12_label)} &middot; data as of {per.as_of:%Y-%m-%d}")
    (d / "index.html").write_text(_page(f"{tenant.tenant_code} program packet", tenant.primary_color,
                                        tenant.secondary_color, header, body), encoding="utf-8")
    has_xlsx = _xlsx_async(d / "program.xlsx", [
        ("program", "Program totals", per.label("program_t12"), tables["program_kpi"], PROGRAM_KPI_COLS),
        ("dealers", "Dealers", per.label("t12"), tables["dealers"], PROGRAM_DEALER_COLS),
        ("claims", PROGRAM_CLAIMS.title, per.label("program_t12"), tables["program_claims"], PROGRAM_CLAIMS.cols),
    ], per)
    meta = {"tenant_code": tenant.tenant_code, "as_of": str(per.as_of), "period": per.t12_label,
            "dealers": tables["dealers"].get("dealer_number", pd.Series(dtype=str)).tolist(), "xlsx": has_xlsx}
    (d / "program.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    return d


def build_packets(con: duckdb.DuckDBPyConnection, out_dir: Path, peer: pd.DataFrame | None = None) -> list[Path]:
    """Write every dealer packet, then one program packet per tenant. Returns the dealer packet dirs."""
    per = periods(con)
    written: list[Path] = []
    by_tenant: dict[str, list[tuple[Packet, dict[str, pd.DataFrame]]]] = {}
    cat, late_rule = explanations(), late_rate_rule(con)
    for p in packets(con):  # tables first: each dealer's tiles are compared with its whole program
        by_tenant.setdefault(p.tenant_code, []).append((p, packet_tables(con, p, per, peer)))
    with workpool.workbook_pool():  # packet workbooks are written in parallel; the block waits for all
        for members in by_tenant.values():
            bench = bench_for(members, late_rule)
            for p, tables in members:
                written.append(write_packet(out_dir, p, tables, per, bench, cat))
            write_program(out_dir, members[0][0], program_tables(con, members, per), per, cat)
    log.info("dealer packets: %d, program packets: %d -> %s", len(written), len(by_tenant), out_dir)
    return written
